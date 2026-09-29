from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional

from flask import Blueprint, jsonify, request, current_app, make_response

from app import db
from app.auth.decorators import login_required

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT

blacklist_report_bp = Blueprint('blacklist_report', __name__, url_prefix='/MasterData')

SEARCH_BY_CUSTOMER_NO = '1'
SEARCH_BY_CUSTOMER_NAME = '2'
SEARCH_BY_REASON = '3'

STATUS_ACTIVE = '0'
STATUS_DELETED = '1'

_FROM_JOINS = """
    FROM MCMOWatchList mw WITH (NOLOCK)
        LEFT JOIN APP_USRMST au WITH (NOLOCK) ON au.USERNAME  = mw.CreateUser
        LEFT JOIN UMFBRANCH  br WITH (NOLOCK) ON br.BRANCH_CD = au.BRANCHCODE
        LEFT JOIN cafless    cf WITH (NOLOCK) ON cf.APLESS    = mw.apless
        LEFT JOIN MSTCOMBO   mt WITH (NOLOCK) ON mt.VALUE      = mw.Reason_Category
                                            AND mt.CAT_VALUE  = 'BLACKLIST_CATEGORY'
"""

_SELECT_COLS = """
    SELECT
        mw.Name       AS Name,
        cf.APLESS     AS APLESS,
        mw.Address    AS Address,
        mw.ID_Card    AS ID_Card,
        mw.NPWP       AS NPWP,
        mw.Phone      AS Phone,
        mt.DESC_VALUE AS Reason,
        mw.CreateUser AS CreateUser,
        mw.CreateDate AS CreateDate,
        mw.Flagdel    AS Flagdel,
        mw.DeleteBy   AS DeleteBy,
        mw.attachment AS Attachment,
        cf.LESSEE_TP  AS LesseeType
"""

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _exec_scalar(sql: str, params: list):
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        row = cursor.fetchone()
        return row[0] if row else 0
    finally:
        cursor.close()

def _resolve_status(existing: str, deleted: str) -> Optional[str]:
   
    existing_on = bool(existing)
    deleted_on = bool(deleted)
    if existing_on == deleted_on:
        return None
    return STATUS_DELETED if deleted_on else STATUS_ACTIVE

def _build_where(as_of: date, searchby: str, searchbytext: str,
                  reason_cat: str, status: Optional[str]) -> tuple[list[str], list]:
    where = ["CONVERT(date, mw.CreateDate) <= ?", "mw.Flag_type = 'B'"]
    params: list = [as_of]

    if searchby == SEARCH_BY_CUSTOMER_NO and searchbytext:
        where.append("cf.APLESS = ?")
        params.append(searchbytext)
    elif searchby == SEARCH_BY_CUSTOMER_NAME and searchbytext:
        where.append("mw.Name LIKE ?")
        params.append(f"%{searchbytext}%")
    elif searchby == SEARCH_BY_REASON and reason_cat:
        where.append("mw.Reason_Category = ?")
        params.append(reason_cat)

    if status is not None:
        where.append("mw.Flagdel = ?")
        params.append(status)

    return where, params

def _query_blacklist(as_of, searchby, searchbytext, reason_cat, existing, deleted,
                      page: Optional[int] = None, page_size: Optional[int] = None):
   
    status = _resolve_status(existing, deleted)
    where, params = _build_where(as_of, searchby, searchbytext, reason_cat, status)
    where_sql = " AND ".join(where)

    total = _exec_scalar(f"SELECT COUNT(*) {_FROM_JOINS} WHERE {where_sql}", params)

    data_sql = f"{_SELECT_COLS} {_FROM_JOINS} WHERE {where_sql} ORDER BY mw.Name ASC"
    data_params = list(params)

    if page is not None and page_size is not None:
        data_sql += " OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
        data_params += [(page - 1) * page_size, page_size]

    rows = _exec_rows(data_sql, data_params)
    return total, rows

def _fmt_date(val) -> str:
    if isinstance(val, (date, datetime)):
        return val.strftime('%d-%m-%Y')
    return str(val) if val is not None else ''

def _serialize_row(r: dict, no: int) -> dict:
    lessee_type = (r.get('LesseeType') or '').strip()
    identitas = r.get('ID_Card') if lessee_type == 'PR' else r.get('NPWP')
    is_deleted = str(r.get('Flagdel')) == STATUS_DELETED

    return {
        'no': no,
        'name': r.get('Name') or '-',
        'customer_no': r.get('APLESS') or '-',
        'address': r.get('Address') or '-',
        'identitas': identitas or '-',
        'phone': r.get('Phone') or '-',
        'reason': r.get('Reason') or '-',
        'create_user': r.get('CreateUser') or '-',
        'create_date': _fmt_date(r.get('CreateDate')),
        'status': 'Deleted' if is_deleted else 'Active',
        'delete_by': (r.get('DeleteBy') or '-') if is_deleted else '-',
        'attachment': r.get('Attachment') or None,
        'lessee_type': lessee_type,
    }

def _describe_search(searchby: str, searchbytext: str, reason_desc: str) -> str:
    if searchby == SEARCH_BY_CUSTOMER_NO:
        return f'Customer No / {searchbytext}'
    if searchby == SEARCH_BY_CUSTOMER_NAME:
        return f'Customer Name / {searchbytext}'
    if searchby == SEARCH_BY_REASON:
        return f'Reason / {reason_desc}'
    return 'All'

def _describe_view(existing: str, deleted: str) -> str:
    parts = []
    if existing:
        parts.append('Existing in Blacklist Report')
    if deleted:
        parts.append('Data Deleted from Blacklist Report')
    return ' / '.join(parts) if parts else 'All'

def _reason_description(reason_cat: str) -> str:
    if not reason_cat:
        return ''
    rows = _exec_rows(
        "SELECT DESC_VALUE FROM MSTCOMBO WHERE CAT_VALUE = ? AND VALUE = ?",
        ['BLACKLIST_CATEGORY', reason_cat],
    )
    return rows[0]['DESC_VALUE'] if rows else ''

def _parse_request_args(args) -> dict:
    as_of_raw = args.get('as_of', '').strip()
    if not as_of_raw:
        raise ValueError('as_of is required')
    try:
        as_of = datetime.strptime(as_of_raw, '%Y-%m-%d').date()
    except ValueError as exc:
        raise ValueError('as_of must be in YYYY-MM-DD format') from exc

    return {
        'as_of': as_of,
        'as_of_raw': as_of_raw,
        'searchby': args.get('searchby', '').strip(),
        'searchbytext': args.get('searchbytext', '').strip(),
        'reason_cat': args.get('reason_cat', '').strip(),
        'existing': args.get('existing', '').strip(),
        'deleted': args.get('deleted', '').strip(),
    }

_THIN = Side(style='thin', color='000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT_DATA = Font(name='Tahoma', size=8)
_FONT_HEADER = Font(name='Tahoma', size=8, bold=True)

_XLS_COLUMNS = [
    ('No', 'no'),
    ('Customer Name', 'name'),
    ('Customer Number', 'customer_no'),
    ('Address', 'address'),
    ('No Identitas', 'identitas'),
    ('Phone', 'phone'),
    ('Reason', 'reason'),
    ('Create By', 'create_user'),
    ('Create Date', 'create_date'),
    ('Status', 'status'),
    ('Delete By', 'delete_by'),
    ('Attachment', 'attachment'),
]
_LAST_COL = get_column_letter(len(_XLS_COLUMNS))

def _apply_default(cell):
    cell.border = _BORDER
    cell.font = _FONT_DATA

def _apply_header(cell):
    cell.border = _BORDER
    cell.font = _FONT_HEADER
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def _build_excel(rows: list[dict], as_of: date, search_desc: str, view_desc: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'BlacklistReport'

    ws.merge_cells(f'A1:{_LAST_COL}1')
    cell_title = ws['A1']
    cell_title.value = 'BLACKLIST REPORT'
    cell_title.font = Font(name='Tahoma', size=12, bold=True)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 20

    ws.merge_cells(f'A2:{_LAST_COL}2')
    cell_sub = ws['A2']
    cell_sub.value = f'As of {as_of.strftime("%d %b %Y")}'
    cell_sub.font = Font(name='Tahoma', size=10, bold=True)
    cell_sub.alignment = Alignment(horizontal='center', vertical='center')

    ws['A4'] = 'Branch'
    ws['B4'] = ':'
    ws['C4'] = 'Consolidate'
    ws['A5'] = 'Search by'
    ws['B5'] = ':'
    ws['C5'] = search_desc
    ws['A6'] = 'View'
    ws['B6'] = ':'
    ws['C6'] = view_desc
    for row in ws['A4:C6']:
        for cell in row:
            cell.font = Font(name='Tahoma', size=9)

    header_row = 8
    for idx, (label, _) in enumerate(_XLS_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=label)
        _apply_header(cell)
        ws.column_dimensions[get_column_letter(idx)].width = max(len(label) + 2, 12)

    for r_idx, row in enumerate(rows, start=1):
        excel_row = header_row + r_idx
        for c_idx, (_, key) in enumerate(_XLS_COLUMNS, start=1):
            cell = ws.cell(row=excel_row, column=c_idx, value=row.get(key, ''))
            _apply_default(cell)

    ws.freeze_panes = f'A{header_row + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

_WORD_HEADERS = ['No', 'Customer Name', 'Customer Number', 'Address', 'No Identitas',
                  'Phone', 'Reason', 'Create By', 'Create Date', 'Status', 'Delete By',
                  'Attachment']
_WORD_KEYS = ['no', 'name', 'customer_no', 'address', 'identitas', 'phone', 'reason',
              'create_user', 'create_date', 'status', 'delete_by', 'attachment']


def _build_word(rows: list[dict], as_of: date, search_desc: str, view_desc: str) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run('BLACKLIST REPORT')
    title_run.bold = True
    title_run.font.size = Pt(14)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f'As of {as_of.strftime("%d %b %Y")}')
    subtitle_run.bold = True

    meta = doc.add_table(rows=3, cols=3)
    meta.cell(0, 0).text = 'Branch'
    meta.cell(0, 1).text = ':'
    meta.cell(0, 2).text = 'Consolidate'
    meta.cell(1, 0).text = 'Search by'
    meta.cell(1, 1).text = ':'
    meta.cell(1, 2).text = search_desc
    meta.cell(2, 0).text = 'View'
    meta.cell(2, 1).text = ':'
    meta.cell(2, 2).text = view_desc

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(_WORD_HEADERS))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(_WORD_HEADERS):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8)

    for row in rows:
        cells = table.add_row().cells
        for i, key in enumerate(_WORD_KEYS):
            value = row.get(key) or ''
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()

@blacklist_report_bp.route('/blacklist-report/reasons')
@login_required
def list_reasons():
    try:
        rows = _exec_rows(
            "SELECT VALUE, DESC_VALUE FROM MSTCOMBO WHERE CAT_VALUE = ? ORDER BY DESC_VALUE",
            ['BLACKLIST_CATEGORY'],
        )
        return jsonify(rows)
    except Exception as e:
        current_app.logger.exception("Blacklist reason list error")
        return jsonify({'error': str(e)}), 500

@blacklist_report_bp.route('/blacklist-report/preview')
@login_required
def preview():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        page = max(int(request.args.get('page', 1)), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = min(max(int(request.args.get('page_size', 20)), 1), 200)
    except (TypeError, ValueError):
        page_size = 20

    try:
        total, rows = _query_blacklist(
            p['as_of'], p['searchby'], p['searchbytext'], p['reason_cat'],
            p['existing'], p['deleted'], page=page, page_size=page_size,
        )
        offset = (page - 1) * page_size
        data = [_serialize_row(r, offset + i) for i, r in enumerate(rows, start=1)]

        return jsonify({
            'data': data,
            'page': page,
            'page_size': page_size,
            'total': total,
            'total_pages': max((total + page_size - 1) // page_size, 1),
            'search_desc': _describe_search(
                p['searchby'], p['searchbytext'], _reason_description(p['reason_cat'])
            ),
            'view_desc': _describe_view(p['existing'], p['deleted']),
            'as_of': p['as_of_raw'],
        })
    except Exception as e:
        current_app.logger.exception("Blacklist preview error")
        return jsonify({'error': str(e)}), 500

@blacklist_report_bp.route('/blacklist-report/excel')
@login_required
def export_excel():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        _, rows = _query_blacklist(
            p['as_of'], p['searchby'], p['searchbytext'], p['reason_cat'],
            p['existing'], p['deleted'],
        )
        data = [_serialize_row(r, i) for i, r in enumerate(rows, start=1)]
        search_desc = _describe_search(
            p['searchby'], p['searchbytext'], _reason_description(p['reason_cat'])
        )
        view_desc = _describe_view(p['existing'], p['deleted'])

        xlsx_bytes = _build_excel(data, p['as_of'], search_desc, view_desc)
        filename = f"BlacklistReport_{p['as_of_raw']}.xlsx"

        response = make_response(xlsx_bytes)
        response.headers['Content-Type'] = (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        current_app.logger.exception("Blacklist Excel export error")
        return jsonify({'error': str(e)}), 500

@blacklist_report_bp.route('/blacklist-report/word')
@login_required
def export_word():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        _, rows = _query_blacklist(
            p['as_of'], p['searchby'], p['searchbytext'], p['reason_cat'],
            p['existing'], p['deleted'],
        )
        data = [_serialize_row(r, i) for i, r in enumerate(rows, start=1)]
        search_desc = _describe_search(
            p['searchby'], p['searchbytext'], _reason_description(p['reason_cat'])
        )
        view_desc = _describe_view(p['existing'], p['deleted'])

        docx_bytes = _build_word(data, p['as_of'], search_desc, view_desc)
        filename = f"BlacklistReport_{p['as_of_raw']}.docx"

        response = make_response(docx_bytes)
        response.headers['Content-Type'] = (
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        current_app.logger.exception("Blacklist Word export error")
        return jsonify({'error': str(e)}), 500