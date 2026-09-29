from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional

from flask import Blueprint, jsonify, request, current_app, make_response

from app import db
from app.auth.decorators import login_required

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT

watchlist_report_bp = Blueprint('watchlist_report', __name__, url_prefix='/MasterData')

SEARCH_BY_CUSTOMER_NO = '1'
SEARCH_BY_CUSTOMER_NAME = '2'

STATUS_ACTIVE = '0'
STATUS_DELETED = '1'

_FROM_JOINS = """
    FROM MCMOWatchList mw WITH (NOLOCK)
    LEFT JOIN APP_USRMST au WITH (NOLOCK) ON au.USERNAME  = mw.CreateUser
    LEFT JOIN UMFBRANCH  br WITH (NOLOCK) ON br.BRANCH_CD = au.BRANCHCODE
    LEFT JOIN cafless    cf WITH (NOLOCK) ON cf.APLESS    = mw.apless
    LEFT JOIN MSTCOMBO   ms WITH (NOLOCK) ON ms.CAT_VALUE = 'BLACKLIST_CATEGORY'
                                          AND ms.VALUE    = mw.Reason_Category
"""

_SELECT_COLS = """
    SELECT
        mw.Name       AS Name,
        cf.APLESS     AS APLESS,
        mw.Address    AS Address,
        mw.ID_Card    AS ID_Card,
        mw.NPWP       AS NPWP,
        mw.Phone      AS Phone,
        CASE WHEN mw.Reason_Category = '8' THEN ms.DESC_VALUE ELSE mw.Reason END AS Reason,
        mw.CreateUser AS CreateUser,
        mw.CreateDate AS CreateDate,
        mw.Flagdel    AS Flagdel,
        mw.DeleteBy   AS DeleteBy,
        mw.attachment AS Attachment,
        cf.LESSEE_TP  AS LesseeType
"""

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _exec_scalar(sql: str, params: list):
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        row = cursor.fetchone()
        return row[0] if row else 0
    finally:
        cursor.close()

def _resolve_status(existing: str, deleted: str) -> Optional[str]:
    existing_on = bool(existing)
    deleted_on = bool(deleted)
    if existing_on == deleted_on:
        return None
    return STATUS_DELETED if deleted_on else STATUS_ACTIVE

def _build_where(as_of: date, searchby: str, searchbytext: str,
                  status: Optional[str]) -> tuple[list[str], list]:
    where = ["CONVERT(date, mw.CreateDate) <= ?", "mw.Flag_type = 'W'"]
    params: list = [as_of]

    if searchby == SEARCH_BY_CUSTOMER_NO and searchbytext:
        where.append("cf.APLESS = ?")
        params.append(searchbytext)
    elif searchby == SEARCH_BY_CUSTOMER_NAME and searchbytext:
        where.append("mw.Name = ?")
        params.append(searchbytext)

    if status is not None:
        where.append("mw.Flagdel = ?")
        params.append(status)

    return where, params

def _query_watchlist(as_of, searchby, searchbytext, existing, deleted,
                      page: Optional[int] = None, page_size: Optional[int] = None):
    status = _resolve_status(existing, deleted)
    where, params = _build_where(as_of, searchby, searchbytext, status)
    where_sql = " AND ".join(where)

    total = _exec_scalar(f"SELECT COUNT(*) {_FROM_JOINS} WHERE {where_sql}", params)

    data_sql = f"{_SELECT_COLS} {_FROM_JOINS} WHERE {where_sql} ORDER BY mw.Name ASC"
    data_params = list(params)

    if page is not None and page_size is not None:
        data_sql += " OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
        data_params += [(page - 1) * page_size, page_size]

    rows = _exec_rows(data_sql, data_params)
    return total, rows

def _fmt_date(val) -> str:
    if isinstance(val, (date, datetime)):
        return val.strftime('%d-%m-%Y')
    return str(val) if val is not None else ''

def _serialize_row(r: dict, no: int) -> dict:
    lessee_type = (r.get('LesseeType') or '').strip()
    identitas = r.get('ID_Card') if lessee_type == 'PR' else r.get('NPWP')
    is_deleted = str(r.get('Flagdel')) == STATUS_DELETED

    return {
        'no': no,
        'name': r.get('Name') or '-',
        'customer_no': r.get('APLESS') or '-',
        'address': r.get('Address') or '-',
        'identitas': identitas or '-',
        'phone': r.get('Phone') or '-',
        'reason': r.get('Reason') or '-',
        'create_user': r.get('CreateUser') or '-',
        'create_date': _fmt_date(r.get('CreateDate')),
        'status': 'Deleted' if is_deleted else 'Active',
        'delete_by': (r.get('DeleteBy') or '-') if is_deleted else '-',
        'attachment': r.get('Attachment') or None,
        'lessee_type': lessee_type,
    }

def _describe_search(searchby: str, searchbytext: str) -> str:
    if searchby == SEARCH_BY_CUSTOMER_NO:
        return f'Customer No / {searchbytext}'
    if searchby == SEARCH_BY_CUSTOMER_NAME:
        return f'Customer Name / {searchbytext}'
    return 'All'

def _describe_view(existing: str, deleted: str) -> str:
    parts = []
    if existing:
        parts.append('Existing in Watchlist Report')
    if deleted:
        parts.append('Data Deleted from Watchlist Report')
    return ' / '.join(parts) if parts else 'All'

def _parse_request_args(args) -> dict:
    as_of_raw = args.get('as_of', '').strip()
    if not as_of_raw:
        raise ValueError('as_of is required')
    try:
        as_of = datetime.strptime(as_of_raw, '%Y-%m-%d').date()
    except ValueError as exc:
        raise ValueError('as_of must be in YYYY-MM-DD format') from exc

    return {
        'as_of': as_of,
        'as_of_raw': as_of_raw,
        'searchby': args.get('searchby', '').strip(),
        'searchbytext': args.get('searchbytext', '').strip(),
        'existing': args.get('existing', '').strip(),
        'deleted': args.get('deleted', '').strip(),
    }
    
_THIN = Side(style='thin', color='000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT_DATA = Font(name='Tahoma', size=8)
_FONT_HEADER = Font(name='Tahoma', size=8, bold=True)

_XLS_COLUMNS = [
    ('No', 'no'),
    ('Customer Name', 'name'),
    ('Customer Number', 'customer_no'),
    ('Address', 'address'),
    ('No Identitas', 'identitas'),
    ('Phone', 'phone'),
    ('Reason', 'reason'),
    ('Create By', 'create_user'),
    ('Create Date', 'create_date'),
    ('Status', 'status'),
    ('Delete By', 'delete_by'),
    ('Attachment', 'attachment'),
]
_LAST_COL = get_column_letter(len(_XLS_COLUMNS))

def _apply_default(cell):
    cell.border = _BORDER
    cell.font = _FONT_DATA

def _apply_header(cell):
    cell.border = _BORDER
    cell.font = _FONT_HEADER
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def _build_excel(rows: list[dict], as_of: date, search_desc: str, view_desc: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'WatchlistReport'

    ws.merge_cells(f'A1:{_LAST_COL}1')
    cell_title = ws['A1']
    cell_title.value = 'WATCHLIST REPORT'
    cell_title.font = Font(name='Tahoma', size=12, bold=True)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 20

    ws.merge_cells(f'A2:{_LAST_COL}2')
    cell_sub = ws['A2']
    cell_sub.value = f'As of {as_of.strftime("%d %b %Y")}'
    cell_sub.font = Font(name='Tahoma', size=10, bold=True)
    cell_sub.alignment = Alignment(horizontal='center', vertical='center')

    ws['A4'] = 'Branch'
    ws['B4'] = ':'
    ws['C4'] = 'Consolidate'
    ws['A5'] = 'Search by'
    ws['B5'] = ':'
    ws['C5'] = search_desc
    ws['A6'] = 'View'
    ws['B6'] = ':'
    ws['C6'] = view_desc
    for row in ws['A4:C6']:
        for cell in row:
            cell.font = Font(name='Tahoma', size=9)

    header_row = 8
    for idx, (label, _) in enumerate(_XLS_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=label)
        _apply_header(cell)
        ws.column_dimensions[get_column_letter(idx)].width = max(len(label) + 2, 12)

    for r_idx, row in enumerate(rows, start=1):
        excel_row = header_row + r_idx
        for c_idx, (_, key) in enumerate(_XLS_COLUMNS, start=1):
            cell = ws.cell(row=excel_row, column=c_idx, value=row.get(key, ''))
            _apply_default(cell)

    ws.freeze_panes = f'A{header_row + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

_WORD_HEADERS = ['No', 'Customer Name', 'Customer Number', 'Address', 'No Identitas',
                  'Phone', 'Reason', 'Create By', 'Create Date', 'Status', 'Delete By',
                  'Attachment']
_WORD_KEYS = ['no', 'name', 'customer_no', 'address', 'identitas', 'phone', 'reason',
              'create_user', 'create_date', 'status', 'delete_by', 'attachment']


def _build_word(rows: list[dict], as_of: date, search_desc: str, view_desc: str) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run('WATCHLIST REPORT')
    title_run.bold = True
    title_run.font.size = Pt(14)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f'As of {as_of.strftime("%d %b %Y")}')
    subtitle_run.bold = True

    meta = doc.add_table(rows=3, cols=3)
    meta.cell(0, 0).text = 'Branch'
    meta.cell(0, 1).text = ':'
    meta.cell(0, 2).text = 'Consolidate'
    meta.cell(1, 0).text = 'Search by'
    meta.cell(1, 1).text = ':'
    meta.cell(1, 2).text = search_desc
    meta.cell(2, 0).text = 'View'
    meta.cell(2, 1).text = ':'
    meta.cell(2, 2).text = view_desc

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(_WORD_HEADERS))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(_WORD_HEADERS):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8)

    for row in rows:
        cells = table.add_row().cells
        for i, key in enumerate(_WORD_KEYS):
            value = row.get(key) or ''
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()

@watchlist_report_bp.route('/watchlist-report/list')
@login_required
def list_watchlist():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        page = max(int(request.args.get('page', 1)), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = min(max(int(request.args.get('page_size', 20)), 1), 200)
    except (TypeError, ValueError):
        page_size = 20

    try:
        total, rows = _query_watchlist(
            p['as_of'], p['searchby'], p['searchbytext'],
            p['existing'], p['deleted'], page=page, page_size=page_size,
        )
        offset = (page - 1) * page_size
        data = [_serialize_row(r, offset + i) for i, r in enumerate(rows, start=1)]

        return jsonify({
            'data': data,
            'page': page,
            'page_size': page_size,
            'total': total,
            'total_pages': max((total + page_size - 1) // page_size, 1),
            'search_desc': _describe_search(p['searchby'], p['searchbytext']),
            'view_desc': _describe_view(p['existing'], p['deleted']),
            'as_of': p['as_of_raw'],
        })
    except Exception as e:
        current_app.logger.exception("Watchlist list error")
        return jsonify({'error': str(e)}), 500

@watchlist_report_bp.route('/watchlist-report/excel')
@login_required
def export_excel():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        _, rows = _query_watchlist(
            p['as_of'], p['searchby'], p['searchbytext'], p['existing'], p['deleted'],
        )
        data = [_serialize_row(r, i) for i, r in enumerate(rows, start=1)]
        search_desc = _describe_search(p['searchby'], p['searchbytext'])
        view_desc = _describe_view(p['existing'], p['deleted'])

        xlsx_bytes = _build_excel(data, p['as_of'], search_desc, view_desc)
        filename = f"WatchlistReport_{p['as_of_raw']}.xlsx"

        response = make_response(xlsx_bytes)
        response.headers['Content-Type'] = (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        current_app.logger.exception("Watchlist Excel export error")
        return jsonify({'error': str(e)}), 500

@watchlist_report_bp.route('/watchlist-report/word')
@login_required
def export_word():
    try:
        p = _parse_request_args(request.args)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400

    try:
        _, rows = _query_watchlist(
            p['as_of'], p['searchby'], p['searchbytext'], p['existing'], p['deleted'],
        )
        data = [_serialize_row(r, i) for i, r in enumerate(rows, start=1)]
        search_desc = _describe_search(p['searchby'], p['searchbytext'])
        view_desc = _describe_view(p['existing'], p['deleted'])

        docx_bytes = _build_word(data, p['as_of'], search_desc, view_desc)
        filename = f"WatchlistReport_{p['as_of_raw']}.docx"

        response = make_response(docx_bytes)
        response.headers['Content-Type'] = (
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        current_app.logger.exception("Watchlist Word export error")
        return jsonify({'error': str(e)}), 500

_PR_DETAIL_QUERY = """
    SELECT
        a.LESSEE_NM, a.ADDRESS1, a.CITY1, a.ZipCode1, a.PHONE1, a.Fax1,
        a.ADDRESS3, a.CITY3, a.PHONE3, a.ZipCode3, a.Fax3,
        a.ADDRESS4, a.CITY4, a.PHONE4, a.ZipCode4, a.Fax4,
        a.ADDRESS5, a.CITY5, a.PHONE5, a.ZipCode5, a.Fax5,
        a.Mobile_phone1, a.Mobile_phone2, a.Mobile_phone3,
        a.EMAIL1, a.EMAIL2, a.NPWP,
        a.Spouse, a.Mother, a.PLACEBIRTH, a.TGLBIRTH, a.Citizen, a.Nationality,
        a.REVIEW, a.ID_CARD, a.FAM_CARD,
        a.OTHER_PHONE1, a.OTHER_PHONE2, a.OTHER_PHONENOTES1, a.OTHER_PHONENOTES2,
        c.NAME AS CustType,
        CASE WHEN d.FLGDEL = '1' THEN '-' ELSE d.GrpName END AS GrpName,
        CASE WHEN a.GENDER = 'M' THEN 'Male' ELSE 'Female' END AS Gender,
        f.name AS Marital, g.name AS Religion, i.MLCI_Desc, mi.Descrip
    FROM CAFLESS a WITH (NOLOCK)
        LEFT JOIN MSTLESSEETYPE  c  WITH (NOLOCK) ON a.Lessee_Cat = c.CUST_TYPE AND c.LESSEE_TP = 'PR'
        LEFT JOIN CLFGROUP       d  WITH (NOLOCK) ON a.GrpCode = d.GrpCode
        LEFT JOIN MSTMARITAL     f  WITH (NOLOCK) ON a.MARITAL = f.Marital
        LEFT JOIN MSTRELIGION    g  WITH (NOLOCK) ON a.RELIGION = g.Religion
        LEFT JOIN MstIndustry_LBPP mi WITH (NOLOCK) ON a.INDUS_CODE = mi.MLCI_CODE
        LEFT JOIN UMFBITYP_LBPP  i  WITH (NOLOCK) ON a.Cust_Type = i.OJK_Code
    WHERE a.APLESS = ?
"""

_PT_DETAIL_QUERY = """
    SELECT
        a.LESSEE_NM, a.Lessee_Cat,
        a.ADDRESS2, a.CITY2, a.PHONE2, a.ZipCode2, a.Fax2,
        a.ADDRESS3, a.CITY3, a.PHONE3, a.ZipCode3, a.Fax3,
        a.ADDRESS4, a.CITY4, a.PHONE4, a.ZipCode4, a.Fax4,
        a.Address5, a.City5, a.Phone5, a.ZipCode5, a.Fax5,
        a.Mobile_phone1, a.Mobile_phone2, a.Mobile_phone3,
        a.OTHER_PHONE1, a.OTHER_PHONE2, a.OTHER_PHONENOTES1, a.OTHER_PHONENOTES2,
        a.EMAIL1, a.EMAIL2, a.NPWP,
        a.TGLBIRTH, a.REVIEW, a.Contact, a.Occupation AS Position, a.LOB,
        b.name AS OfficeStatus, c.NAME AS CustType,
        CASE WHEN d.FLGDEL = '1' THEN '-' ELSE d.GrpName END AS GrpName,
        i.MLCI_Desc, mi.Descrip
    FROM CAFLESS a WITH (NOLOCK)
        LEFT JOIN MSTADDRESSSTATUS b  WITH (NOLOCK) ON a.Address_Status = b.AddressStatus
        LEFT JOIN MSTLESSEETYPE    c  WITH (NOLOCK) ON a.Lessee_Cat = c.CUST_TYPE AND c.LESSEE_TP = 'PT'
        LEFT JOIN MstIndustry_LBPP mi WITH (NOLOCK) ON a.INDUS_CODE = mi.MLCI_CODE
        LEFT JOIN CLFGROUP         d  WITH (NOLOCK) ON a.GrpCode = d.GrpCode
        LEFT JOIN UMFBITYP_LBPP    i  WITH (NOLOCK) ON a.Cust_Type = i.OJK_Code
    WHERE a.APLESS = ?
"""

_REFERENCES_QUERY = """
    SELECT ID, REF_NAME, REF_PHONE, REMARK
    FROM CAFREF WITH (NOLOCK)
    WHERE APLESS = ?
    ORDER BY ID
"""

_DOCUMENT_QUERY = """
    SELECT TOP 1 file_document
    FROM CUSTOMER_DOCUMENT WITH (NOLOCK)
    WHERE APLESS = ? AND organization_type = ? AND doc_type = 'non'
    ORDER BY Upload_Dt DESC
"""

_PR_CITIZEN_DOC_LABELS = {
    'WNI': [
        'KTP a/n Pelanggan',
        'KTP a/n Suami/Istri Pelanggan',
        'Kartu Keluarga a/n Pelanggan',
        'Kartu Keluarga a/n Suami/Istri Pelanggan',
    ],
    'WNA': [
        'KITAS/KITAP a/n Pelanggan',
        'KITAS/KITAP a/n Suami/Istri Pelanggan',
        'SK. Susunan Keluarga Pendatang (SKSP) WNA',
        'Surat Referensi Kerja dari Perusahaan WNA Bekerja',
    ],
}

_PT_ORG_TYPE = {'CO': 'KOPERASI', 'CV': 'CV', 'FD': 'YAYASAN', 'PT': 'PT'}
_PT_ORG_DOC_LABELS = {
    'CO': [
        'Susunan Pengurus & Pengawas Koperasi',
        'KTP a/n Pengurus Penanda Tangan Kontrak',
        'KTP a/n Pengawas Penanda Tangan Kontrak',
        'SIUP',
        'NPWP',
    ],
    'CV': [
        'Akta Pendirian CV + Perubahan',
        'KTP a/n Persero Aktif (Direktur)',
        'KTP a/n Persero Komanditer',
        'SIUP',
        'NPWP',
    ],
    'FD': [
        'KTP a/n Pengurus',
        'KTP a/n Pengawas',
        'NPWP',
    ],
    'PT': [
        'SIUP',
        'NPWP',
        'Surat Kuasa Asli dari Direktur',
    ],
}

def _format_npwp(npwp: Optional[str]) -> str:
    if not npwp:
        return '-'
    return (f"{npwp[0:2]}.{npwp[2:5]}.{npwp[5:8]}.{npwp[8:9]}"
            f"-{npwp[9:12]}.{npwp[12:16]}")

def _address_block(address, city, zipcode, phone, fax):
    if not address:
        return None
    return {
        'address': address,
        'city': city or '',
        'zipcode': zipcode or '',
        'phone': phone or '',
        'fax': fax or None,
    }

def _join_nonempty(*vals) -> str:
    parts = [v for v in vals if v]
    return ', '.join(parts) if parts else '-'

def _fmt_detail_date(val) -> str:
    if isinstance(val, (date, datetime)):
        return val.strftime('%d-%m-%Y')
    return '-'


def _fetch_references(apless: str) -> list[dict]:
    rows = _exec_rows(_REFERENCES_QUERY, [apless])
    return [
        {
            'no': i,
            'name': r.get('REF_NAME') or '',
            'phone': r.get('REF_PHONE') or '',
            'remark': r.get('REMARK') or '',
        }
        for i, r in enumerate(rows, start=1)
    ]

def _fetch_document_filenames(apless: str, organization_type: str) -> list[str]:
    rows = _exec_rows(_DOCUMENT_QUERY, [apless, organization_type])
    if not rows or not rows[0].get('file_document'):
        return []
    return [f for f in rows[0]['file_document'].split(';')]


def _build_doc_rows(labels: list[str], filenames: list[str], start_index: int = 0) -> list[dict]:
    return [
        {'label': label, 'filename': (filenames[start_index + i]
                                       if start_index + i < len(filenames) else '') or None}
        for i, label in enumerate(labels)
    ]

def _fetch_lessee_type(apless: str) -> Optional[str]:
    rows = _exec_rows("SELECT LESSEE_TP FROM cafless WITH (NOLOCK) WHERE APLESS = ?", [apless])
    if not rows:
        return None
    return (rows[0].get('LESSEE_TP') or '').strip()

def _fetch_pr_detail(apless: str) -> Optional[dict]:
    rows = _exec_rows(_PR_DETAIL_QUERY, [apless])
    if not rows:
        return None
    r = rows[0]

    citizen = (r.get('Citizen') or '').strip()
    citizen_labels = _PR_CITIZEN_DOC_LABELS.get(citizen, [])
    filenames = _fetch_document_filenames(apless, 'PR')

    documents = {
        'citizen_type': citizen or None,
        'citizen_docs': _build_doc_rows(citizen_labels, filenames, 0),
        'additional_docs': _build_doc_rows(['Slip Gaji', 'NPWP'], filenames, 4),
        'bpkb_docs': _build_doc_rows(['KTP an. BPKB', 'Kartu Keluarga an. BPKB'], filenames, 6),
        'other_doc': _build_doc_rows(['Other'], filenames, 8)[0],
    }

    return {
        'lessee_type': 'PR',
        'name': r.get('LESSEE_NM') or '-',
        'customer_type': r.get('CustType') or '-',
        'group_name': r.get('GrpName') or '-',
        'id_address': _address_block(r.get('ADDRESS1'), r.get('CITY1'), r.get('ZipCode1'),
                                      r.get('PHONE1'), r.get('Fax1')),
        'correspondence': _address_block(r.get('ADDRESS3'), r.get('CITY3'), r.get('ZipCode3'),
                                          r.get('PHONE3'), r.get('Fax3')),
        'additional_1': _address_block(r.get('ADDRESS4'), r.get('CITY4'), r.get('ZipCode4'),
                                        r.get('PHONE4'), r.get('Fax4')),
        'additional_2': _address_block(r.get('ADDRESS5'), r.get('CITY5'), r.get('ZipCode5'),
                                        r.get('PHONE5'), r.get('Fax5')),
        'mobile': _join_nonempty(r.get('Mobile_phone1'), r.get('Mobile_phone2'), r.get('Mobile_phone3')),
        'email': _join_nonempty(r.get('EMAIL1'), r.get('EMAIL2')),
        'npwp': _format_npwp(r.get('NPWP')),
        'spouse': r.get('Spouse') or '-',
        'mothers_maiden_name': r.get('Mother') or '-',
        'place_of_birth': r.get('PLACEBIRTH') or '-',
        'date_of_birth': _fmt_detail_date(r.get('TGLBIRTH')),
        'gender': r.get('Gender') or '-',
        'marital_status': r.get('Marital') or '-',
        'religion': r.get('Religion') or '-',
        'citizenship': citizen or '-',
        'nationality': r.get('Nationality') or '-',
        'id_card_no': r.get('ID_CARD') or '-',
        'family_card_no': r.get('FAM_CARD') or '-',
        'industrial_code': r.get('Descrip') or '-',
        'bi_customer_type': r.get('MLCI_Desc') or '-',
        'other_phone_1': r.get('OTHER_PHONE1') or '',
        'other_phone_1_notes': r.get('OTHER_PHONENOTES1') or '',
        'other_phone_2': r.get('OTHER_PHONE2') or '',
        'other_phone_2_notes': r.get('OTHER_PHONENOTES2') or '',
        'review_date': _fmt_detail_date(r.get('REVIEW')),
        'references': _fetch_references(apless),
        'documents': documents,
    }

def _fetch_pt_detail(apless: str) -> Optional[dict]:
    rows = _exec_rows(_PT_DETAIL_QUERY, [apless])
    if not rows:
        return None
    r = rows[0]

    lessee_cat = (r.get('Lessee_Cat') or '').strip()
    org_labels = _PT_ORG_DOC_LABELS.get(lessee_cat, [])
    org_title = _PT_ORG_TYPE.get(lessee_cat)
    filenames = _fetch_document_filenames(apless, org_title or '')

    bpkb_start = len(org_labels)
    documents = {
        'org_type': org_title,
        'org_docs': _build_doc_rows(org_labels, filenames, 0),
        'bpkb_docs': _build_doc_rows(['KTP an. BPKB', 'Kartu Keluarga an. BPKB'], filenames, bpkb_start),
        'other_doc': _build_doc_rows(['Other'], filenames, bpkb_start + 2)[0],
    }

    return {
        'lessee_type': 'PT',
        'name': r.get('LESSEE_NM') or '-',
        'customer_type': r.get('CustType') or '-',
        'group_name': r.get('GrpName') or '-',
        'id_address': _address_block(r.get('ADDRESS2'), r.get('CITY2'), r.get('ZipCode2'),
                                      r.get('PHONE2'), r.get('Fax2')),
        'correspondence': _address_block(r.get('ADDRESS3'), r.get('CITY3'), r.get('ZipCode3'),
                                          r.get('PHONE3'), r.get('Fax3')),
        'additional_1': _address_block(r.get('ADDRESS4'), r.get('CITY4'), r.get('ZipCode4'),
                                        r.get('PHONE4'), r.get('Fax4')),
        'additional_2': _address_block(r.get('Address5'), r.get('City5'), r.get('ZipCode5'),
                                        r.get('Phone5'), r.get('Fax5')),
        'mobile': _join_nonempty(r.get('Mobile_phone1'), r.get('Mobile_phone2'), r.get('Mobile_phone3')),
        'email': _join_nonempty(r.get('EMAIL1'), r.get('EMAIL2')),
        'npwp': _format_npwp(r.get('NPWP')),
        'contact': r.get('Contact') or '-',
        'position': r.get('Position') or '-',
        'line_of_business': r.get('LOB') or '-',
        'establishment_date': _fmt_detail_date(r.get('TGLBIRTH')),
        'industrial_code': r.get('Descrip') or '-',
        'bi_customer_type': r.get('MLCI_Desc') or '-',
        'office_status': r.get('OfficeStatus') or '-',
        'other_phone_1': r.get('OTHER_PHONE1') or '',
        'other_phone_1_notes': r.get('OTHER_PHONENOTES1') or '',
        'other_phone_2': r.get('OTHER_PHONE2') or '',
        'other_phone_2_notes': r.get('OTHER_PHONENOTES2') or '',
        'review_date': _fmt_detail_date(r.get('REVIEW')),
        'references': _fetch_references(apless),
        'documents': documents,
    }

@watchlist_report_bp.route('/watchlist-report/detail/<apless>')
@login_required
def customer_detail(apless):
    try:
        lessee_type = _fetch_lessee_type(apless)
        if lessee_type is None:
            return jsonify({'error': 'Customer not found'}), 404

        data = _fetch_pr_detail(apless) if lessee_type == 'PR' else _fetch_pt_detail(apless)
        if data is None:
            return jsonify({'error': 'Customer not found'}), 404

        return jsonify(data)
    except Exception as e:
        current_app.logger.exception("Customer detail review error")
        return jsonify({'error': str(e)}), 500