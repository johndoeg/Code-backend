from flask import Blueprint, request, send_file, current_app
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text
import io
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

form10_bp = Blueprint('form10', __name__, url_prefix='/Form10')

_THIN         = Side(style='thin')
_BORDER       = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT_DATA    = Font(name='Tahoma', size=7)
_FILL_HEADER  = PatternFill('solid', fgColor='1F4E79')
_ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_ALIGN_LEFT   = Alignment(horizontal='left',   vertical='center')
_ALIGN_RIGHT  = Alignment(horizontal='right',  vertical='center')

COLUMNS = [
    ('A',  'No. Debitur',                              'No Debitur',                       _ALIGN_LEFT),
    ('B',  'Nama Debitur',                             'Nama Debitur',                     _ALIGN_LEFT),
    ('C',  'Golongan Debitur',                         'Golongan Debitur',                 _ALIGN_RIGHT),
    ('D',  'Sektor Ekonomi',                           'Sektor Ekonomi',                   _ALIGN_RIGHT),
    ('E',  'Lokasi Proyek',                            'Lokasi Proyek',                    _ALIGN_LEFT),
    ('F',  'No. Kontrak',                              'No kontrak',                       _ALIGN_LEFT),
    ('G',  'Jenis Pembiayaan',                         'Jenis Pembiayaan',                 _ALIGN_LEFT),
    ('H',  'Skema Pembiayaan',                         'Skema Pembiayaan',                 _ALIGN_LEFT),
    ('I',  'Tujuan Pembiayaan',                        'Tujuan Pembiayaan',                _ALIGN_LEFT),
    ('J',  'Tgl Mulai Pembiayaan',                     'Tgl Mulai Pembiayaan',             _ALIGN_LEFT),
    ('K',  'Tgl Jth Tempo Pembiayaan',                 'Tgl Jth Tempo Pembiayaan',         _ALIGN_LEFT),
    ('L',  'Nilai Awal Pembiayaan',                    'Nilai Awal Pembiayaan',            _ALIGN_LEFT),
    ('M',  'Uang Muka',                                'Uang Muka',                        _ALIGN_LEFT),
    ('N',  'Biaya Komisi',                             'Biaya Komisi',                     _ALIGN_LEFT),
    ('O',  'Jenis Suku Bunga',                         'Jenis Suku Bunga',                 _ALIGN_LEFT),
    ('P',  'Tingkat Bunga',                            'Tingkat Bunga',                    _ALIGN_LEFT),
    ('Q',  'Tingkat Bunga Hasil Bagi',                 'Tingkat Bunga Hasil Bagi',         _ALIGN_LEFT),
    ('R',  'Biaya Adm',                                'Biaya Adm',                        _ALIGN_LEFT),
    ('S',  'Biaya Provisi',                            'Biaya Provisi',                    _ALIGN_LEFT),
    ('T',  'Kategori Piutang',                         'Kategori Piutang',                 _ALIGN_LEFT),
    ('U',  'Tgl Jth Tempo Terakhir',                   'tgl jth tempo terakhir',           _ALIGN_LEFT),
    ('V',  'Tgl Bayar Terakhir',                       'tgl byr terakhir',                 _ALIGN_LEFT),
    ('W',  'Periode Bayar Terakhir',                   'periode byr terakhir',             _ALIGN_LEFT),
    ('X',  'Nilai Bayar Terakhir',                     'nilai byr terakhir',               _ALIGN_LEFT),
    ('Y',  'Jenis Barang & Jasa yg Dibiayai',          'Jenis Barang & Jasa yg Dibiayai',  _ALIGN_LEFT),
    ('Z',  'Nilai Barang & Jasa yg Dibiayai',          'Nilai Barang & Jasa yg Dibiayai',  _ALIGN_LEFT),
    ('AA', 'Metode CKPN',                              'metode ckpn',                      _ALIGN_LEFT),
    ('AB', 'Perusahaan Asuransi',                      'Perusahaan Asuransi',              _ALIGN_LEFT),
    ('AC', 'Jangka Waktu Asuransi',                    'Jangka Waktu Asuransi',            _ALIGN_LEFT),
    ('AD', 'Premi Debitur',                            'Premi Debitur',                    _ALIGN_LEFT),
    ('AE', 'Identitas Agunan',                         'Identitas Agunan',                 _ALIGN_LEFT),
    ('AF', 'Jenis Agunan',                             'Jenis Agunan',                     _ALIGN_LEFT),
    ('AG', 'Nilai Agunan',                             'Nilai Agunan',                     _ALIGN_LEFT),
    ('AH', 'Jenis Pengikatan',                         'Jenis Pengikatan',                 _ALIGN_LEFT),
    ('AI', 'No. Sertifikasi Kepemilikan',              'No Sertifikat Kepemilikan',        _ALIGN_LEFT),
    ('AJ', 'No. Sertifikasi Pengikatan',               'No Sertifikat Pengikatan',         _ALIGN_LEFT),
    ('AK', 'Tgl Sertifikasi Pengikatan',               'Tgl Sertifikat Pengikatan',        _ALIGN_LEFT),
    ('AL', 'Posisi Penyimpanan Sertifikat',            'Posisi Penyimpanan Sertifikat',    _ALIGN_LEFT),
    ('AM', 'Kategori Usaha Debitur',                   'Kategori Usaha Debitur',           _ALIGN_LEFT),
    ('AN', 'Business Size',                            'Business Size',                    _ALIGN_LEFT),
    ('AO', 'Porsi Perusahaan Pada Pembiayaan Bersama', 'Porsi Perusahaan',                 _ALIGN_LEFT),
]

def _build_workbook(rows: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Form 10'

    for col_letter, header, _, _ in COLUMNS:
        cell = ws[f'{col_letter}1']
        cell.value = header
        cell.font = Font(name='Tahoma', size=7, bold=True, color='FFFFFF')
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _BORDER
        ws.column_dimensions[col_letter].width = max(len(header) + 2, 8)

    ws.row_dimensions[1].height = 30

    for i, row in enumerate(rows, start=2):
        for col_letter, _, db_key, align in COLUMNS:
            cell = ws[f'{col_letter}{i}']
            cell.value = row.get(db_key)
            cell.font = _FONT_DATA
            cell.alignment = align
            cell.border = _BORDER

    ws.freeze_panes = 'A2'
    return wb

@form10_bp.route('/excel')
@login_required
def get_form10_excel():
    try:
        date_from = request.args.get('date_from')
        if not date_from:
            return {'error': 'date_from is required'}, 400

        result = db.session.execute(
            text("EXEC USP_OJK_FORM10 :date_from"),
            {'date_from': date_from},
        )
        rows = [dict(zip(result.keys(), r)) for r in result.fetchall()]

        output = io.BytesIO()
        _build_workbook(rows).save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'Form_10_{date_from}.xlsx',
        )

    except Exception:
        current_app.logger.exception('Form10 excel export error')
        return {'error': 'Internal server error'}, 500