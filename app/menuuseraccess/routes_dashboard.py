from datetime import date as date_cls, datetime
from io import BytesIO

from flask import Blueprint, jsonify, request, current_app, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app import db
from app.auth.decorators import login_required

dashboard_bp = Blueprint('dashboard', __name__, url_prefix='/SystemAdmin')

API_LIST = ['ESPAY', 'DELIMA', 'RAPINDO', 'INSURANCE', 'SMS']

def _parse_date(raw: str):
    raw = (raw or '').strip()
    if not raw:
        return date_cls.today()
    return datetime.strptime(raw, "%Y-%m-%d").date()

def _fetch_proc(sql: str, params: tuple) -> list:
    
    raw = db.session.connection().connection
    cur = raw.cursor()
    try:
        cur.execute(sql, params)
        while cur.description is None:
            if not cur.nextset():
                return []
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, row)) for row in cur.fetchall()]
    finally:
        cur.close()

def _val(row: dict, key: str):
    if key in row:
        return row[key]
    lowered = {k.lower(): v for k, v in row.items()}
    return lowered.get(key.lower())

def _cell(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date_cls):
        return value.strftime("%Y-%m-%d")
    return "" if value is None else value

@dashboard_bp.route('/dashboard')
@login_required
def get_dashboard():
    try:
        run_date = _parse_date(request.args.get('date', ''))
    except ValueError:
        return jsonify({"data": [], "message": "Date must be in YYYY-MM-DD format"}), 400

    try:
        rows = []
        for api_name in API_LIST:
            for row in _fetch_proc("EXEC USP_DASHBOARD_API ?, ?", (run_date, api_name)):
                rows.append({
                    "api":         api_name,
                    "requesttype": _val(row, "RequestType") or "",
                    "status_code": _val(row, "Status_Code") or "",
                    "jlh":         _val(row, "jlh") or 0,
                })

        return jsonify({"data": rows, "date": run_date.strftime("%Y-%m-%d")})

    except Exception as e:
        current_app.logger.exception("Dashboard summary error")
        return jsonify({"data": [], "error": str(e)}), 500

@dashboard_bp.route('/dashboard/detail')
@login_required
def get_dashboard_detail():
    api_name     = (request.args.get('api') or '').strip()
    request_type = (request.args.get('request_type') or '').strip()

    if not api_name:
        return jsonify({"data": [], "message": "API is required"}), 400

    try:
        run_date = _parse_date(request.args.get('date', ''))
    except ValueError:
        return jsonify({"data": [], "message": "Date must be in YYYY-MM-DD format"}), 400

    try:
        rows = [
            {
                "lease_no":  _val(row, "Lease_No")  or "",
                "lessee_nm": _val(row, "Lessee_Nm") or "",
                "param1":    _val(row, "Param1")    or "",
            }
            for row in _fetch_proc(
                "EXEC USP_DASHBOARD_API_DETAIL ?, ?, ?",
                (run_date, api_name, request_type)
            )
        ]
        return jsonify({"data": rows})

    except Exception as e:
        current_app.logger.exception("Dashboard detail error")
        return jsonify({"data": [], "error": str(e)}), 500

EXPORT_HEADERS = [
    ("A", "No.",           None),
    ("B", "Contract No.",  "LEASE_NO"),
    ("C", "Customer Name", "LESSEE_NM"),
    ("D", "Disb. Date",    "execution"),
    ("E", "Cover Date",    "DelimaCoverDt"),
    ("F", "Batch ID",      "BatchID"),
    ("G", "Success Date",  "DelimaSuccessDt"),
    ("H", "Name",          "Ins_co"),
    ("I", "Cover Date",    "InsCoverDt"),
    ("J", "Status",        "InsStatus"),
    ("K", "Claim Date",    "ClaimDt"),
    ("L", "Claim Status",  "ClaimStatus"),
    ("M", "CLFPAY",        "CLFPAY"),
    ("N", "BGIRO",         "BGIRO"),
    ("O", "Generate VA",   "VA"),
    ("P", "Invoice Type",  "TransInvoice"),
]

GROUP_HEADERS = [
    ("E2:G2", "DELIMA"),
    ("H2:J2", "INSURANCE"),
    ("K2:L2", "RAPINDO"),
    ("M2:N2", "AMORTIZATION"),
    ("O2:P2", "ESPAY"),
]

@dashboard_bp.route('/dashboard/export')
@login_required
def export_dashboard():
    try:
        run_date = _parse_date(request.args.get('date', ''))
    except ValueError:
        return jsonify({"message": "Date must be in YYYY-MM-DD format"}), 400

    try:
        rows = _fetch_proc("EXEC USP_DASHBOARD_DETAIL ?", (run_date,))

        wb = Workbook()
        ws = wb.active
        ws.title = "Dashboard"

        thin = Side(style="thin", color="FF000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        title_font = Font(name="Tahoma", size=12, bold=True)
        body_font = Font(name="Tahoma", size=7)
        center = Alignment(horizontal="center", vertical="center")

        for span, label in GROUP_HEADERS:
            ws.merge_cells(span)
            cell = ws[span.split(':')[0]]
            cell.value = label
            cell.font = title_font
            cell.alignment = center
        for col, _, _ in EXPORT_HEADERS:
            ws[f"{col}2"].border = border

        for col, label, _ in EXPORT_HEADERS:
            cell = ws[f"{col}3"]
            cell.value = label
            cell.font = title_font
            cell.alignment = center
            cell.border = border

        row_idx = 4
        for no, row in enumerate(rows, start=1):
            for col, _, key in EXPORT_HEADERS:
                cell = ws[f"{col}{row_idx}"]
                cell.value = no if key is None else _cell(_val(row, key))
                cell.font = body_font
                cell.border = border
            row_idx += 1

        for idx, (col, label, _) in enumerate(EXPORT_HEADERS, start=1):
            longest = len(label)
            for r in range(4, row_idx):
                longest = max(longest, len(str(ws.cell(row=r, column=idx).value or "")))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 6), 40)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        return send_file(
            stream,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Dashboard_{run_date.strftime('%Y-%m-%d')}.xlsx",
        )

    except Exception as e:
        current_app.logger.exception("Dashboard export error")
        return jsonify({"message": "Failed", "error": str(e)}), 500