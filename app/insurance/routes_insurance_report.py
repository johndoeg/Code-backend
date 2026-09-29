from flask import Blueprint, request, send_file, current_app, jsonify
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text

import io
import re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import date

insurance_report_bp = Blueprint('insurance_report', __name__, url_prefix='/InsuranceReport')

_THIN         = Side(style='thin')
_BORDER       = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FILL_HEADER  = PatternFill('solid', fgColor='1F4E79')
_FONT_HEADER  = Font(name='Tahoma', size=7, bold=True, color='FFFFFF')
_FONT_TITLE   = Font(name='Tahoma', size=12, bold=True)
_FONT_DATA    = Font(name='Tahoma', size=7)
_ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_ALIGN_LEFT   = Alignment(horizontal='left',   vertical='center')

COLUMNS = [
    ('A', 'No.',               None),
    ('B', 'Contract No.',      'Lease_No'),
    ('C', 'Customer Name',     'LESSEE_NM'),
    ('D', 'Insurance Company', 'INS_CO'),
    ('E', 'Policy No.',        'PolicyNo'),
    ('F', 'Status',            'ResponseStatus'),
    ('G', 'Message',           'ResponseMessage'),
    ('H', 'Success Date',      'success_dt'),
]

FILTER_LABELS = {'1': 'Contract No. ', '2': 'Disbursement Date '}

def _clean_message(raw: str) -> str:
    if not raw:
        return '-'
    cleaned = re.sub(r'[\r\n\t]+', ' ', raw).strip()
    return cleaned or '-'

def _run_sp(ins_cd, search_by, search_val, search_val1, search_val2):
    result = db.session.execute(
        text("EXEC USP_INSURANCE_REPORT :ins_cd, :search_by, :search_val, :search_val1, :search_val2"),
        {
            "ins_cd":      ins_cd,
            "search_by":   search_by,
            "search_val":  search_val,
            "search_val1": search_val1,
            "search_val2": search_val2,
        },
    )
    return result

@insurance_report_bp.route('/companies')
@login_required
def get_companies():
    try:
        result = db.session.execute(text("""
            SELECT INS_CD, INS_CO
            FROM UMFINSCD
            WHERE INS_CD IN ('01', '02')
        """))
        return jsonify([{"INS_CD": r[0], "INS_CO": r[1]} for r in result])
    except Exception:
        current_app.logger.exception("companies error")
        return jsonify([])

@insurance_report_bp.route('/list')
@login_required
def get_list():
    ins_cd      = request.args.get('ins_cd',      '')
    search_by   = request.args.get('search_by',   '')
    search_val  = request.args.get('search_val',  '')
    search_val1 = request.args.get('search_val1', '')
    search_val2 = request.args.get('search_val2', '')
    page        = request.args.get('page',  1,  type=int)
    limit       = request.args.get('limit', 25, type=int)

    try:
        result = _run_sp(ins_cd, search_by, search_val, search_val1, search_val2)
        all_rows = [
            {
                "lease_no":   row.Lease_No         or "",
                "lessee_nm":  row.LESSEE_NM        or "",
                "ins_co":     row.INS_CO            or "",
                "policy_no":  row.PolicyNo          or "",
                "status":     row.ResponseStatus    or "",
                "message":    _clean_message(row.ResponseMessage),
                "success_dt": str(row.success_dt)   if row.success_dt else "",
            }
            for row in result
        ]

        total     = len(all_rows)
        start     = (page - 1) * limit
        paginated = all_rows[start:start + limit]

        return jsonify({"data": paginated, "total": total})

    except Exception:
        current_app.logger.exception("insurance report list error")
        return jsonify({"error": "Failed to load"}), 500

@insurance_report_bp.route('/excel')
@login_required
def get_excel():
    ins_cd      = request.args.get('ins_cd',      '')
    search_by   = request.args.get('search_by',   '')
    search_val  = request.args.get('search_val',  '')
    search_val1 = request.args.get('search_val1', '')
    search_val2 = request.args.get('search_val2', '')

    try:
        result = _run_sp(ins_cd, search_by, search_val, search_val1, search_val2)
        rows   = result.fetchall()

        wb = _build_workbook(rows, search_by, search_val, search_val1, search_val2)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'InsuranceReport_{date.today().isoformat()}.xlsx',
        )

    except Exception:
        current_app.logger.exception("insurance report excel error")
        return jsonify({"error": "Export failed"}), 500

def _build_workbook(rows, search_by, search_val, search_val1, search_val2):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'InsuranceReport'

    today_str = date.today().strftime('%d %b %Y')

    ws.merge_cells('A1:G1')
    ws['H1'] = f'Date : {today_str}'
    ws['H1'].font = _FONT_DATA
    ws['H1'].alignment = _ALIGN_LEFT

    ws.merge_cells('A2:G2')
    ws['A2'].font      = _FONT_DATA
    ws['A2'].alignment = _ALIGN_LEFT
    ws['H2'].font      = _FONT_DATA
    ws['H2'].alignment = _ALIGN_LEFT

    ws.merge_cells('A3:H3')
    ws['A3']           = 'INSURANCE REPORT'
    ws['A3'].font      = _FONT_TITLE
    ws['A3'].alignment = _ALIGN_CENTER

    ws.merge_cells('A4:H4')
    filter_label = FILTER_LABELS.get(search_by, '')
    if search_by == '1':
        ws['A4'] = filter_label + search_val
    elif search_by == '2':
        ws['A4'] = f"{filter_label}{search_val1} - {search_val2}"
    ws['A4'].alignment = _ALIGN_CENTER

    for col_letter, header, _ in COLUMNS:
        cell            = ws[f'{col_letter}5']
        cell.value      = f' {header} '
        cell.font       = _FONT_HEADER
        cell.fill       = _FILL_HEADER
        cell.alignment  = _ALIGN_CENTER
        cell.border     = _BORDER
        ws.column_dimensions[col_letter].width = max(len(header) + 4, 8)

    ws.row_dimensions[5].height = 25
    ws.column_dimensions['A'].width = 6

    for i, row in enumerate(rows, start=1):
        r = i + 5
        data = {
            'Lease_No':        row.Lease_No        or '',
            'LESSEE_NM':       row.LESSEE_NM       or '',
            'INS_CO':          row.INS_CO           or '',
            'PolicyNo':        row.PolicyNo         or '',
            'ResponseStatus':  row.ResponseStatus   or '',
            'ResponseMessage': _clean_message(row.ResponseMessage),
            'success_dt':      str(row.success_dt)  if row.success_dt else '',
        }
        for col_letter, _, db_key in COLUMNS:
            cell           = ws[f'{col_letter}{r}']
            cell.value     = i if db_key is None else data.get(db_key, '')
            cell.font      = _FONT_DATA
            cell.alignment = _ALIGN_LEFT
            cell.border    = _BORDER

    ws.freeze_panes = 'A6'
    return wb