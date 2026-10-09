import calendar
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Optional

import httpx
from flask import Blueprint, jsonify, request, current_app, session
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text

insurance_bp = Blueprint('insurance', __name__, url_prefix='/Insurance')

INSURERS = {
    'raksa': ('USP_INVOICE_RAKSA', '01'),
    'simas': ('USP_INVOICE_SIMAS', '02'),
}
RAKSA_INS_CD = INSURERS['raksa'][1]
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
_MONTHS = {m: i for i, m in enumerate(
    ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'], 1)}

def _flag(value) -> str:
    return "1" if str(value).strip().lower() in ("1", "true") else "0"

def _page_args():
    page = max(request.args.get('page', 1, type=int), 1)
    limit = min(max(request.args.get('limit', 25, type=int), 1), 200)
    return page, limit

def _invoice_summary(insurer: str):
    proc, _ = INSURERS[insurer]
    page, limit = _page_args()
    try:
        result = db.session.execute(
            text(f"EXEC {proc} :mode, :param2, :param3"),
            {"mode": "Summary", "param2": "", "param3": ""},
        )

        all_rows = []
        for row in result:
            item = {
                "asof": str(row.tanggal) if row.tanggal else "",
                "disb_month": row.disb_month or "",
                "tot_kontrak_inv": str(row.tot_kontrak_inv) if row.tot_kontrak_inv is not None else "0",
                "tot_biaya_inv": str(row.tot_nett_inv) if row.tot_nett_inv is not None else "0",
                "tot_kontrak_sys": str(row.tot_kontrak_sys) if row.tot_kontrak_sys is not None else "0",
                "tot_biaya_sys": str(row.tot_nett_sys) if row.tot_nett_sys is not None else "0",
                "confirmby": row.ConfirmBy or "",
                "confirmdate": str(row.ConfirmDate) if row.ConfirmDate else "",
            }
            if insurer == 'simas':
                item["payment_dt"] = str(row.Payment_Dt) if row.Payment_Dt else ""
                item["batchno"] = str(row.BatchNo) if row.BatchNo is not None else ""
            all_rows.append(item)

        total = len(all_rows)
        start = (page - 1) * limit
        return jsonify({"data": all_rows[start:start + limit], "total": total})

    except Exception:
        current_app.logger.exception("%s invoice summary error", insurer)
        return jsonify({"data": [], "total": 0}), 500

def _invoice_detail(insurer: str):
    proc, _ = INSURERS[insurer]
    invoice_no = request.args.get('id')
    filter_param = request.args.get('filter') or '1'
    if not invoice_no:
        return jsonify([])

    try:
        result = db.session.execute(
            text(f"EXEC {proc} :mode, :param2, :param3"),
            {"mode": "Detail", "param2": invoice_no, "param3": filter_param},
        )

        details = []
        for row in result:
            totalnett = row.TotalNett
            gross = row.TotalNettSystem

            totalnett_str = "-" if not totalnett or str(totalnett).strip() == "" else str(totalnett)
            gross_str = "-" if not gross or str(gross).strip() == "" else str(gross)
            batchdt = "-" if not row.CoverDt or str(row.CoverDt).strip() == "" else str(row.CoverDt)

            remark = ""
            try:
                nett_val = float(totalnett) if totalnett else 0
                sys_val = float(gross) if gross else 0
                if sys_val - nett_val > 10000:
                    remark = "Net Premi<br>"
            except (ValueError, TypeError):
                pass

            if row.Invoice != row.System:
                remark += "No. kontrak"

            details.append({
                "lease_no": row.lease_no or "",
                "lessee_nm": row.LESSEE_NM or "",
                "execution": str(row.EXECUTION) if row.EXECUTION else "",
                "batchdt": batchdt,
                "totalnett": totalnett_str,
                "gross": gross_str,
                "invoice": row.Invoice or "",
                "system": row.System or "",
                "remark": remark,
                "branch": row.BRANCH_NAME or "",
                "isconfirm": _flag(row.isConfirm),
                "confirmby": row.ConfirmBy or "",
                "confirmdate": str(row.ConfirmDate) if row.ConfirmDate else "",
                "hasil": row.Hasil or "",
            })

        return jsonify(details)

    except Exception:
        current_app.logger.exception("%s invoice detail error", insurer)
        return jsonify([])

def _confirm_invoice(insurer: str):
    ins_cd = INSURERS[insurer][1]
    body = request.get_json(silent=True) or {}
    asof = str(body.get('asof') or '').strip()
    emp_id = session.get('emp_id')

    seen, leases = set(), []
    for v in body.get('leases') or []:
        v = str(v).strip()
        if v and v != '0' and v not in seen:
            seen.add(v)
            leases.append(v)

    if not asof:
        return jsonify({"success": False, "message": "asof is required"}), 400
    if not leases:
        return jsonify({"success": False, "message": "Please, select at least one contract!"}), 400

    try:
        done = db.session.execute(
            text("SELECT 1 FROM INVOICE_INS_CONFIRMATION WHERE INS_CD = :ins_cd AND Asof = :asof"),
            {"ins_cd": ins_cd, "asof": asof},
        ).first()
        if done:
            return jsonify({"success": False, "message": "Already confirmed"}), 409

        db.session.execute(
            text("INSERT INTO INVOICE_INS_CONFIRMDETAIL (INS_CD, Asof, Lease_No) "
                 "VALUES (:ins_cd, :asof, :lease_no)"),
            [{"ins_cd": ins_cd, "asof": asof, "lease_no": l} for l in leases],
        )
        db.session.execute(
            text("INSERT INTO INVOICE_INS_CONFIRMATION (INS_CD, Asof, ConfirmBy, ConfirmDate) "
                 "VALUES (:ins_cd, :asof, :emp_id, GETDATE())"),
            {"ins_cd": ins_cd, "asof": asof, "emp_id": emp_id},
        )
        db.session.commit()
        return jsonify({"success": True, "message": "Successful"})

    except Exception:
        db.session.rollback()
        current_app.logger.exception("%s confirm error", insurer)
        return jsonify({"success": False, "message": "Failed"}), 500

@insurance_bp.route('/filter-options')
@login_required
def get_filter_options():
    try:
        result = db.session.execute(text("""
            SELECT [VALUE], DESC_VALUE
            FROM MSTCOMBO WITH (NOLOCK)
            WHERE CAT_VALUE = 'INSURANCE_FILTER_BY'
            ORDER BY [VALUE]
        """))
        return jsonify([{"value": str(r[0]), "label": r[1] or ""} for r in result])
    except Exception:
        current_app.logger.exception("insurance filter options error")
        return jsonify([])

@insurance_bp.route('/simas/invoice')
@login_required
def get_simas_invoice_summary():
    return _invoice_summary('simas')

@insurance_bp.route('/simas/detail')
@login_required
def get_simas_invoice_detail():
    return _invoice_detail('simas')

@insurance_bp.route('/simas/confirm', methods=['POST'])
@login_required
def confirm_simas_invoice():
    return _confirm_invoice('simas')

def _parse_pmt_date(raw) -> Optional[str]:
    raw = (raw or '').strip()
    for fmt in ('%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(raw, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return None

@insurance_bp.route('/simas/payment', methods=['POST'])
@login_required
def set_simas_payment():
    body = request.get_json(silent=True) or {}
    batchno = str(body.get('batchno') or '').strip()
    pmt_dt = _parse_pmt_date(body.get('pmt_dt'))

    if not batchno:
        return jsonify({"success": False, "message": "batchno is required"}), 400
    if not pmt_dt:
        return jsonify({"success": False, "message": "Please fill in Payment Date"}), 400

    url = current_app.config.get("INSURANCE_SIMAS_PAYMENT_URL",
                                 "https://ifc.sfintech.id:38888/Simas/Payment")
    verify_ssl = current_app.config.get("INSURANCE_H2H_VERIFY_SSL", True)
    timeout = current_app.config.get("INSURANCE_H2H_TIMEOUT", 30)

    try:
        with httpx.Client(verify=verify_ssl, timeout=timeout) as client:
            resp = client.post(url, data={"batchno": batchno, "pmt_dt": pmt_dt})
        data = resp.json()
    except Exception:
        current_app.logger.exception("Simas payment call failed for batchno=%s", batchno)
        return jsonify({"success": False, "message": "Payment service is unreachable"}), 502

    if not isinstance(data, dict):
        data = {}
    if str(data.get('status')) == '1':
        return jsonify({"success": True, "message": "Successful"})
    return jsonify({"success": False, "message": data.get('message') or "Failed"}), 400

@insurance_bp.route('/raksa/invoice')
@login_required
def get_raksa_invoice_summary():
    return _invoice_summary('raksa')

@insurance_bp.route('/raksa/detail')
@login_required
def get_raksa_invoice_detail():
    return _invoice_detail('raksa')

@insurance_bp.route('/raksa/confirm', methods=['POST'])
@login_required
def confirm_raksa_invoice():
    return _confirm_invoice('raksa')

def _parse_period(raw) -> Optional[date]:
    raw = (raw or '').strip()
    year = month = None
    m = re.fullmatch(r'(\d{4})-(\d{1,2})', raw)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r'([A-Za-z]{3})[A-Za-z]*-(\d{4})', raw)
        if m:
            month, year = _MONTHS.get(m.group(1).lower()), int(m.group(2))
    if not year or not month or not 1 <= month <= 12:
        return None
    return date(year, month, calendar.monthrange(year, month)[1])

def _read_sheet_rows(data: bytes, ext: str) -> list:
    if ext == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()

    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            elif cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value)
        rows.append(tuple(row))
    return rows

def _col(row, idx):
    return row[idx] if idx < len(row) else None

def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else format(Decimal(repr(v)), 'f')
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    return str(v)

def _date_str(v) -> Optional[str]:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date().isoformat()
        except (OverflowError, ValueError):
            return None
    if isinstance(v, str) and v.strip():
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d-%b-%Y', '%d %b %Y'):
            try:
                return datetime.strptime(v.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None

@insurance_bp.route('/raksa/upload', methods=['POST'])
@login_required
def upload_raksa_invoice():
    upload = request.files.get('invoiceRaksa')
    period_raw = request.form.get('period', '')

    if not upload or not upload.filename or not period_raw.strip():
        return jsonify({"success": False, "message": "No file uploaded."}), 400

    ext = upload.filename.rsplit('.', 1)[-1].lower() if '.' in upload.filename else ''
    if ext not in ('xls', 'xlsx'):
        return jsonify({"success": False, "message": "Only .xls and .xlsx files are allowed."}), 400

    as_of = _parse_period(period_raw)
    if as_of is None:
        return jsonify({"success": False, "message": "Invalid As Of period."}), 400

    data = upload.read(MAX_UPLOAD_BYTES + 1)
    if not data:
        return jsonify({"success": False, "message": "Error uploading the file."}), 400
    if len(data) > MAX_UPLOAD_BYTES:
        return jsonify({"success": False, "message": "File size exceeds the 25MB limit."}), 400

    try:
        rows = _read_sheet_rows(data, ext)
    except ImportError:
        return jsonify({"success": False,
                        "message": "Reading .xls needs the 'xlrd' package. Save the file as .xlsx or install xlrd."}), 400
    except Exception:
        current_app.logger.exception("Raksa upload: cannot read workbook")
        return jsonify({"success": False, "message": "Error loading file."}), 400

    empty = (None, '', 0, '0', False)
    last_row = 0
    for n, row in enumerate(rows, 1):
        if _col(row, 0) not in empty:
            last_row = n

    as_of_str = as_of.isoformat()
    params = []
    for n in range(7, last_row + 1):
        row = rows[n - 1]
        lease_no = _cell_str(_col(row, 2))
        if not lease_no:
            continue
        params.append({
            "asof": as_of_str,
            "policy_no": _cell_str(_col(row, 0)),
            "insured": _cell_str(_col(row, 1)),
            "lease_no": lease_no,
            "due_date": _date_str(_col(row, 3)),
            "gross": _cell_str(_col(row, 4)),
            "comm": _cell_str(_col(row, 5)),
            "ppn": _cell_str(_col(row, 6)),
            "pph": _cell_str(_col(row, 7)),
            "adm": _cell_str(_col(row, 9)),
            "net": _cell_str(_col(row, 10)),
        })

    if not params:
        return jsonify({"success": False,
                        "message": "No data rows found (data must start at row 7)."}), 400

    try:
        db.session.execute(text("DELETE FROM INVOICE_RAKSA WHERE AsOf = :asof"), {"asof": as_of_str})
        db.session.execute(
            text("INSERT INTO INVOICE_RAKSA "
                 "(AsOf, PolicyNo, InsuredName, Lease_No, DueDate, GrossPremi, BrokerComm, PPN, PPh23_21, AdmCost, Net) "
                 "VALUES (:asof, :policy_no, :insured, :lease_no, :due_date, :gross, :comm, :ppn, :pph, :adm, :net)"),
            params,
        )
        db.session.execute(text("DELETE FROM INVOICE_RAKSA WHERE Lease_No = ''"))
        db.session.commit()
        return jsonify({"success": True, "inserted": len(params)})

    except Exception:
        db.session.rollback()
        current_app.logger.exception("Raksa upload error")
        return jsonify({"success": False, "message": "Failed to save data."}), 500
