from flask import Blueprint, jsonify, request, current_app, session, make_response
from app import db
from app.auth.decorators import login_required
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils  import get_column_letter

jurnal_financing_bp = Blueprint('jurnal_financing', __name__, url_prefix='/JurnalFinancing')

def _exec_sp(sp_name: str, params: list) -> list[dict]:
   
    placeholders = ', '.join(['?'] * len(params))
    sql          = f"EXEC {sp_name} {placeholders}"

    raw_conn = db.session.connection().connection
    cursor   = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        while cursor.description is None:
            if not cursor.nextset():
                return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

_THIN = Side(style='thin', color='000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT_DATA = Font(name='Tahoma', size=7)

def _apply_default(cell):
    
    cell.border = _BORDER
    cell.font   = _FONT_DATA

def _apply_right(cell):
    
    cell.border    = _BORDER
    cell.font      = _FONT_DATA
    cell.alignment = Alignment(horizontal='right', vertical='center')

def _apply_left(cell):
    
    cell.border    = _BORDER
    cell.font      = _FONT_DATA
    cell.alignment = Alignment(horizontal='left', vertical='center')

def _apply_header(cell):
    
    cell.border    = _BORDER
    cell.font      = Font(name='Tahoma', size=7, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

_COLUMNS = [
    ('A',  ' No. ',                         None,                  _apply_default),
    ('B',  ' Tgl Pelunasan ',               'Tanggal',             _apply_default),
    ('C',  ' No. CAM ',                     'BookingNo',           _apply_right  ),
    ('D',  ' No. Kontrak ',                 'ContractNo',          _apply_right  ),
    ('E',  ' Nama Debitur ',                'NamaDebitur',         _apply_default),
    ('F',  ' NIK Debitur ',                 'NIK',                 _apply_default),
    ('G',  ' Group ',                       'Grup',                _apply_default),
    ('H',  ' Tempat, Tgl Lahir ',           '__BIRTH__',           _apply_default),
    ('I',  ' Tenor ',                       'tenor',               _apply_default),
    ('J',  ' Merek ',                       'Brand',               _apply_default),
    ('K',  ' Tipe ',                        'TypeSeries',          _apply_default),
    ('L',  ' Tahun Kendaraan ',             'TAHUN',               _apply_default),
    ('M',  ' No. Mesin ',                   'ENGINE',              _apply_default),
    ('N',  ' No. Rangka ',                  'CHASIS',              _apply_default),
    ('O',  ' Roda Kendaraan ',              'WHEELS',              _apply_default),
    ('P',  ' No. BPKB ',                    'NoBPKB',              _apply_default),
    ('Q',  ' Nama BPKB ',                   'NamaBPKB',            _apply_default),
    ('R',  ' No. Polisi ',                  'POLICENO',            _apply_default),
    ('S',  ' Jenis Pembiayaan ',            'JenisPembiayaan',     _apply_default),
    ('T',  ' Arr/ Adv ',                    'AdvArr',              _apply_default),
    ('U',  ' NTF ',                         'ntf',                 _apply_default),
    ('V',  ' Harga Kendaraan ',             'HargaKendaraan',      _apply_default),
    ('W',  ' Down Payment ',                'DownPayment',         _apply_default),
    ('X',  ' Biaya Provisi ',               'BiayaProvisi',        _apply_default),
    ('Y',  ' Biaya Admin ',                 'BiayaAdmin',          _apply_default),
    ('Z',  ' Biaya Notaris ',               'Notary',              _apply_default),
    ('AA', ' Premi Asuransi ',              'PremiIns',            _apply_default),
    ('AB', ' Jenis Asuransi ',              'JenisAsuransi',       _apply_default),
    ('AC', ' CLI ',                         'CLI',                 _apply_default),
    ('AD', ' Bunga Jual ',                  'SaleInterest',        _apply_default),
    ('AE', ' Total DP ',                    'TotalDP',             _apply_default),
    ('AF', ' Struktur Pembiayaan ',         'StrukturPembiayaan',  _apply_default),
    ('AG', ' Angsuran ',                    'Angsuran',            _apply_left   ),
    ('AH', ' Pelunasan Dealer ',            'PelunasanDealer',     _apply_default),
    ('AI', ' Total Insentif Sales ',        'tot_komisi_sales',    _apply_default),
    ('AJ', ' Total Insentif Dealer ',       'tot_komisi_dealer',   _apply_default),
    ('AK', ' Asuransi ',                    'Asuransi',            _apply_default),
    ('AL', ' Tipe Asuransi ',               'TipeAsuransi',        _apply_default),
    ('AM', ' Contract Receivable (Dr) ',    'OUTS_REC',            _apply_right  ),
    ('AN', ' UI Income (Cr) ',              'UN_INCOME',           _apply_right  ),
    ('AO', ' A/P Customers (Cr) ',          'APCust',              _apply_right  ),
    ('AP', ' Transaction Code ',            'cont_type',           _apply_default),
]

LAST_COL = _COLUMNS[-1][0]

def _fmt_date(val) -> str:
   
    if isinstance(val, (date, datetime)):
        return val.strftime('%d-%m-%Y')
    return str(val) if val is not None else ''

def _build_excel(rows: list[dict], date_from: str, date_to: str,
                 company: str, username: str) -> bytes:
   
    wb = Workbook()
    ws = wb.active
    ws.title = 'JurnalFinancing'

    ws.merge_cells(f'A1:AO1')
    ws.merge_cells(f'A2:AO2')
    cell_company           = ws['A2']
    cell_company.value     = company
    cell_company.font      = Font(name='Tahoma', size=10)
    cell_company.alignment = Alignment(horizontal='left', vertical='center')

    ws['AP1'].value     = f'Date : {date.today().strftime("%d %b %Y")}'
    ws['AP1'].font      = Font(name='Tahoma', size=9)
    ws['AP1'].alignment = Alignment(horizontal='left')

    ws['AP2'].value     = f'Print By : {username}'
    ws['AP2'].font      = Font(name='Tahoma', size=9)
    ws['AP2'].alignment = Alignment(horizontal='left')

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 16

    ws.merge_cells(f'A3:{LAST_COL}3')
    cell_title            = ws['A3']
    cell_title.value      = 'JURNAL FINANCING'
    cell_title.font       = Font(name='Tahoma', bold=True, size=12)
    cell_title.alignment  = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 20

    def _disp(d: str) -> str:
        try:
            return datetime.strptime(d, '%Y-%m-%d').strftime('%d-%m-%Y')
        except ValueError:
            return d

    ws.merge_cells(f'A4:{LAST_COL}4')
    cell_period            = ws['A4']
    cell_period.value      = f'Disbursement Date {_disp(date_from)} to {_disp(date_to)}'
    cell_period.font       = Font(name='Tahoma', size=10)
    cell_period.alignment  = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[4].height = 16

    ws.row_dimensions[5].height = 30
    for col_idx, (col_letter, label, _, __) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=5, column=col_idx, value=label)
        _apply_header(cell)
       
        width = max(len(label.strip()) + 2, 5)
        ws.column_dimensions[col_letter].width = width

    ws.column_dimensions['A'].width = 5

    for no, r in enumerate(rows, start=1):
        excel_row = no + 5

        for col_idx, (_, __, sp_key, style_fn) in enumerate(_COLUMNS, start=1):
            if sp_key is None:
                value = no
            elif sp_key == '__BIRTH__':
                place = r.get('PLACEBIRTH', '') or ''
                tgl   = r.get('TGLBIRTH',   '') or ''
                value = f"{place}, {tgl}" if place or tgl else ''
            else:
                raw   = r.get(sp_key)
                value = _fmt_date(raw) if isinstance(raw, (date, datetime)) else (raw if raw is not None else '')

            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            style_fn(cell)

    ws.freeze_panes = 'A6'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

@jurnal_financing_bp.route('/excel')
@login_required
def export_excel():
    
    date_from = request.args.get('date_from', '').strip()
    date_to   = request.args.get('date_to',   '').strip()

    if not date_from or not date_to:
        return jsonify({'error': 'date_from and date_to are required'}), 400

    username = session.get('username', '')
    company  = session.get('company',  '')

    try:
        
        rows = _exec_sp('USP_JURNAL_FINANCING', [date_from, date_to])

        xlsx_bytes = _build_excel(rows, date_from, date_to, company, username)
        filename   = 'JurnalFinancing.xlsx'

        response = make_response(xlsx_bytes)
        response.headers['Content-Type']        = (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except Exception as e:
        current_app.logger.exception("Jurnal Financing export error")
        return jsonify({'error': str(e)}), 500