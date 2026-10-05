from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from flask import Blueprint, request, jsonify, send_file, current_app, make_response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app import db
from app.auth.decorators import login_required

amortization_report_bp = Blueprint("amortization_report", __name__, url_prefix="/AmortizationReport")

_LEASE_RE = re.compile(r"^[A-Za-z0-9._/\-]{1,30}$")

_MONEY = ("RENTAL", "ACCRUAL1", "ACCRUAL2", "L_INCOME", "I_RECOVERY", "OUTS_PRINC", "UN_INCOME", "OUTS_REC",
          "L_INCOME_PSAK", "I_RECOVERY_PSAK", "OUTS_PRINC_PSAK", "UN_INCOME_PSAK", "ATRIBUTE")

_SQL_SCHEDULE = """
SELECT a.PERIOD, a.DATE_DUE, a.RENTAL, a.ACCRUAL1, a.ACCRUAL2, a.L_INCOME, a.I_RECOVERY,
       a.OUTS_PRINC, a.UN_INCOME, a.OUTS_REC, a.DATE_PAY,
       a.L_INCOME_PSAK, a.I_RECOVERY_PSAK, a.OUTS_PRINC_PSAK, a.UN_INCOME_PSAK, a.ATRIBUTE,
       c.lessee_nm, d.declrate, e.irr
FROM CLFPAY a WITH (NOLOCK)
    INNER JOIN clfcont_data b WITH (NOLOCK) ON a.lease_no = b.lease_no
    INNER JOIN clfless c WITH (NOLOCK) ON b.lessee_no = c.lessee_no
    INNER JOIN cafcont d WITH (NOLOCK) ON b.appl_no = d.appl_no
    INNER JOIN clfatr e WITH (NOLOCK) ON a.lease_no = e.lease_no
WHERE a.lease_no = :lease_no
ORDER BY a.PERIOD
"""

_SQL_HEADER = """
SELECT a.LESSEE_NM, a.tenor, a.tipe_pmt, a.L_AMOUNT, a.dp, a.FLATRATE, a.DECLRATE,
       CONCAT(a.BRAND_NM, ' ', a.tipe_mobil) AS fin_item, a.ins_loan,
       b.Otherloan, b.OVERATE, c.tot_install, c.tot_interest, c.tot_recovery
FROM V_DataKredit a
    LEFT JOIN CLFCONT b ON a.lease_no = b.lease_no
    LEFT JOIN (SELECT SUM(RENTAL) AS tot_install, SUM(L_INCOME) AS tot_interest,
                    SUM(I_RECOVERY) AS tot_recovery, lease_no
            FROM CLFPAY WITH (NOLOCK) GROUP BY lease_no) c ON a.lease_no = c.lease_no
WHERE a.lease_no = :lease_no
"""

def _lease_no():
    raw = (request.args.get("lease_no") or "").strip()
    if not _LEASE_RE.match(raw):
        return None, (jsonify({"success": False, "message": "Please enter a valid Contract No."}), 400)
    return raw, None

def _g(row: dict, key: str):
    return row.get(key.lower())

def _num(v) -> float:
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0

def _idr(v) -> str:
    if v is None or v == "":
        return ""
    n = int(Decimal(str(v)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return f"{n:,}".replace(",", ".")

def _dmy(v) -> str:
    if isinstance(v, (datetime, date)):
        mon = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")[v.month - 1]
        return f"{v.day:02d}-{mon}-{v.year}"
    return "" if v is None else str(v)

def _rate(v) -> str:
    if v in (None, ""):
        return ""
    return f"{round(_num(v), 2):g}"

def _company() -> str:
    return current_app.config.get("COMPANY_NAME", "") or ""

def _load_schedule(lease_no: str):
    res = db.session.execute(text(_SQL_SCHEDULE), {"lease_no": lease_no})
    rows = [{str(k).lower(): v for k, v in dict(r).items()} for r in res.mappings().all()]
    meta = {"customerName": "", "declrate": 0, "irr": 0}
    if rows:
        last = rows[-1]
        meta = {
            "customerName": _g(last, "lessee_nm") or "",
            "declrate": round(_num(_g(last, "declrate")), 2),
            "irr": round(_num(_g(last, "irr")), 2),
        }
    return meta, rows

def _fail(msg: str, code: int):
    return jsonify({"success": False, "message": msg}), code

def _get_schedule_or_error(lease_no: str):
    try:
        meta, rows = _load_schedule(lease_no)
    except Exception:
        db.session.rollback()
        current_app.logger.exception("amortization query failed")
        return None, None, _fail("Failed to load amortization data", 500)
    if not rows:
        return None, None, _fail("Contract not found or has no amortization schedule", 404)
    return meta, rows, None

@amortization_report_bp.route("/get-data", methods=["GET"])
@login_required
def amortization_preview():
    lease_no, err = _lease_no()
    if err:
        return err
    meta, rows, err = _get_schedule_or_error(lease_no)
    if err:
        return err

    out = []
    for r in rows:
        item = {
            "PERIOD": int(_num(_g(r, "PERIOD"))),
            "DueDate": _dmy(_g(r, "DATE_DUE")),
            "Payment": _dmy(_g(r, "DATE_PAY")),
        }
        for k in _MONEY:
            item[k] = _idr(_g(r, k))
        out.append(item)
    return jsonify({**meta, "rows": out})

_XL_HEAD1 = [("Period", 2), ("Due Date", 2), ("Installment", 2)]
_XL_SUB = ["Accrual 1", "Accrual 2", "Interest Income", "Recovery", "Outstanding Principal",
           "Unearned Income", "Outstanding Receivable"]
_XL_PSAK_SUB = ["Interest Income", "Recovery", "Outstanding Principal", "Unearned Income", "Attribute"]
_XL_KEYS = ["PERIOD", "DATE_DUE", "RENTAL", "ACCRUAL1", "ACCRUAL2", "L_INCOME", "I_RECOVERY", "OUTS_PRINC",
            "UN_INCOME", "OUTS_REC", "DATE_PAY", "L_INCOME_PSAK", "I_RECOVERY_PSAK", "OUTS_PRINC_PSAK",
            "UN_INCOME_PSAK", "ATRIBUTE"]

@amortization_report_bp.route("/export", methods=["GET"])
@login_required
def amortization_export():
    lease_no, err = _lease_no()
    if err:
        return err
    meta, rows, err = _get_schedule_or_error(lease_no)
    if err:
        return err

    thin = Side(style="thin")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    head_font = Font(name="Tahoma", size=10, bold=True)
    body_font = Font(name="Tahoma", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Amortization"

    def head(col, row, value, rowspan=1, colspan=1):
        if rowspan > 1 or colspan > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row + rowspan - 1, end_column=col + colspan - 1)
        for rr in range(row, row + rowspan):
            for cc in range(col, col + colspan):
                ws.cell(row=rr, column=cc).border = box
        c = ws.cell(row=row, column=col, value=value)
        c.font = head_font
        c.alignment = center

    head(1, 1, "Period", rowspan=2)
    head(2, 1, "Due Date", rowspan=2)
    head(3, 1, "Installment", rowspan=2)
    head(4, 1, f"Contractual ({meta['declrate']:g}%)", colspan=7)
    head(11, 1, "Payment Date", rowspan=2)
    head(12, 1, f"PSAK ({meta['irr']:g}%)", colspan=5)
    for i, t in enumerate(_XL_SUB):
        head(4 + i, 2, t)
    for i, t in enumerate(_XL_PSAK_SUB):
        head(12 + i, 2, t)

    r = 3
    for row in rows:
        for i, key in enumerate(_XL_KEYS, start=1):
            v = _g(row, key)
            if key in ("DATE_DUE", "DATE_PAY"):
                v = _dmy(v)
            elif key == "PERIOD":
                v = int(_num(v))
            else:
                v = _num(v) if v is not None else None
            c = ws.cell(row=r, column=i, value=v)
            c.font = body_font
            c.border = box
            if key not in ("PERIOD", "DATE_DUE", "DATE_PAY"):
                c.number_format = "#,##0"
        r += 1

    for i in range(1, 17):
        ws.column_dimensions[get_column_letter(i)].width = 8 if i == 1 else 16
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", lease_no)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"Amortization_{safe}.xlsx",
    )

def _pdf_response(data: bytes, filename: str):
    resp = make_response(data)
    resp.headers["Content-Type"] = "application/pdf"
    resp.headers["Content-Disposition"] = f'inline; filename="{filename}"'
    resp.headers["Content-Length"] = str(len(data))
    resp.headers["Cache-Control"] = "no-store"
    return resp

def _build_pdf(kind: str, lease_no: str, meta: dict, rows: list, h: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=8 * mm, rightMargin=8 * mm,
                            topMargin=8 * mm, bottomMargin=8 * mm, title="Amortization Report")
    base = getSampleStyleSheet()["Normal"]
    title_s = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=10, alignment=TA_CENTER, spaceAfter=4)
    info_s = ParagraphStyle("i", parent=base, fontName="Helvetica", fontSize=8)
    GREY = colors.HexColor("#f2f2f2")

    def style(n_cols, spans, bold_last=False, fs=6.5, aligns=None):
        st = [
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), fs),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.black),
            ("BACKGROUND", (0, 0), (-1, 1), GREY),
            ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, 1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ]
        st += [("SPAN",) + sp for sp in spans]
        if bold_last:
            st.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        for col, al in (aligns or {}).items():
            st.append(("ALIGN", (col, 2), (col, -1), al))
        return TableStyle(st)

    declr, irr = f"{meta['declrate']:g}", f"{meta['irr']:g}"
    story = []

    if kind == "all":
        d = [["Period", "Due Date", "Installment", f"Contractual ({declr}%)", "", "", "", "", "", "",
              "Payment Date", f"PSAK ({irr}%)", "", "", "", ""],
             ["", "", ""] + _XL_SUB + [""] + _XL_PSAK_SUB]
        for r in rows:
            d.append([str(int(_num(_g(r, "PERIOD")))) if k == "PERIOD" else
                      (_dmy(_g(r, k)) if k in ("DATE_DUE", "DATE_PAY") else _idr(_g(r, k)))
                      for k in _XL_KEYS])
        spans = [((0, 0), (0, 1)), ((1, 0), (1, 1)), ((2, 0), (2, 1)), ((3, 0), (9, 0)),
                 ((10, 0), (10, 1)), ((11, 0), (15, 0))]

        d[1][10] = ""
        aligns = {i: "RIGHT" for i in list(range(2, 10)) + list(range(11, 16))}
        aligns.update({0: "CENTER", 1: "CENTER", 10: "CENTER"})

        hs = ParagraphStyle("hs", parent=base, fontName="Helvetica-Bold", fontSize=5.5, leading=6.5, alignment=TA_CENTER)
        d[1] = ["", "", ""] + [Paragraph(x, hs) for x in _XL_SUB] + [""] + [Paragraph(x, hs) for x in _XL_PSAK_SUB]
        t = Table(d, colWidths=[12 * mm, 20 * mm, 22 * mm] + [17 * mm] * 7 + [20 * mm] + [17 * mm] * 5, repeatRows=2)
        t.setStyle(style(16, spans, aligns=aligns, fs=6))
        story.append(Paragraph(f"AMORTIZATION SCHEDULE - {escape_xml(lease_no)} ({escape_xml(meta['customerName'])})", title_s))
        story.append(t)
        doc.build(story)
        return buf.getvalue()

    def n0(v):
        return _idr(v if v is not None else 0)

    def info(l1, v1, l2, v2):
        return [l1, f": {v1 if v1 is not None else ''}", l2, f": {v2 if v2 is not None else ''}"]

    overate = _num(h.get("overate")) / 100
    info_rows = [
        info("Contract No.", lease_no, "Tenor", h.get("tenor")),
        info("Customer Name", h.get("lessee_nm") or meta["customerName"], "Flat Rate (% p.a.)", _rate(h.get("flatrate"))),
        info("Financing Item", h.get("fin_item"), "Decl Rate (% p.a.)", declr),
        info("Method of Payment", h.get("tipe_pmt"), "Overdue Rate", f"{overate:g}"),
        info("Asset Value", n0(h.get("l_amount")), "Insurance Loan", n0(h.get("ins_loan"))),
        info("Down Payment", n0(h.get("dp")), "Other Loan", n0(h.get("otherloan"))),
    ]
    it = Table(info_rows, colWidths=[35 * mm, 100 * mm, 40 * mm, 60 * mm])
    it.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                            ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
    story += [Paragraph("Amortization Schedule in Indonesian Rupiah", title_s), it, Spacer(1, 4 * mm)]

    keys = ["PERIOD", "DATE_DUE", "RENTAL", "L_INCOME", "I_RECOVERY", "OUTS_PRINC", "UN_INCOME", "OUTS_REC", "DATE_PAY"]
    d = [["Period", "Due Date", f"Contractual ({declr}%)", "", "", "Outstanding Principal", "Unearned Income",
          "Outstanding Receivable", "Payment Date"],
         ["", "", "Installment", "Interest Income", "Recovery", "", "", "", ""]]
    for r in rows:
        d.append([str(int(_num(_g(r, k)))) if k == "PERIOD" else
                  (_dmy(_g(r, k)) if k in ("DATE_DUE", "DATE_PAY") else _idr(_g(r, k))) for k in keys])
    d.append(["TOTAL", "", n0(h.get("tot_install")), n0(h.get("tot_interest")), n0(h.get("tot_recovery")), "", "", "", ""])
    last = len(d) - 1
    spans = [((0, 0), (0, 1)), ((1, 0), (1, 1)), ((2, 0), (4, 0)), ((5, 0), (5, 1)), ((6, 0), (6, 1)),
             ((7, 0), (7, 1)), ((8, 0), (8, 1)), ((0, last), (1, last)), ((5, last), (8, last))]
    aligns = {2: "RIGHT", 3: "RIGHT", 4: "RIGHT", 5: "RIGHT", 6: "RIGHT", 7: "RIGHT", 0: "CENTER", 1: "CENTER", 8: "CENTER"}
    t = Table(d, colWidths=[18 * mm, 28 * mm, 32 * mm, 32 * mm, 32 * mm, 36 * mm, 32 * mm, 36 * mm, 28 * mm], repeatRows=2)
    ts = style(9, spans, bold_last=True, aligns=aligns, fs=7)
    ts.add("ALIGN", (0, last), (-1, last), "RIGHT")
    t.setStyle(ts)
    story.append(t)
    doc.build(story)
    return buf.getvalue()

def escape_xml(v) -> str:
    return str(v if v is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

@amortization_report_bp.route("/print", methods=["GET"])
@login_required
def amortization_print():
    lease_no, err = _lease_no()
    if err:
        return err
    kind = (request.args.get("type") or "all").strip().lower()
    if kind not in ("all", "contractual"):
        return _fail("Invalid print type", 400)

    meta, rows, err = _get_schedule_or_error(lease_no)
    if err:
        return err

    h = {}
    if kind == "contractual":
        try:
            r = db.session.execute(text(_SQL_HEADER), {"lease_no": lease_no}).mappings().first()
        except Exception:
            db.session.rollback()
            current_app.logger.exception("amortization header query failed")
            return _fail("Failed to load amortization data", 500)
        h = {str(k).lower(): v for k, v in dict(r).items()} if r else {}

    try:
        data = _build_pdf(kind, lease_no, meta, rows, h)
    except Exception as e:
        current_app.logger.exception("amortization pdf build failed")
        import traceback
        return jsonify({"success": False, "message": f"PDF error: {type(e).__name__}: {e}",
                        "trace": traceback.format_exc()}), 500
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", lease_no)
    return _pdf_response(data, f"Amortization_{safe}_{kind}.pdf")