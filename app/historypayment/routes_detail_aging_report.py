from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from flask import Blueprint, request, jsonify, session, current_app, make_response

from app import db
from app.auth.decorators import login_required

detail_aging_report_bp = Blueprint("detail_aging_report_bp", __name__, url_prefix="/DetailAgingReport")

AGING_DETAIL_COLUMNS = [
    ("NO.", "no_urut", "int", 8),
    ("NAMA CUSTOMER", "LESSEE_NM", "text", 30),
    ("NO. CUSTOMER", "LESSEE_NO", "text", 14),
    ("NO. KONTRAK", "LEASE_NO", "text", 18),
    ("ALAMAT", "ALAMAT", "text", 36),
    ("NO. HP", "HP", "text", 17),
    ("TGL JTH TEMPO", "GIRODT", "date", 14),
    ("ANGSURAN TERTUNGGAK", "nominal", "money", 14),
    ("DAYS", "HARI", "int", 8),
    ("BUCKET", "BUCKET", "text", 10),
    ("OUTSTANDING PRINCIPAL", "OUTS_PRINC", "money", 16),
    ("OUTSTANDING PRINCIPAL (CONTRACTUAL)", "os_contractual", "money", 16),
    ("INST", "periode", "int", 7),
    ("TENOR", "TENOR", "int", 8),
    ("MEREK/ TIPE KENDARAAN", "KENDARAAN", "text", 26),
    ("NO. POLISI", "POLICENO", "text", 14),
    ("RED CODE", "RedCode", "text", 10),
    ("CONDITION", "CONDITION", "text", 12),
    ("PENALTI", "penalty", "money", 11),
    ("DEALER", "DEALER", "text", 26),
    ("CMO", "CMO", "text", 20),
    ("KODE POS", "KODEPOS", "text", 10),
    ("COLLECTOR", None, "blank", 12),
    ("LATE CONDITION", "Catatan", "text", 18),
    ("BRANCH NAME", "BRANCH_NAME", "text", 18),
    ("CONTRACT TYPE", "cont_type", "text", 14),
    ("PURPOSE OF FINANCE", "fin_purpose", "text", 16),
    ("STATUS JFG", "Status_JFG", "text", 10),
    ("INSTALLMENT", "RENTAL", "money", 13),
    ("TANGGAL DISBURSED", "EXECUTION", "date", 14),
    ("NTF", "ntf", "money", 14),
    ("Type Of Payment", "PaymentType", "text", 14),
    ("OUTSTANDING INCOME", "UN_INCOME", "money", 14),
    ("OUTSTANDING INCOME (CONTRACTUAL)", "income_contractual", "money", 16),
]
_NCOLS = len(AGING_DETAIL_COLUMNS)
_TOTAL_KEYS = ("nominal", "OUTS_PRINC", "os_contractual", "penalty", "UN_INCOME", "income_contractual")

def _parse_asof(raw: str):
    raw = (raw or "").strip()
    for fmt in ("%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None

def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

def _exec_sp(asof: date, branch_cd: str, type_of: str):
    cursor = db.session.connection().connection.cursor()
    try:
        cursor.execute("EXEC USP_AGING_DETAIL ?, ?, ?", [asof.isoformat(), branch_cd, type_of])
        while cursor.description is None:
            if not cursor.nextset():
                return {}, []
        idx = {str(c[0]).lower(): i for i, c in enumerate(cursor.description)}
        return idx, cursor.fetchall()
    finally:
        cursor.close()

def _plan(idx: dict):
    return [idx.get(key.lower()) if key else None for _h, key, _k, _w in AGING_DETAIL_COLUMNS]

def _totals(rows, pos):
    out = {}
    for n, (_h, key, _k, _w) in enumerate(AGING_DETAIL_COLUMNS):
        if key in _TOTAL_KEYS:
            p = pos[n]
            out[key] = sum(_num(r[p]) for r in rows) if p is not None else 0.0
    return out

def _fmt_date(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%d-%m-%Y")
    return "" if v is None else str(v)

def _round_half_up(n: float) -> str:
    return f"{int(Decimal(str(n)).quantize(Decimal('1'), rounding=ROUND_HALF_UP)):,}"

def _text(v) -> str:
    return "" if v is None else str(v)

def _validate_request():
    asof = _parse_asof(request.args.get("asof", ""))
    if not asof:
        return None, None, None, (jsonify({"success": False, "message": "Please select date first"}), 400)
    return asof, (request.args.get("branch_cd") or "").strip(), (request.args.get("type_of") or "").strip(), None

def _company() -> str:
    return current_app.config.get("COMPANY_NAME", "") or session.get("company", "") or ""

def _load(asof, branch_cd, type_of):
    try:
        idx, rows = _exec_sp(asof, branch_cd, type_of)
    except Exception:
        db.session.rollback()
        current_app.logger.exception("aging detail query failed")
        return None, None, (jsonify({"success": False, "message": "Failed to generate the report"}), 500)
    return _plan(idx), rows, None

def _file_response(data: bytes, mimetype: str, disposition: str):
    resp = make_response(data)
    resp.headers["Content-Type"] = mimetype
    resp.headers["Content-Disposition"] = disposition
    resp.headers["Content-Length"] = str(len(data))
    resp.headers["Cache-Control"] = "no-store"
    return resp

def _build_xlsx(rows, pos, totals, asof: date) -> bytes:
    try:
        import xlsxwriter
    except ImportError:
        return _build_xlsx_openpyxl(rows, pos, totals, asof)

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"constant_memory": True})
    ws = wb.add_worksheet("DetailAging")

    f_head = wb.add_format({"font_name": "Tahoma", "font_size": 8, "border": 1})
    f_text = wb.add_format({"font_name": "Tahoma", "font_size": 7, "border": 1})
    f_wrap = wb.add_format({"font_name": "Tahoma", "font_size": 7, "border": 1, "text_wrap": True})
    f_right = wb.add_format({"font_name": "Tahoma", "font_size": 7, "border": 1, "align": "right"})
    f_money = wb.add_format({"font_name": "Tahoma", "font_size": 7, "border": 1, "num_format": "#,##0"})
    f_title = wb.add_format({"font_name": "Tahoma", "font_size": 10, "bold": True, "align": "center", "valign": "vcenter"})
    f_center = wb.add_format({"font_name": "Tahoma", "font_size": 8, "align": "center"})
    f_plain = wb.add_format({"font_name": "Tahoma", "font_size": 8})

    for i, (header, key, _kind, _w) in enumerate(AGING_DETAIL_COLUMNS):
        ws.set_column(i, i, 5 if i == 0 else (20 if i == 4 else max(len(header) + 2, 8)))

    last = _NCOLS - 1
    ws.merge_range(0, 0, 0, last - 1, "", f_plain)
    ws.write(0, last, "Date : " + date.today().strftime("%d %b %Y"), f_plain)
    ws.merge_range(1, 0, 1, last - 1, _company(), f_plain)
    ws.write(1, last, "Print By : " + (session.get("username") or ""), f_plain)
    ws.merge_range(2, 0, 2, last, "DETAIL AGING REPORT", f_title)
    ws.merge_range(3, 0, 3, last, "As of " + asof.strftime("%d-%m-%Y"), f_center)
    for i, (header, _k, _kd, _w) in enumerate(AGING_DETAIL_COLUMNS):
        ws.write_string(4, i, f" {header} ", f_head)

    plan = [(n, p, kind, key) for n, (_h, key, kind, _w) in enumerate(AGING_DETAIL_COLUMNS)
            for p in (pos[n],)]
    r = 5
    for row in rows:
        for n, p, kind, key in plan:
            if p is None or kind == "blank":
                ws.write_blank(r, n, None, f_text)
                continue
            v = row[p]
            if v is None or v == "":
                ws.write_blank(r, n, None, f_text)
            elif kind == "money":
                ws.write_number(r, n, _num(v), f_money)
            elif kind == "int":
                ws.write_number(r, n, int(_num(v)), f_text)
            elif kind == "date":
                ws.write_string(r, n, _fmt_date(v), f_text)
            elif key == "LEASE_NO":
                ws.write_string(r, n, str(v), f_right)
            elif key == "ALAMAT":
                ws.write_string(r, n, str(v), f_wrap)
            else:
                ws.write_string(r, n, str(v), f_text)
        r += 1
    for n, (_h, key, _kind, _w) in enumerate(AGING_DETAIL_COLUMNS):
        if key in _TOTAL_KEYS:
            ws.write_number(r, n, totals[key], f_money)
    wb.close()
    return buf.getvalue()

def _build_xlsx_openpyxl(rows, pos, totals, asof: date) -> bytes:
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("DetailAging")
    ws.append([""] * (_NCOLS - 1) + ["Date : " + date.today().strftime("%d %b %Y")])
    ws.append([_company()] + [""] * (_NCOLS - 2) + ["Print By : " + (session.get("username") or "")])
    ws.append(["DETAIL AGING REPORT"])
    ws.append(["As of " + asof.strftime("%d-%m-%Y")])
    ws.append([f" {h} " for h, _k, _kd, _w in AGING_DETAIL_COLUMNS])
    for row in rows:
        out = []
        for n, (_h, _key, kind, _w) in enumerate(AGING_DETAIL_COLUMNS):
            p = pos[n]
            v = None if p is None or kind == "blank" else row[p]
            if v is None or v == "":
                out.append(None)
            elif kind == "money":
                out.append(_num(v))
            elif kind == "int":
                out.append(int(_num(v)))
            elif kind == "date":
                out.append(_fmt_date(v))
            else:
                out.append(str(v))
        ws.append(out)
    ws.append([totals.get(key) if key in _TOTAL_KEYS else None for _h, key, _k, _w in AGING_DETAIL_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

@detail_aging_report_bp.route("/export/aging-detail", methods=["GET"])
@login_required
def export_aging_detail():
    asof, branch_cd, type_of, err = _validate_request()
    if err:
        return err
    pos, rows, err = _load(asof, branch_cd, type_of)
    if err:
        return err
    try:
        data = _build_xlsx(rows, pos, _totals(rows, pos), asof)
    except Exception:
        current_app.logger.exception("aging detail excel build failed")
        return jsonify({"success": False, "message": "Failed to generate the Excel file"}), 500
    return _file_response(
        data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        'attachment; filename="DetailAging.xlsx"')

def _build_pdf(rows, pos, totals, asof: date) -> bytes:
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    W, H = landscape(A3)
    margin = 8 * mm
    avail = W - 2 * margin
    wsum = sum(c[3] for c in AGING_DETAIL_COLUMNS)
    widths = [avail * c[3] / wsum for c in AGING_DETAIL_COLUMNS]
    xs = [margin]
    for w in widths:
        xs.append(xs[-1] + w)

    FONT, FONT_B, FS, HS = "Helvetica", "Helvetica-Bold", 4.6, 4.0
    ROW_H = 6.2
    pad = 1.2
    company = _company() or "SFI"
    user = session.get("username") or ""
    today = date.today().strftime("%d %b %Y")

    head_lines = [simpleSplit(h, FONT_B, HS, w - 2 * pad) or [""] for (h, *_), w in zip(AGING_DETAIL_COLUMNS, widths)]
    n_head = max(len(x) for x in head_lines)
    HEAD_H = n_head * (HS + 1.2) + 3

    def clip(s: str, maxw: float, font=FONT, size=FS) -> str:
        if not s or stringWidth(s, font, size) <= maxw:
            return s
        lo, hi = 0, len(s)
        while lo < hi:
            mid = (lo + hi + 1) // 2
            if stringWidth(s[:mid], font, size) <= maxw:
                lo = mid
            else:
                hi = mid - 1
        return s[:lo]

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(W, H), pageCompression=1)
    c.setTitle("Detail Aging Report")
    page = [0]
    state = {"y": 0.0}

    def new_page():
        if page[0]:
            c.showPage()
        page[0] += 1
        y = H - margin
        c.setFont(FONT_B, 9)
        c.drawString(margin, y - 8, company)
        c.drawCentredString(W / 2, y - 8, "DETAIL AGING REPORT")
        c.setFont(FONT, 6)
        c.drawRightString(W - margin, y - 6, f"Date : {today}")
        c.drawRightString(W - margin, y - 14, f"Print By : {user}")
        c.drawRightString(W - margin, y - 22, f"Page {page[0]}")
        c.setFont(FONT_B, 8)
        c.drawCentredString(W / 2, y - 20, "As of : " + asof.strftime("%d-%m-%Y"))
        top = y - 30
        c.setLineWidth(0.3)
        c.setFont(FONT_B, HS)
        c.rect(margin, top - HEAD_H, avail, HEAD_H, stroke=1, fill=0)
        for i, lines in enumerate(head_lines):
            if i:
                c.line(xs[i], top, xs[i], top - HEAD_H)
            ty = top - HS - 1.5
            for ln in lines:
                c.drawCentredString((xs[i] + xs[i + 1]) / 2, ty, ln)
                ty -= HS + 1.2
        state["y"] = top - HEAD_H
        c.setFont(FONT, FS)

    def draw_row(vals, bold=False):
        y = state["y"]
        if y - ROW_H < margin + 12:
            close_page_grid()
            new_page()
            y = state["y"]
        font = FONT_B if bold else FONT
        c.setFont(font, FS)
        base = y - ROW_H + 1.8
        for i, (txt, right) in enumerate(vals):
            if txt:
                w = widths[i] - 2 * pad
                txt = clip(txt, w, font)
                if right:
                    c.drawRightString(xs[i + 1] - pad, base, txt)
                else:
                    c.drawString(xs[i] + pad, base, txt)
        c.line(margin, y - ROW_H, margin + avail, y - ROW_H)
        state["y"] = y - ROW_H

    def close_page_grid():
        top = H - margin - 30 - HEAD_H
        c.setLineWidth(0.3)
        for x in xs:
            c.line(x, top, x, state["y"])

    new_page()
    plan = [(n, pos[n], kind, key) for n, (_h, key, kind, _w) in enumerate(AGING_DETAIL_COLUMNS)]
    for row in rows:
        vals = []
        for n, p, kind, key in plan:
            if p is None or kind == "blank":
                vals.append(("", False))
                continue
            v = row[p]
            if v is None:
                vals.append(("", False))
            elif kind == "money":
                n_ = _num(v)
                vals.append((_round_half_up(n_) if n_ else "0", True))
            elif kind == "int":
                vals.append((str(v), key != "no_urut"))
            elif kind == "date":
                vals.append((_fmt_date(v), False))
            else:
                vals.append((str(v), False))
        draw_row(vals)

    foot = []
    for n, (_h, key, _kind, _w) in enumerate(AGING_DETAIL_COLUMNS):
        foot.append((_round_half_up(totals[key]), True) if key in _TOTAL_KEYS else ("", False))
    draw_row(foot, bold=True)
    close_page_grid()
    c.save()
    return buf.getvalue()

@detail_aging_report_bp.route("/print/aging-detail", methods=["GET"])
@login_required
def print_aging_detail():
    asof, branch_cd, type_of, err = _validate_request()
    if err:
        return err
    pos, rows, err = _load(asof, branch_cd, type_of)
    if err:
        return err
    try:
        data = _build_pdf(rows, pos, _totals(rows, pos), asof)
    except Exception as e:
        current_app.logger.exception("aging detail pdf build failed")
        import traceback
        return jsonify({"success": False, "message": f"PDF error: {type(e).__name__}: {e}",
                        "trace": traceback.format_exc()}), 500
    return _file_response(data, "application/pdf", f'inline; filename="DetailAging_{asof.isoformat()}.pdf"')