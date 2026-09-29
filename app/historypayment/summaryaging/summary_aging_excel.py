import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

_THIN   = Side(style='thin', color='000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUM    = '#,##0'

def _ft(bold=False, size=7, name='Tahoma', color='000000'):
    return Font(bold=bold, size=size, name=name, color=color)

def _al(h='left', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _title(cell, value):
    cell.value     = value
    cell.font      = _ft(bold=True, size=12)
    cell.alignment = _al('left', 'bottom')

def _sub(cell, value=None):
    if value is not None:
        cell.value = value
    cell.font      = _ft(bold=True, size=8)
    cell.alignment = _al('center', 'center', wrap=True)
    cell.border    = _BORDER

def _data(cell, value=None, num=False):
    if value is not None:
        cell.value = value
    cell.font      = _ft()
    cell.border    = _BORDER
    if num:
        cell.number_format = _NUM

def _bold_data(cell, value=None, num=False):
    if value is not None:
        cell.value = value
    cell.font      = _ft(bold=True, size=7)
    cell.border    = _BORDER
    if num:
        cell.number_format = _NUM

def _fmt_date(d: str) -> str:

    try:
        return datetime.strptime(d, '%Y-%m-%d').strftime('%d-%m-%Y')
    except ValueError:
        return d

def _safe_pct(num, denom, dec=2):
   
    try:
        return round((num or 0) / denom * 100, dec)
    except (ZeroDivisionError, TypeError):
        return 0.0

def _apply_range(ws, row_start, row_end, col_start, col_end, fn, **kw):
   
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            fn(ws.cell(row=r, column=c), **kw)

_LABELS_4 = ['Acc', 'Amount', 'Outstanding Principal', 'AR Principal']
_LABELS_2 = ['Acc', '%']

_PERIODS_T1 = [
    'Current', '01 s/d 10 Days', '11 s/d 30 Days',
    '31 s/d 60 Days', '61 s/d 90 Days', '> 90 Days',
]
_PERIODS_T2 = [
    '91 s/d 120 Days', '121 s/d 150 Days', '151 s/d 180 Days',
    '>180 Days', 'Total Overdue', 'Total',
]

_KEYS_T1  = ['curr', 'b1', 'b2', 'b3', 'b4', '90up']
_KEYS_T2  = ['b5',   'b6', 'b7', '180up', 'od', 'tot']

def _write_title_block(ws, company, subtitle, fmt_date):
    _title(ws['A1'], company)
    _title(ws['A2'], subtitle)
    _title(ws['A3'], 'CURRENCY : IDR')

def _period_header_4col(ws, row_merge, row_sub, col_start, periods, dim_col=None):
   
    c = col_start
    for period in periods:
        ws.merge_cells(
            start_row=row_merge, start_column=c,
            end_row=row_merge, end_column=c + 3
        )
        _sub(ws.cell(row=row_merge, column=c), period)
        for lbl in _LABELS_4:
            _sub(ws.cell(row=row_sub, column=c), lbl)
            c += 1

def _period_header_2col(ws, row_merge, row_sub, col_start, periods):
   
    c = col_start
    for period in periods:
        ws.merge_cells(
            start_row=row_merge, start_column=c,
            end_row=row_merge, end_column=c + 1
        )
        _sub(ws.cell(row=row_merge, column=c), period)
        for lbl in _LABELS_2:
            _sub(ws.cell(row=row_sub, column=c), lbl)
            c += 1

def _row_t1(r, keys=_KEYS_T1):
    
    vals = []
    for k in keys:
        vals += [r.get(f'acc_{k}', 0), r.get(f'amt_{k}', 0),
                 r.get(f'op_{k}', 0),  r.get(f'ar_{k}', 0)]
    return vals

def _row_t2(r, keys=_KEYS_T2):
    return _row_t1(r, keys)

def _row_outstanding(r):
    return [r.get('tot_acc', 0), r.get('acc_persen', 0),
            r.get('tot_amt', 0), r.get('amt_persen', 0),
            r.get('tot_op', 0),  r.get('op_persen', 0),
            r.get('tot_ar', 0),  r.get('ar_persen', 0)]

def _row_comp(r, keys=_KEYS_T1):

    vals = []
    for k in keys:
        vals += [r.get(f'acc_{k}_tot', 0), r.get(f'amt_{k}_tot', 0),
                 r.get(f'op_{k}_tot', 0),  r.get(f'ar_{k}_tot', 0)]
    return vals

def build_excel_branch(rows: list[dict], asof: str, company: str) -> bytes:
  
    wb  = Workbook()
    ws1 = wb.active;     ws1.title = 'Summary'
    ws2 = wb.create_sheet('Comparison')
    fmt = _fmt_date(asof)
    n   = len(rows)

    W1 = {'A':25,'B':15,'C':15,'D':25,'E':20,'F':25,'G':15,'H':25,'I':20,
          'J':15,'K':15,'L':25,'M':20,'N':15,'O':15,'P':25,'Q':20,
          'R':15,'S':15,'T':25,'U':20,'W':15,'X':25,'Y':20}
    W2 = {'A':25,'C':15,'D':25,'E':20,'G':15,'H':25,'I':20,'K':15,
          'L':25,'M':20,'O':15,'P':25,'Q':20,'S':15,'T':25,'U':20,
          'W':15,'X':25,'Y':20}
    for col, w in W1.items(): ws1.column_dimensions[col].width = w
    for col, w in W2.items(): ws2.column_dimensions[col].width = w

    _write_title_block(ws1, company,
        f'SUMMARY OF AGING DAILY LATE PAYMENT REPORT AS OF {fmt}', fmt)
    _write_title_block(ws2, company,
        f'SUMMARY OF AGING DAILY LATE PAYMENT REPORT AS OF {fmt}', fmt)

    ROW_H1A = 5;  ROW_H1B = 6;  ROW_D1 = 7

    ws1.merge_cells(f'A{ROW_H1A}:A{ROW_H1B}');  _sub(ws1[f'A{ROW_H1A}'], 'Branch')
    _period_header_4col(ws1, ROW_H1A, ROW_H1B, 2, _PERIODS_T1)

    tots_t1 = {f: 0 for k in _KEYS_T1 for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']}
    for ri, r in enumerate(rows):
        row = ROW_D1 + ri
        _data(ws1.cell(row=row, column=1), r.get('BRANCH_NAME', ''))
        for ci, v in enumerate(_row_t1(r), start=2):
            _data(ws1.cell(row=row, column=ci), v, num=True)
        for k in _KEYS_T1:
            for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']:
                tots_t1[f] += r.get(f, 0) or 0

    TOT1 = ROW_D1 + n
    _bold_data(ws1.cell(row=TOT1, column=1), 'Total')
    for ci, v in enumerate([tots_t1[f] for k in _KEYS_T1
                             for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']], start=2):
        _bold_data(ws1.cell(row=TOT1, column=ci), v, num=True)

    ROW_H2A = TOT1 + 2;  ROW_H2B = ROW_H2A + 1;  ROW_D2 = ROW_H2B + 1

    ws1.merge_cells(f'A{ROW_H2A}:A{ROW_H2B}');  _sub(ws1[f'A{ROW_H2A}'], 'Branch')
    _period_header_4col(ws1, ROW_H2A, ROW_H2B, 2, _PERIODS_T2)

    tots_t2 = {f: 0 for k in _KEYS_T2 for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']}
    for ri, r in enumerate(rows):
        row = ROW_D2 + ri
        _data(ws1.cell(row=row, column=1), r.get('BRANCH_NAME', ''))
        for ci, v in enumerate(_row_t2(r), start=2):
            _data(ws1.cell(row=row, column=ci), v, num=True)
        for k in _KEYS_T2:
            for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']:
                tots_t2[f] += r.get(f, 0) or 0

    TOT2 = ROW_D2 + n
    _bold_data(ws1.cell(row=TOT2, column=1), 'Total')
    for ci, v in enumerate([tots_t2[f] for k in _KEYS_T2
                             for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']], start=2):
        _bold_data(ws1.cell(row=TOT2, column=ci), v, num=True)

    tot_out = {k: rows[0].get(f'{k}_tot_out', 0) if rows else 0
               for k in ['acc','amt','op','ar']}

    ROW_OUT_TITLE = TOT2 + 2
    ws1[f'A{ROW_OUT_TITLE}'] = f'TOTAL OUTSTANDING AS OF {fmt}'
    ws1[f'A{ROW_OUT_TITLE}'].font = _ft(bold=True, size=10)
    ws1.merge_cells(f'A{ROW_OUT_TITLE}:G{ROW_OUT_TITLE}')

    ROW_H3A = ROW_OUT_TITLE + 2;  ROW_H3B = ROW_H3A + 1;  ROW_D3 = ROW_H3B + 1
    ws1.merge_cells(f'A{ROW_H3A}:A{ROW_H3B}');  _sub(ws1[f'A{ROW_H3A}'], 'Branch')
    ws1.merge_cells(f'B{ROW_H3A}:I{ROW_H3A}');  _sub(ws1[f'B{ROW_H3A}'], 'OUTSTANDING')
    for ci, lbl in enumerate(['Acc','%','Amount','%','Outstanding Principal','%','AR Principal','%'], start=2):
        _sub(ws1.cell(row=ROW_H3B, column=ci), lbl)

    for ri, r in enumerate(rows):
        row = ROW_D3 + ri
        _data(ws1.cell(row=row, column=1), r.get('BRANCH_NAME', ''))
        for ci, v in enumerate(_row_outstanding(r), start=2):
            c = ws1.cell(row=row, column=ci)
            _data(c, v)
            if ci in (3, 5, 7):
                c.number_format = _NUM

    TOT3 = ROW_D3 + n
    _bold_data(ws1.cell(row=TOT3, column=1), 'Total')
    for ci, v in enumerate([tot_out['acc'],100, tot_out['amt'],100,
                             tot_out['op'],100, tot_out['ar'],100], start=2):
        c = ws1.cell(row=TOT3, column=ci)
        _bold_data(c, v)
        if ci in (3, 5, 7):
            c.number_format = _NUM

    ws2['A5'] = 'COMPARISON TO TOTAL OUTSTANDING'
    ws2['A5'].font = _ft(bold=True, size=10)
    ws2.merge_cells('A5:E5')

    ROW_C1A = 7;  ROW_C1B = 8;  ROW_CD1 = 9
    ws2.merge_cells(f'A{ROW_C1A}:A{ROW_C1B}');  _sub(ws2[f'A{ROW_C1A}'], 'Branch')
    _period_header_4col(ws2, ROW_C1A, ROW_C1B, 2, _PERIODS_T1)

    for ri, r in enumerate(rows):
        row = ROW_CD1 + ri
        _data(ws2.cell(row=row, column=1), r.get('BRANCH_NAME', ''))
        for ci, v in enumerate(_row_comp(r, _KEYS_T1), start=2):
            _data(ws2.cell(row=row, column=ci), v)

    CTOT1 = ROW_CD1 + n
    _bold_data(ws2.cell(row=CTOT1, column=1), 'Total')
    for ci, k in enumerate(_KEYS_T1, start=0):
        base = 2 + ci * 4
        _bold_data(ws2.cell(row=CTOT1, column=base),   _safe_pct(tots_t1[f'acc_{k}'], tot_out['acc']))
        _bold_data(ws2.cell(row=CTOT1, column=base+1), _safe_pct(tots_t1[f'amt_{k}'], tot_out['amt']))
        _bold_data(ws2.cell(row=CTOT1, column=base+2), _safe_pct(tots_t1[f'op_{k}'],  tot_out['op']))
        _bold_data(ws2.cell(row=CTOT1, column=base+3), _safe_pct(tots_t1[f'ar_{k}'],  tot_out['ar']))

    ROW_C2A = CTOT1 + 2;  ROW_C2B = ROW_C2A + 1;  ROW_CD2 = ROW_C2B + 1
    ws2.merge_cells(f'A{ROW_C2A}:A{ROW_C2B}');  _sub(ws2[f'A{ROW_C2A}'], 'Branch')
    _period_header_4col(ws2, ROW_C2A, ROW_C2B, 2, _PERIODS_T2)

    for ri, r in enumerate(rows):
        row = ROW_CD2 + ri
        _data(ws2.cell(row=row, column=1), r.get('BRANCH_NAME', ''))
        for ci, v in enumerate(_row_comp(r, _KEYS_T2), start=2):
            _data(ws2.cell(row=row, column=ci), v)

    CTOT2 = ROW_CD2 + n
    _bold_data(ws2.cell(row=CTOT2, column=1), 'Total')
    for ci, k in enumerate(_KEYS_T2, start=0):
        base = 2 + ci * 4
        _bold_data(ws2.cell(row=CTOT2, column=base),   _safe_pct(tots_t2[f'acc_{k}'], tot_out['acc']))
        _bold_data(ws2.cell(row=CTOT2, column=base+1), _safe_pct(tots_t2[f'amt_{k}'], tot_out['amt']))
        _bold_data(ws2.cell(row=CTOT2, column=base+2), _safe_pct(tots_t2[f'op_{k}'],  tot_out['op']))
        _bold_data(ws2.cell(row=CTOT2, column=base+3), _safe_pct(tots_t2[f'ar_{k}'],  tot_out['ar']))

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

_PERIODS_CMO = [
    'Current', '01 s/d 10 Days', '11 s/d 30 Days',
    '31 s/d 60 Days', '61 s/d 90 Days', '> 90 Days',
    'Total Overdue', 'Total',
]
_KEYS_CMO = ['curr', 'b1', 'b2', 'b3', 'b4', '90up', 'od', 'tot']

_W_CMO_S1 = {
    'A':25,'B':15,'C':15,'D':20,'E':20,'F':15,'G':15,'H':20,'I':20,
    'J':15,'K':15,'L':20,'M':20,'N':15,'O':15,'P':20,'Q':20,'R':15,
    'S':15,'T':20,'U':20,'V':15,'W':15,'X':20,'Y':20,'Z':15,'AA':15,
    'AB':20,'AC':20,'AD':15,'AE':15,'AF':20,'AG':20,
}
_W_CMO_S2 = {k: 10 for k in list(_W_CMO_S1.keys())}
_W_CMO_S2['A'] = 25

def build_excel_by_person(rows: list[dict], rows_out: list[dict],
                          asof: str, company: str,
                          row_label: str, title_suffix: str) -> bytes:

    wb  = Workbook()
    ws1 = wb.active;     ws1.title = 'Summary'
    ws2 = wb.create_sheet('Comparison')
    fmt = _fmt_date(asof)

    for col, w in _W_CMO_S1.items(): ws1.column_dimensions[col].width = w
    for col, w in _W_CMO_S2.items(): ws2.column_dimensions[col].width = w

    subtitle = (f'SUMMARY OF AGING DAILY LATE PAYMENT REPORT '
                f'{title_suffix} AS OF {fmt}')
    _write_title_block(ws1, company, subtitle, fmt)
    _write_title_block(ws2, company, subtitle, fmt)

    ws1.merge_cells('A5:A6');  _sub(ws1['A5'], row_label)
    _period_header_4col(ws1, 5, 6, 2, _PERIODS_CMO)

    ws2['A5'] = 'COMPARISON TO TOTAL OUTSTANDING'
    ws2['A5'].font = _ft(bold=True, size=10)
    ws2.merge_cells('A5:E5')
    ws2.merge_cells('A7:A8');  _sub(ws2['A7'], row_label)
    _period_header_4col(ws2, 7, 8, 2, _PERIODS_CMO)

    tots = {f: 0 for k in _KEYS_CMO
            for f in [f'acc_{k}',f'amt_{k}',f'op_{k}',f'ar_{k}']}

    tot_out = {}
    ROW_D = 7
    for ri, r in enumerate(rows):
        row = ROW_D + ri
        if ri == 0:
            tot_out = {k: r.get(f'{k}_tot_out', 0) for k in ['acc','amt','op','ar']}

        _data(ws1.cell(row=row, column=1), r.get('FULLNAME', ''))
        col = 2
        for k in _KEYS_CMO:
            for fld in ['acc','amt','op','ar']:
                _data(ws1.cell(row=row, column=col), r.get(f'{fld}_{k}', 0), num=True)
                col += 1
                tots[f'{fld}_{k}'] += r.get(f'{fld}_{k}', 0) or 0

        ws2_row = 9 + ri
        _data(ws2.cell(row=ws2_row, column=1), r.get('FULLNAME', ''))
        col2 = 2
        for k in _KEYS_CMO:
            for fld in ['acc','amt','op','ar']:
                _data(ws2.cell(row=ws2_row, column=col2), r.get(f'{fld}_{k}_tot', 0))
                col2 += 1

    n = len(rows)
    TOT1 = ROW_D + n

    _bold_data(ws1.cell(row=TOT1, column=1), 'Total')
    col = 2
    for k in _KEYS_CMO:
        for fld in ['acc','amt','op','ar']:
            _bold_data(ws1.cell(row=TOT1, column=col), tots[f'{fld}_{k}'], num=True)
            col += 1

    ROW_OUT_TITLE = TOT1 + 2
    ws1[f'A{ROW_OUT_TITLE}'] = f'TOTAL OUTSTANDING AS OF {fmt}'
    ws1[f'A{ROW_OUT_TITLE}'].font = _ft(bold=True, size=10)
    ws1.merge_cells(f'A{ROW_OUT_TITLE}:I{ROW_OUT_TITLE}')

    ROW_OH1 = ROW_OUT_TITLE + 2;  ROW_OH2 = ROW_OH1 + 1;  ROW_OD = ROW_OH2 + 1
    ws1.merge_cells(f'A{ROW_OH1}:A{ROW_OH2}');  _sub(ws1[f'A{ROW_OH1}'], row_label)
    ws1.merge_cells(f'B{ROW_OH1}:I{ROW_OH1}');  _sub(ws1[f'B{ROW_OH1}'], 'OUTSTANDING')
    for ci, lbl in enumerate(['Acc','%','Amount','%','Outstanding Principal','%','AR Principal','%'], start=2):
        _sub(ws1.cell(row=ROW_OH2, column=ci), lbl)

    for ri, r in enumerate(rows_out):
        row = ROW_OD + ri
        _data(ws1.cell(row=row, column=1), r.get('FULLNAME', ''))
        for ci, v in enumerate([r.get('tot_acc',0), r.get('acc_persen',0),
                                 r.get('tot_amt',0), r.get('amt_persen',0),
                                 r.get('tot_op',0),  r.get('op_persen',0),
                                 r.get('tot_ar',0),  r.get('ar_persen',0)], start=2):
            c = ws1.cell(row=row, column=ci)
            _data(c, v)
            if ci in (3, 5, 7):
                c.number_format = _NUM

    OTOT = ROW_OD + len(rows_out)
    _bold_data(ws1.cell(row=OTOT, column=1), 'Total')
    for ci, v in enumerate([tot_out.get('acc',0),100,
                             tot_out.get('amt',0),100,
                             tot_out.get('op',0),100,
                             tot_out.get('ar',0),100], start=2):
        c = ws1.cell(row=OTOT, column=ci)
        _bold_data(c, v)
        if ci in (3, 5, 7):
            c.number_format = _NUM

    ws2_tot = 9 + n
    _bold_data(ws2.cell(row=ws2_tot, column=1), 'Total')
    col2 = 2
    for k in _KEYS_CMO:
        _bold_data(ws2.cell(row=ws2_tot, column=col2),   _safe_pct(tots[f'acc_{k}'], tot_out.get('acc')))
        _bold_data(ws2.cell(row=ws2_tot, column=col2+1), _safe_pct(tots[f'amt_{k}'], tot_out.get('amt')))
        _bold_data(ws2.cell(row=ws2_tot, column=col2+2), _safe_pct(tots[f'op_{k}'],  tot_out.get('op')))
        _bold_data(ws2.cell(row=ws2_tot, column=col2+3), _safe_pct(tots[f'ar_{k}'],  tot_out.get('ar')))
        col2 += 4

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

_PERIODS_OTHER = [
    'Current', '01 s/d 10 Days', '11 s/d 30 Days',
    '31 s/d 60 Days', '61 s/d 90 Days', '> 90 Days',
    'Total Overdue', 'Total',
]
_KEYS_OTHER = ['curr', 'b1', 'b2', 'b3', 'b4', '90up', 'od', 'tot']

_W_OTHER = {
    'A':25,'B':15,'C':15,'D':20,'E':20,'F':15,'G':15,'H':20,'I':20,
    'J':15,'K':15,'L':20,'M':20,'N':15,'O':15,'P':20,'Q':20,
}

def _write_other_section(ws, row_title_merge, row_data_start,
                         section_label, rows, label_key):
  
    R1 = row_title_merge;  R2 = R1 + 1

    ws.merge_cells(start_row=R1, start_column=1, end_row=R2, end_column=1)
    _sub(ws.cell(row=R1, column=1), section_label)

    col = 2
    for period in _PERIODS_OTHER:
        ws.merge_cells(start_row=R1, start_column=col, end_row=R1, end_column=col+1)
        _sub(ws.cell(row=R1, column=col), period)
        _sub(ws.cell(row=R2, column=col),   'Acc')
        _sub(ws.cell(row=R2, column=col+1), '%')
        col += 2

    for ri, r in enumerate(rows):
        row = row_data_start + ri
        _data(ws.cell(row=row, column=1), r.get(label_key, ''))
        col = 2
        for k in _KEYS_OTHER:
            _data(ws.cell(row=row, column=col),   r.get(f'acc_{k}', 0), num=True)
            _data(ws.cell(row=row, column=col+1), r.get(f'acc_{k}_tot', 0))
            col += 2

    end_row = row_data_start + len(rows) - 1
    if end_row >= row_data_start:
        for r in range(row_data_start, end_row + 1):
            for c in range(1, 18):
                ws.cell(row=r, column=c).border = _BORDER

    return row_data_start + len(rows)

def build_excel_other(rows_karoseri, rows_tenor, rows_ltv,
                      asof: str, company: str) -> bytes:

    wb  = Workbook()
    ws1 = wb.active;  ws1.title = 'Summary by Other'
    fmt = _fmt_date(asof)

    for col, w in _W_OTHER.items(): ws1.column_dimensions[col].width = w

    ws1['A1'] = company;     ws1['A1'].font = _ft(bold=True, size=12); ws1['A1'].alignment = _al('left')
    ws1['A2'] = f'SUMMARY OF AGING DAILY LATE PAYMENT REPORT AS OF {fmt}'
    ws1['A2'].font = _ft(bold=True, size=12); ws1['A2'].alignment = _al('left')
    ws1['A3'] = 'CURRENCY : IDR'
    ws1['A3'].font = _ft(bold=True, size=12); ws1['A3'].alignment = _al('left')

    nk = _write_other_section(ws1, 5, 7, 'Karoseri', rows_karoseri, 'Karoseri')

    nt_start = nk + 2
    nt = _write_other_section(ws1, nt_start, nt_start + 2, 'Tenor', rows_tenor, 'tenor')

    nl_start = nt + 2
    _write_other_section(ws1, nl_start, nl_start + 2, 'LTV', rows_ltv, 'LTV')

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

def build_excel_dealer(rows_tot_acc, rows_tot_od, rows_tot_od90,
                       asof: str, company: str) -> bytes:
    
    wb  = Workbook()
    ws1 = wb.active;  ws1.title = 'Summary by Dealer'
    fmt = _fmt_date(asof)

    for col, w in _W_OTHER.items(): ws1.column_dimensions[col].width = w

    ws1['A1'] = company;     ws1['A1'].font = _ft(bold=True, size=12); ws1['A1'].alignment = _al('left')
    ws1['A2'] = f'SUMMARY OF AGING DAILY LATE PAYMENT REPORT (BY DEALER) AS OF {fmt}'
    ws1['A2'].font = _ft(bold=True, size=12); ws1['A2'].alignment = _al('left')
    ws1['A3'] = 'CURRENCY : IDR'
    ws1['A3'].font = _ft(bold=True, size=12); ws1['A3'].alignment = _al('left')

    n1 = _write_other_section(ws1, 5, 7, 'Dealer (By Total Account)', rows_tot_acc, 'Dealer')

    n2_start = max(n1 + 2, 29)
    n2 = _write_other_section(ws1, n2_start, n2_start + 2,
                               'Dealer (By Total Overdue)', rows_tot_od, 'Dealer')

    n3_start = max(n2 + 2, 53)
    _write_other_section(ws1, n3_start, n3_start + 2,
                          'Dealer (By Total Overdue > 90 Days)', rows_tot_od90, 'Dealer')

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()