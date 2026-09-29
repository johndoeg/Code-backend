from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

approved_cam_by_dealer_bp = Blueprint('approved_cam_by_dealer', __name__, url_prefix='/CAM/Approved')

PAGE_SIZE_DEFAULT = 20
DEALER_PAGE_SIZE_DEFAULT = 20

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd or branch_cd == '000':
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

@approved_cam_by_dealer_bp.route('/cam-by-dealer/branches')
@login_required
def list_branches():
    own_branch_cd = session.get('branchcode')
    try:
        if own_branch_cd in ('000', '100'):
            rows = _exec_rows(
                "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
                [],
            )
        else:
            rows = _exec_rows(
                "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?",
                [own_branch_cd],
            )
        return jsonify({
            'branches': [{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows],
            'restricted': own_branch_cd not in ('000', '100'),
            'default_branch_cd': '000' if own_branch_cd in ('000', '100') else own_branch_cd,
        })
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@approved_cam_by_dealer_bp.route('/cam-by-dealer/dealers')
@login_required
def list_dealers():
    search = request.args.get('search', '').strip()
    branch_cd = request.args.get('branch_cd', '000').strip()
    page = int(request.args.get('page', 1) or 1)
    page_size = int(request.args.get('page_size', DEALER_PAGE_SIZE_DEFAULT) or DEALER_PAGE_SIZE_DEFAULT)

    where = ["Supp_Type <> 'KR'", "NICKNAME LIKE ?"]
    params: list = [f"%{search}%"]
    if branch_cd != '000':
        where.append("BRANCH_CD = ?")
        params.append(branch_cd)

    where_sql = ' AND '.join(where)
    try:
        count_rows = _exec_rows(f"SELECT COUNT(*) AS cnt FROM CLFSUPP WITH (NOLOCK) WHERE {where_sql}", params)
        total_count = count_rows[0]['cnt'] if count_rows else 0

        offset = max(page - 1, 0) * page_size
        rows = _exec_rows(
            f"""
            SELECT SUPP, NAME, NICKNAME, ADDRESS, PHONE, CONTACT
            FROM CLFSUPP WITH (NOLOCK)
            WHERE {where_sql}
            ORDER BY NAME
            OFFSET ? ROWS FETCH NEXT ? ROWS ONLY
            """,
            params + [offset, page_size],
        )
        return jsonify({
            'data': [
                {
                    'supp': r.get('SUPP'), 'name': r.get('NAME'), 'nickname': r.get('NICKNAME'),
                    'address': r.get('ADDRESS'), 'phone': r.get('PHONE'), 'contact': r.get('CONTACT'),
                }
                for r in rows
            ],
            'total_count': total_count,
            'page': page,
            'page_size': page_size,
        })
    except Exception as e:
        current_app.logger.exception("Dealer list error")
        return jsonify({'error': str(e)}), 500

def _build_filters(start_date: Optional[date], end_date: Optional[date], branch_cd_search: str, dealer_nickname: str):
    where = ["ccont.EXECUTION IS NOT NULL"]
    params: list = []

    if start_date and end_date:
        where.append("ccont.EXECUTION >= ? AND ccont.EXECUTION <= ?")
        params += [start_date, end_date]

    if branch_cd_search and branch_cd_search != '000':
        where.append("ccont.BRANCH_CD = ?")
        params.append(branch_cd_search)

    if dealer_nickname:
        where.append("cls.NICKNAME LIKE ?")
        params.append(f"%{dealer_nickname}%")

    return ' AND '.join(where), params

_BASE_CTE = """
    WITH Base AS (
        SELECT
            ccont.APPL_NO, ccont.Cons_Leas, ccont.LEASE_NO,
            clfl.LESSEE_NM, cls.NICKNAME, ub.BRANCH_NAME, ccont.EXECUTION,
            au.FULLNAME, clfq.SUPP, v.Net_Finance AS NETFINANCE, clfl.APLESS, ccont.TENOR,
            LTRIM(RTRIM(clfq.CONDITION)) AS CONDITION, ca.HAS_GUARANTOR, clfl.LESSEE_TP,
            mr.REVISI,
            ROW_NUMBER() OVER (PARTITION BY ccont.APPL_NO ORDER BY ccont.EXECUTION ASC) AS dedup_rn
        FROM CLFCONT ccont WITH (NOLOCK)
            LEFT JOIN CLFLESS clfl WITH (NOLOCK) ON ccont.LESSEE_NO = clfl.LESSEE_NO
            LEFT JOIN CLFEQUIP clfq WITH (NOLOCK) ON clfq.LEASE_NO = ccont.LEASE_NO
            LEFT JOIN CLFSUPP cls WITH (NOLOCK) ON cls.SUPP = clfq.SUPP
            LEFT JOIN UMFBRANCH ub WITH (NOLOCK) ON ccont.BRANCH_CD = ub.BRANCH_CD
            LEFT JOIN APP_USRMST au WITH (NOLOCK) ON ccont.CMO = au.EMPLOYEE_ID
            LEFT JOIN CAFINSU cfu WITH (NOLOCK) ON ccont.APPL_NO = cfu.APPL_NO
            LEFT JOIN CAFCONT ca WITH (NOLOCK) ON ccont.APPL_NO = ca.APPL_NO
            LEFT JOIN V_NetFinance v WITH (NOLOCK) ON ccont.LEASE_NO = v.LEASE_NO
            LEFT JOIN MCMORevisi mr WITH (NOLOCK) ON mr.NewAppl_No = ccont.APPL_NO
        WHERE {where_sql}
    ),
    Deduped AS (
        SELECT *, ROW_NUMBER() OVER (ORDER BY EXECUTION ASC) AS ro_number
        FROM Base WHERE dedup_rn = 1
    )
"""

def _fetch_detail_rows(where_sql: str, params: list, page: Optional[int], page_size: int) -> tuple[list[dict], int]:
    count_sql = _BASE_CTE.format(where_sql=where_sql) + "SELECT COUNT(*) AS cnt FROM Deduped"
    count_rows = _exec_rows(count_sql, params)
    total_count = count_rows[0]['cnt'] if count_rows else 0

    if page is None:
        list_sql = _BASE_CTE.format(where_sql=where_sql) + "SELECT * FROM Deduped ORDER BY ro_number"
        rows = _exec_rows(list_sql, params)
    else:
        offset = max(page - 1, 0) * page_size
        list_sql = (
            _BASE_CTE.format(where_sql=where_sql)
            + "SELECT * FROM Deduped ORDER BY ro_number OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
        )
        rows = _exec_rows(list_sql, params + [offset, page_size])

    return rows, total_count

def _fetch_summary(where_sql: str, params: list) -> dict:
    sql = _BASE_CTE.format(where_sql=where_sql) + """
        SELECT
            SUM(CASE WHEN REVISI IS NULL THEN 1 ELSE 0 END) AS jlh_cam,
            SUM(CASE WHEN REVISI IS NULL THEN NETFINANCE ELSE 0 END) AS tot_cam,
            SUM(CASE WHEN REVISI IS NOT NULL THEN 1 ELSE 0 END) AS jlh_revisi,
            SUM(CASE WHEN REVISI IS NOT NULL THEN NETFINANCE ELSE 0 END) AS tot_revisi,
            COUNT(*) AS grand_count,
            SUM(NETFINANCE) AS grand_net
        FROM Deduped
    """
    rows = _exec_rows(sql, params)
    r = rows[0] if rows else {}
    return {
        'new_cam_count': int(_num(r.get('jlh_cam'))),
        'new_cam_net_finance': _num(r.get('tot_cam')),
        'revised_count': int(_num(r.get('jlh_revisi'))),
        'revised_net_finance': _num(r.get('tot_revisi')),
        'grand_count': int(_num(r.get('grand_count'))),
        'grand_net_finance': _num(r.get('grand_net')),
    }

def _row_to_detail(row: dict) -> dict:
    execution = row.get('EXECUTION')
    revisi = row.get('REVISI')
    revision_suffix = f"R{revisi}" if revisi is not None else ''

    has_guarantor = row.get('HAS_GUARANTOR')
    guarantor = 2 if has_guarantor == '1' else 1
    ind_cor = 2 if row.get('LESSEE_TP') == 'PT' else 1
    c2c = 1 if row.get('SUPP') == 999999 else 0
    condition = (row.get('CONDITION') or '').strip()
    new_car = 2 if condition == 'NEW' else 1

    return {
        'appl_no': row.get('APPL_NO'),
        'cons_leas': row.get('Cons_Leas'),
        'revision_suffix': revision_suffix,
        'customer_name': row.get('LESSEE_NM') or '-',
        'dealer_nickname': row.get('NICKNAME') or '-',
        'branch_name': row.get('BRANCH_NAME') or '-',
        'disbursement_date': execution.strftime('%d %b %Y') if execution else '',
        'cmo_name': row.get('FULLNAME') or '-',
        'edit_link_params': {
            'id1': row.get('Cons_Leas'), 'id2': ind_cor, 'id3': 1, 'id4': guarantor,
            'id5': new_car, 'id6': 'Edit', 'id7': row.get('TENOR'), 'id8': c2c,
            'id9': '', 'apless': row.get('APLESS'), 'applno': row.get('APPL_NO'),
        },
    }

def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

@approved_cam_by_dealer_bp.route('/cam-by-dealer/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    branch_cd_search = request.args.get('branch_cd_search', '000').strip()
    dealer_nickname = request.args.get('dealer_nickname', '').strip()
    page = int(request.args.get('page', 1) or 1)
    page_size = int(request.args.get('page_size', PAGE_SIZE_DEFAULT) or PAGE_SIZE_DEFAULT)

    where_sql, params = _build_filters(start_date, end_date, branch_cd_search, dealer_nickname)
    try:
        rows, total_count = _fetch_detail_rows(where_sql, params, page, page_size)
        summary = _fetch_summary(where_sql, params)
        return jsonify({
            'data': [_row_to_detail(r) for r in rows],
            'total_count': total_count,
            'page': page,
            'page_size': page_size,
            'summary': summary,
            'show_branch_column': branch_cd_search == '000',
            'show_dealer_column': dealer_nickname == '',
        })
    except Exception as e:
        current_app.logger.exception("CAM by Dealer list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=16)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

def _build_workbook(start_date, end_date, branch_cd_search, dealer_nickname, rows, summary, show_branch, show_dealer) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'CamByDealer'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List Approved by Dealer per Disbursement Date'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = _period_label(start_date, end_date)
    ws['A3'].alignment = _CENTER

    ws.merge_cells('A4:D4')
    ws['A4'] = f"Branch : {_branch_name(branch_cd_search)}"

    ws.merge_cells('A5:D5')
    ws['A5'] = f"Dealer : {dealer_nickname or 'Consolidate'}"

    headers = ['No.', 'CAM No.', 'Customer Name']
    if show_dealer:
        headers.append('Dealer Nickname')
    if show_branch:
        headers.append('Branch')
    headers += ['Disbursement Date', 'CMO']

    header_row = 7
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.border = _ALL_BORDERS
        cell.alignment = _CENTER

    excel_row = header_row + 1
    if not rows:
        ws.merge_cells(f'A{excel_row}:{chr(64 + len(headers))}{excel_row}')
        cell = ws[f'A{excel_row}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        excel_row += 1
    else:
        for i, r in enumerate(rows, start=1):
            d = _row_to_detail(r)
            values = [i, f"{d['cons_leas']} - {d['appl_no']}{d['revision_suffix']}", d['customer_name']]
            if show_dealer:
                values.append(d['dealer_nickname'])
            if show_branch:
                values.append(d['branch_name'])
            values += [d['disbursement_date'], d['cmo_name']]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = _ALL_BORDERS
                if col_idx == 1:
                    cell.alignment = _CENTER
            excel_row += 1

    excel_row += 1
    labels_values = [
        ('New CAM', summary['new_cam_count'], 'Net Finance', summary['new_cam_net_finance']),
        ('Revisi', summary['revised_count'], 'Net Finance', summary['revised_net_finance']),
        ('Total', summary['grand_count'], 'Total Net Finance', summary['grand_net_finance']),
    ]
    for label1, val1, label2, val2 in labels_values:
        ws.cell(row=excel_row, column=1, value=label1)
        ws.cell(row=excel_row, column=2, value=f"= {val1:,.0f}")
        ws.cell(row=excel_row, column=3, value=label2)
        ws.cell(row=excel_row, column=4, value=f"= {val2:,.0f}")
        excel_row += 1

    ws.column_dimensions['A'].width = 6
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 20

    return wb

@approved_cam_by_dealer_bp.route('/cam-by-dealer/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    branch_cd_search = request.args.get('branch_cd_search', '000').strip()
    dealer_nickname = request.args.get('dealer_nickname', '').strip()

    where_sql, params = _build_filters(start_date, end_date, branch_cd_search, dealer_nickname)
    try:
        rows, _total = _fetch_detail_rows(where_sql, params, None, PAGE_SIZE_DEFAULT)
        summary = _fetch_summary(where_sql, params)
        wb = _build_workbook(
            start_date, end_date, branch_cd_search, dealer_nickname, rows, summary,
            branch_cd_search == '000', dealer_nickname == '',
        )
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='CamByDealer.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("CAM by Dealer export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Credit Approval Memorandum</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();window.close();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
        </tr>
        <tr>
            <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List Approved by Dealer per Disbursement Date</strong></td>
        </tr>
        <tr><td colspan="4" align="center">{period}</td></tr>
        <tr><td colspan="4">Branch : {branch_name}</td></tr>
        <tr><td colspan="4" nowrap="nowrap">Dealer : {dealer_label}</td></tr>
        <tr>
            <td colspan="4">
                <table border="1" cellpadding="2" cellspacing="0" width="100%">
                    <tr align="center">{header_html}</tr>
                    {rows_html}
                </table>
                <br>
                <table cellpadding="2" cellspacing="0" width="100%">
                    <tr>
                        <td width="17%">New CAM</td>
                        <td>= {new_cam_count}</td>
                        <td width="17%">Net Finance</td>
                        <td>= {new_cam_fin}</td>
                    </tr>
                    <tr>
                        <td width="17%">Revisi</td>
                        <td>= {revised_count}</td>
                        <td width="17%">Net Finance</td>
                        <td>= {revised_fin}</td>
                    </tr>
                    <tr>
                        <td width="17%">Total</td>
                        <td>= {grand_count}</td>
                        <td width="17%">Total Net Finance</td>
                        <td>= {grand_fin}</td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""

def _render_header_html(show_branch: bool, show_dealer: bool) -> str:
    parts = [
        '<td width="3%" style="border:thin solid #000">No.</td>',
        '<td width="14%" style="border:thin solid #000">CAM No.</td>',
        '<td style="border:thin solid #000">Customer Name</td>',
    ]
    if show_dealer:
        parts.append('<td width="17%" style="border:thin solid #000">Dealer Nickname</td>')
    if show_branch:
        parts.append('<td width="12%" style="border:thin solid #000">Branch</td>')
    parts += [
        '<td width="11%" style="border:thin solid #000">Disbursement Date</td>',
        '<td width="9%" style="border:thin solid #000">CMO</td>',
    ]
    return ''.join(parts)

def _render_detail_rows_html(rows: list[dict], show_branch: bool, show_dealer: bool) -> str:
    col_count = 5 + int(show_branch) + int(show_dealer)
    if not rows:
        return f'<tr><td colspan="{col_count}" align="center"><font color="red"><b>Data Not Found</b></font></td></tr>'
    parts = []
    for i, r in enumerate(rows, start=1):
        d = _row_to_detail(r)
        cells = [
            f'<td align="center">{i}</td>',
            f"<td>{d['cons_leas']} - {d['appl_no']} {d['revision_suffix']}</td>",
            f"<td>{d['customer_name']}</td>",
        ]
        if show_dealer:
            cells.append(f"<td>{d['dealer_nickname']}</td>")
        if show_branch:
            cells.append(f"<td>{d['branch_name']}</td>")
        cells += [f"<td>{d['disbursement_date']}</td>", f"<td>{d['cmo_name']}</td>"]
        parts.append(f"<tr>{''.join(cells)}</tr>")
    return ''.join(parts)

@approved_cam_by_dealer_bp.route('/cam-by-dealer/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    branch_cd_search = request.args.get('branch_cd_search', '000').strip()
    dealer_nickname = request.args.get('dealer_nickname', '').strip()

    where_sql, params = _build_filters(start_date, end_date, branch_cd_search, dealer_nickname)
    try:
        rows, _total = _fetch_detail_rows(where_sql, params, None, PAGE_SIZE_DEFAULT)
        summary = _fetch_summary(where_sql, params)
        show_branch = branch_cd_search == '000'
        show_dealer = dealer_nickname == ''

        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            period=_period_label(start_date, end_date),
            branch_name=_branch_name(branch_cd_search),
            dealer_label=dealer_nickname or 'Consolidate',
            header_html=_render_header_html(show_branch, show_dealer),
            rows_html=_render_detail_rows_html(rows, show_branch, show_dealer),
            new_cam_count=summary['new_cam_count'],
            new_cam_fin=f"{summary['new_cam_net_finance']:,.0f}",
            revised_count=summary['revised_count'],
            revised_fin=f"{summary['revised_net_finance']:,.0f}",
            grand_count=summary['grand_count'],
            grand_fin=f"{summary['grand_net_finance']:,.0f}",
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("CAM by Dealer print error")
        return jsonify({'error': str(e)}), 500