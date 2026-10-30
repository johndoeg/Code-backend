from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

cam_by_approver_bp = Blueprint('cam_by_approver', __name__, url_prefix='/CAM/Approved')

_APPROVER_ROLES = [
    ('CRH', 'CREDIT REVIEW UNIT HEAD'),
    ('CRDH', 'CREDIT REVIEW DIVISION HEAD'),
    ('DR', 'BUSINESS OPERATION DIRECTOR'),
    ('PD', 'PRESIDENT DIRECTOR'),
]

_COMMITTEE_ROLE_CODES = ('CRH', 'CRDH', 'DR', 'PD')

_STATUS_LABELS = {
    'A': 'Approved',
    'F': 'Approved & Send to Admin',
    'R': 'Reject',
    'N': 'Send for Approval',
}

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _exec_multi_resultset(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    all_rows: list[dict] = []
    try:
        cursor.execute(sql, params)
        while True:
            if cursor.description is not None:
                columns = [col[0] for col in cursor.description]
                all_rows += [dict(zip(columns, row)) for row in cursor.fetchall()]
            if not cursor.nextset():
                break
        return all_rows
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _resolve_branch(requested_branch_cd: str) -> dict:
    own_branch_cd = session.get('branchcode')
    if own_branch_cd == '100':
        return {'disabled': False, 'search_branch': requested_branch_cd}
    return {'disabled': True, 'search_branch': own_branch_cd}

@cam_by_approver_bp.route('/cam-by-approver/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@cam_by_approver_bp.route('/cam-by-approver/context')
@login_required
def get_context():
    resolved = _resolve_branch(request.args.get('branch_cd', '').strip())
    return jsonify({'disabled': resolved['disabled'], 'forced_branch_cd': resolved['search_branch']})

def _exec_sp_approver(start_date: Optional[date], end_date: Optional[date], branch_cd: str, mode: str) -> list[dict]:
    return _exec_rows(
        "EXEC USP_CAM_By_Approver ?, ?, ?, ?",
        [start_date.isoformat() if start_date else '', end_date.isoformat() if end_date else '', branch_cd, mode],
    )

def _edit_link_params(r: dict) -> dict:
    cons_leas = r.get('Cons_Leas')
    fin_type = 1 if cons_leas == 'C' else 2 if cons_leas == 'L' else None
    guarantor = 2 if r.get('HAS_GUARANTOR') == '1' else 1
    ind_cor = 2 if r.get('LESSEE_TP') == 'PT' else 1
    c2c = 1 if str(r.get('SUPP')) == '999999' else 0
    condition = (r.get('CONDITION') or '').strip()
    new_car = 2 if condition == 'NEW' else 1

    return {
        'id1': fin_type, 'id2': ind_cor, 'id3': 1, 'id4': guarantor, 'id5': new_car,
        'id6': 'Edit', 'id7': r.get('TENOR'), 'id8': c2c,
        'apless': r.get('APLESS'), 'applno': r.get('APPL_NO'),
    }

def _fetch_link_extra(appl_nos: list[str]) -> dict[str, dict]:
    result: dict[str, dict] = {}
    unique = list(dict.fromkeys(appl_nos))
    for i in range(0, len(unique), 1000):
        chunk = unique[i:i + 1000]
        placeholders = ','.join('?' for _ in chunk)
        rows = _exec_rows(
            f"""
            SELECT a.APPL_NO, a.APLESS, a.TENOR, a.HAS_GUARANTOR, b.LESSEE_TP, g.SUPP,
                   LTRIM(RTRIM(g.[CONDITION])) AS CONDITION
            FROM CAFCONT a WITH (NOLOCK)
                LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
                LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
            WHERE a.APPL_NO IN ({placeholders})
            """,
            chunk,
        )
        for row in rows:
            result.setdefault(row['APPL_NO'], row)
    return result

def _row_to_detail(r: dict, no: int, link_extra: Optional[dict] = None) -> dict:
    revisi = r.get('Revisi')
    revision_suffix = f"R{revisi}" if revisi else ''
    appr_date = r.get('APPR_DATE')
    detail = {
        'no': no,
        'branch_name': r.get('BRANCH_NAME') or '-',
        'customer_name': r.get('LESSEE_NM') or '-',
        'cam_no': f"{r.get('Cons_Leas')} - {r.get('APPL_NO')}{revision_suffix}",
        'appl_no': r.get('APPL_NO'),
        'approved_date': appr_date.strftime('%d-%m-%Y') if appr_date else '',
        'model': r.get('model'),
        'brand_type': r.get('tipe_mobil'),
        'car_status': r.get('CONDITION'),
        'contract_status': r.get('ContractStatus'),
        'outstanding_old_cam': _num(r.get('Out_OldCAM')),
        'net_finance_new_cam': _num(r.get('Out_NewCAM')),
        'total_outstanding': _num(r.get('TotOut')),
        'cmo': r.get('CMO') or '-',
    }
    if link_extra is not None:
        detail['edit_link_params'] = _edit_link_params({**r, **link_extra})
    return detail

def _fetch_report(start_date: Optional[date], end_date: Optional[date], branch_cd: str,
                  with_links: bool = False) -> dict:
    raw_sections = [
        (role_code, role_label, _exec_sp_approver(start_date, end_date, branch_cd, role_code))
        for role_code, role_label in _APPROVER_ROLES
    ]

    link_extra: dict[str, dict] = {}
    if with_links:
        appl_nos = [r['APPL_NO'] for _, _, rows in raw_sections for r in rows if r.get('APPL_NO')]
        link_extra = _fetch_link_extra(appl_nos)

    sections = []
    for role_code, role_label, rows in raw_sections:
        sections.append({
            'role_code': role_code,
            'role_label': role_label,
            'rows': [
                _row_to_detail(r, i, link_extra.get(r.get('APPL_NO'), {}) if with_links else None)
                for i, r in enumerate(rows, start=1)
            ],
        })

    total_rows = _exec_sp_approver(start_date, end_date, branch_cd, 'TOTAL')
    total = total_rows[0] if total_rows else {}
    return {
        'sections': sections,
        'grand_total': {
            'net_finance': _num(total.get('tot_ntf')),
            'units': int(_num(total.get('tot_unit'))),
        },
    }

@cam_by_approver_bp.route('/cam-by-approver/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd', '').strip()

    scope = _resolve_branch(requested_branch)
    try:
        result = _fetch_report(start_date, end_date, scope['search_branch'], with_links=True)
        return jsonify({**result, 'disabled': scope['disabled'], 'search_branch': scope['search_branch']})
    except Exception as e:
        current_app.logger.exception("CAM by Approver list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=16)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

_COLUMNS = [
    'No.', 'Branch', 'Customer Name', 'CAM No.', 'Approved Date', 'Model', 'Brand/Type',
    'Car Status', 'Contract Status', 'Outstanding', 'Net Finance', 'Total Outstanding', 'CMO',
]

def _build_workbook(start_date, end_date, branch_cd, result: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'CamByApprover'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d/%m/%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List CAM by Approver'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {_branch_name(branch_cd)}"

    ws.merge_cells('A4:D4')
    ws['A4'] = _period_label(start_date, end_date)

    row_num = 6
    for section in result['sections']:
        ws.merge_cells(f'A{row_num}:M{row_num}')
        ws[f'A{row_num}'] = f"Approved by : {section['role_label']}"
        ws[f'A{row_num}'].border = _ALL_BORDERS
        row_num += 1

        for col_idx, label in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=label)
            cell.border = _ALL_BORDERS
            cell.alignment = _CENTER
            cell.font = _HEADER_FONT
        row_num += 1

        if not section['rows']:
            ws.merge_cells(f'A{row_num}:M{row_num}')
            cell = ws[f'A{row_num}']
            cell.value = 'Data Not Found'
            cell.font = Font(color='FFFF0000', bold=True)
            cell.alignment = _CENTER
            row_num += 1
        else:
            for d in section['rows']:
                values = [
                    d['no'], d['branch_name'], d['customer_name'], d['cam_no'], d['approved_date'],
                    d['model'], d['brand_type'], d['car_status'], d['contract_status'],
                    d['outstanding_old_cam'], d['net_finance_new_cam'], d['total_outstanding'], d['cmo'],
                ]
                for col_idx, val in enumerate(values, start=1):
                    cell = ws.cell(row=row_num, column=col_idx, value=val)
                    cell.border = _ALL_BORDERS
                    if col_idx in (10, 11, 12):
                        cell.alignment = _RIGHT
                row_num += 1

    row_num += 1
    ws.merge_cells(f'A{row_num}:B{row_num}')
    ws[f'A{row_num}'] = f"Grand Total* ({_period_label(start_date, end_date).replace('Period : ', '')})"
    ws[f'A{row_num}'].font = _HEADER_FONT
    row_num += 1
    ws.cell(row=row_num, column=1, value='Net Finance')
    ws.cell(row=row_num, column=2, value=result['grand_total']['net_finance']).alignment = _RIGHT
    row_num += 1
    ws.cell(row=row_num, column=1, value='Units')
    ws.cell(row=row_num, column=2, value=result['grand_total']['units']).alignment = _RIGHT
    row_num += 2
    ws.merge_cells(f'A{row_num}:D{row_num}')
    ws[f'A{row_num}'] = '*Grand Total Amount and Unit Exclude CAM Revision'
    ws[f'A{row_num}'].font = Font(italic=True)

    ws.column_dimensions['A'].width = 6
    for col in 'BCDEFGHIJKLM':
        ws.column_dimensions[col].width = 16

    return wb

@cam_by_approver_bp.route('/cam-by-approver/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd', '').strip()

    scope = _resolve_branch(requested_branch)
    try:
        result = _fetch_report(start_date, end_date, scope['search_branch'])
        wb = _build_workbook(start_date, end_date, scope['search_branch'], result)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ListCamByApprover.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("CAM by Approver export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{company} - Credit Approval Memorandum</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%">Date : {today}</td>
        </tr>
        <tr>
            <td></td>
            <td align="center" style="font-size:20px; letter-spacing:0.2em;"><strong>List CAM by Approver</strong></td>
            <td></td>
            <td nowrap="nowrap">Print By : {printed_by}</td>
        </tr>
        <tr><td>Branch : {branch_name}</td></tr>
        <tr><td nowrap="nowrap">{period}</td></tr>
        <tr>
            <td colspan="4">
                <table cellpadding="2" cellspacing="0" width="100%">
                    <tr align="center">
                        <td style="border:thin solid #000">No.</td><td style="border:thin solid #000">Branch</td>
                        <td style="border:thin solid #000">Customer Name</td><td style="border:thin solid #000">CAM No.</td>
                        <td style="border:thin solid #000">Approved Date</td><td style="border:thin solid #000">Model</td>
                        <td style="border:thin solid #000">Brand/Type</td><td style="border:thin solid #000">Car Status</td>
                        <td style="border:thin solid #000">Contract Status</td><td style="border:thin solid #000">Outstanding</td>
                        <td style="border:thin solid #000">Net Finance</td><td style="border:thin solid #000">Total Outstanding</td>
                        <td style="border:thin solid #000">CMO</td>
                    </tr>
                {rows_html}
                </table>
            </td>
        </tr>
        <tr><td colspan="2" height="30px"></td></tr>
        <tr>
            <td colspan="2">
                <table cellpadding="2" cellspacing="0" width="100%">
                        <tr><td width="40%" colspan="2"><u><strong>Grand Total* ({period_plain}):</strong></u></td></tr>
                        <tr><td width="10%">Net Finance</td><td>: {grand_net_finance}</td></tr>
                        <tr><td width="10%">Units</td><td> : {grand_units}</td></tr>
                        <tr><td colspan="2">&nbsp;</td></tr>
                        <tr><td colspan="2"><i>*Grand Total Amount and Unit Exclude CAM Revision</i></td></tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>"""

def _render_section_rows_html(sections: list[dict]) -> str:
    parts = []
    for section in sections:
        parts.append(f'<tr><td colspan="13" style="border:thin solid #000">Approved by : {section["role_label"]}</td></tr>')
        if not section['rows']:
            parts.append('<tr><td style="background-color:#E5E5E5; border:thin solid #000" colspan="13" align="center"><font color="red"><b>Data Not Found</b></font></td></tr>')
            continue
        for i, d in enumerate(section['rows']):
            warna = '#E5E5E5' if i % 2 == 0 else '#F5F5F5'
            parts.append(f"""
                <tr style="background-color:{warna}">
                    <td style="border:thin solid #000" align="center">{d['no']}</td>
                    <td style="border:thin solid #000">{d['branch_name']}</td>
                    <td style="border:thin solid #000">{d['customer_name']}</td>
                    <td style="border:thin solid #000">{d['cam_no']}</td>
                    <td style="border:thin solid #000">{d['approved_date']}</td>
                    <td style="border:thin solid #000">{d['model']}</td>
                    <td style="border:thin solid #000">{d['brand_type']}</td>
                    <td style="border:thin solid #000">{d['car_status']}</td>
                    <td style="border:thin solid #000">{d['contract_status']}</td>
                    <td style="border:thin solid #000" align="right">{d['outstanding_old_cam']:,.0f}</td>
                    <td style="border:thin solid #000" align="right">{d['net_finance_new_cam']:,.0f}</td>
                    <td style="border:thin solid #000" align="right">{d['total_outstanding']:,.0f}</td>
                    <td style="border:thin solid #000">{d['cmo']}</td>
                </tr>""")
    return ''.join(parts)

@cam_by_approver_bp.route('/cam-by-approver/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd', '').strip()

    scope = _resolve_branch(requested_branch)
    try:
        result = _fetch_report(start_date, end_date, scope['search_branch'])
        period = _period_label(start_date, end_date)
        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_branch_name(scope['search_branch']),
            period=period,
            period_plain=period.replace('Period : ', ''),
            rows_html=_render_section_rows_html(result['sections']),
            grand_net_finance=f"{result['grand_total']['net_finance']:,.0f}",
            grand_units=f"{result['grand_total']['units']:,}",
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("CAM by Approver print error")
        return jsonify({'error': str(e)}), 500

def _fetch_cam_basic_info(appl_no: str) -> dict:
    rows = _exec_rows(
        """
        SELECT LTRIM(RTRIM(a.CMO)) AS CMO, b.LESSEE_NM, b.apless, c.USERNAME
        FROM CAFCONT a WITH (NOLOCK)
            LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
            LEFT JOIN APP_USRMST c WITH (NOLOCK) ON a.CMO = c.EMPLOYEE_ID
        WHERE a.APPL_NO = ?
        """,
        [appl_no],
    )
    return rows[0] if rows else {}

def _fetch_revision_suffix(appl_no: str) -> str:
    rows = _exec_rows("SELECT REVISI FROM MCMOREVISI WITH (NOLOCK) WHERE NEWAPPL_NO = ?", [appl_no])
    return f"R{rows[0]['REVISI']}" if rows else ''

def _fetch_approval_history(appl_no: str, require_comment: bool) -> list[dict]:
    comment_filter = "AND a.comment IS NOT NULL AND a.comment <> ''" if require_comment else ""
    sql = f"""
        SELECT * FROM (
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMOUTSSTATUS a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
            
            UNION
            
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMStatus a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
        ) a
        WHERE a.APPR_ID IN (SELECT username FROM app_usrmst WHERE ACCESSCAM IN ('CRH', 'CRDH', 'DR', 'PD'))
            {comment_filter}
        ORDER BY a.[No]
    """
    rows = _exec_rows(sql, [appl_no, appl_no])
    history = []
    for r in rows:
        appr_date = r.get('APPR_DATE')
        history.append({
            'date': appr_date.strftime('%d-%m-%Y %H:%M:%S') if appr_date else '',
            'approved_by': r.get('APPR_ID'),
            'status': _STATUS_LABELS.get(r.get('APPR_STS'), ''),
            'comment': r.get('COMMENT') or '',
        })
    return history

def _fetch_related_cams(appl_no: str, apless: Optional[str]) -> list[dict]:
    related_rows = _exec_multi_resultset("EXEC USP_ListCamCommittee ?, ?", [appl_no, apless])
    related = []
    for r in related_rows:
        related_appl_no = r.get('appl_no')
        if not related_appl_no:
            continue
        existence_check = _fetch_approval_history(related_appl_no, require_comment=True)
        if not existence_check:
            continue
        history = _fetch_approval_history(related_appl_no, require_comment=False)
        info = _fetch_cam_basic_info(related_appl_no)
        suffix = _fetch_revision_suffix(related_appl_no)
        related.append({
            'appl_no_display': f"{related_appl_no}{suffix}",
            'customer_name': info.get('LESSEE_NM'),
            'cmo_name': info.get('USERNAME'),
            'history': history,
        })
    return related

@cam_by_approver_bp.route('/cam-by-approver/comments')
@login_required
def get_comments():
    appl_no = request.args.get('appl_no', '').strip()
    if not appl_no:
        return jsonify({'error': 'appl_no is required'}), 400

    try:
        info = _fetch_cam_basic_info(appl_no)
        suffix = _fetch_revision_suffix(appl_no)
        history = _fetch_approval_history(appl_no, require_comment=False)
        related = _fetch_related_cams(appl_no, info.get('apless'))
        return jsonify({
            'main': {
                'appl_no_display': f"{appl_no}{suffix}",
                'customer_name': info.get('LESSEE_NM'),
                'cmo_name': info.get('USERNAME'),
                'history': history,
            },
            'related': related,
        })
    except Exception as e:
        current_app.logger.exception("Credit committee comments error")
        return jsonify({'error': str(e)}), 500