from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

cam_by_car_condition_bp = Blueprint('cam_by_car_condition', __name__, url_prefix='/CAM/Approved')

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _resolve_scope(requested_branch_cd: str) -> dict:
    access_level = session.get('accesslevel')
    own_branch_cd = session.get('branchcode')

    if access_level in ('CMO', 'CMH', 'BM'):
        return {'disabled': True, 'search_branch': own_branch_cd, 'sp_branch_param': ''}
    if own_branch_cd == '100':
        return {'disabled': False, 'search_branch': requested_branch_cd, 'sp_branch_param': requested_branch_cd}
    return {'disabled': True, 'search_branch': own_branch_cd, 'sp_branch_param': own_branch_cd}

@cam_by_car_condition_bp.route('/cam-by-car-condition/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@cam_by_car_condition_bp.route('/cam-by-car-condition/context')
@login_required
def get_context():
    resolved = _resolve_scope(request.args.get('branch_cd_search', '').strip())
    return jsonify({'disabled': resolved['disabled'], 'forced_branch_cd': resolved['search_branch']})

def _fetch_and_process(start_date: Optional[date], end_date: Optional[date], scope: dict) -> dict:
    sql = "EXEC USP_S_APPR_COND ?, ?, ?, ?, ?, ?"
    params = [
        start_date.isoformat() if start_date else '',
        end_date.isoformat() if end_date else '',
        session.get('emp_id'),
        session.get('branchcode'),
        session.get('accesslevel'),
        scope['sp_branch_param'],
    ]
    rows = _exec_rows(sql, params)

    detail = []
    totals = {
        'new_nontruck_unit': 0.0, 'new_nontruck_amt': 0.0,
        'new_truck_unit': 0.0, 'new_truck_amt': 0.0,
        'used_nontruck_unit': 0.0, 'used_nontruck_amt': 0.0,
        'used_truck_unit': 0.0, 'used_truck_amt': 0.0,
    }
    temp_date = None

    for r in rows:
        adm_date = r.get('AdmDate')
        date_str = adm_date.strftime('%d-%m-%Y') if adm_date else ''
        vals = {
            'new_nontruck_unit': _num(r.get('unit_newnontruck')),
            'new_nontruck_amt': _num(r.get('Net_Finance_newnontruck')),
            'new_truck_unit': _num(r.get('unit_newtruck')),
            'new_truck_amt': _num(r.get('Net_Finance_newtruck')),
            'used_nontruck_unit': _num(r.get('unit_used_non')),
            'used_nontruck_amt': _num(r.get('Net_Finance_usednon')),
            'used_truck_unit': _num(r.get('unit_used_truck')),
            'used_truck_amt': _num(r.get('Net_Finance_usedtruck')),
        }
        detail.append({
            'approval_date': date_str if date_str != temp_date else '',
            'new_nontruck_unit': int(vals['new_nontruck_unit']), 'new_nontruck_amt': vals['new_nontruck_amt'],
            'new_truck_unit': int(vals['new_truck_unit']), 'new_truck_amt': vals['new_truck_amt'],
            'used_nontruck_unit': int(vals['used_nontruck_unit']), 'used_nontruck_amt': vals['used_nontruck_amt'],
            'used_truck_unit': int(vals['used_truck_unit']), 'used_truck_amt': vals['used_truck_amt'],
        })
        temp_date = date_str
        for k in totals:
            totals[k] += vals[k]

    return {'detail': detail, 'summary': _derive_summary(totals), 'has_data': len(rows) > 0}

def _derive_summary(totals: dict) -> dict:
    def pct(n, d):
        return round((n / d) * 100) if d else 0

    total_amt_new = totals['new_nontruck_amt'] + totals['new_truck_amt']
    total_amt_used = totals['used_nontruck_amt'] + totals['used_truck_amt']
    total_new_unit = totals['new_nontruck_unit'] + totals['new_truck_unit']
    total_used_unit = totals['used_nontruck_unit'] + totals['used_truck_unit']

    total_nontruck_unit = totals['new_nontruck_unit'] + totals['used_nontruck_unit']
    total_nontruck_amt = totals['new_nontruck_amt'] + totals['used_nontruck_amt']
    total_truck_unit = totals['new_truck_unit'] + totals['used_truck_unit']
    total_truck_amt = totals['new_truck_amt'] + totals['used_truck_amt']

    grand_unit_tc = total_nontruck_unit + total_truck_unit
    grand_amt_tc = total_nontruck_amt + total_truck_amt
    grand_unit = total_new_unit + total_used_unit
    grand_amt = total_amt_new + total_amt_used

    persen_non_truck = pct(total_nontruck_unit, grand_unit_tc)
    persen_truck = pct(total_truck_unit, grand_unit_tc)
    persen_non_truck_amt = pct(total_nontruck_amt, grand_amt_tc)
    persen_truck_amt = pct(total_truck_amt, grand_amt_tc)
    persen_new = pct(total_new_unit, grand_unit)
    persen_used = pct(total_used_unit, grand_unit)
    persen_amt_new = pct(total_amt_new, grand_amt)
    persen_amt_used = pct(total_amt_used, grand_amt)

    return {
        'by_condition': {
            'new': {'units': int(total_new_unit), 'pct_units': persen_new, 'net_finance': total_amt_new, 'pct_net_finance': persen_amt_new},
            'used': {'units': int(total_used_unit), 'pct_units': persen_used, 'net_finance': total_amt_used, 'pct_net_finance': persen_amt_used},
            'total': {
                'units': int(grand_unit), 'pct_units': persen_new + persen_used,
                'net_finance': grand_amt, 'pct_net_finance': persen_amt_new + persen_amt_used,
            },
        },
        'by_type': {
            'non_truck': {'units': int(total_nontruck_unit), 'pct_units': persen_non_truck, 'net_finance': total_nontruck_amt, 'pct_net_finance': persen_non_truck_amt},
            'truck': {'units': int(total_truck_unit), 'pct_units': persen_truck, 'net_finance': total_truck_amt, 'pct_net_finance': persen_truck_amt},
            'total': {
                'units': int(grand_unit_tc), 'pct_units': persen_non_truck + persen_truck,
                'net_finance': grand_amt_tc, 'pct_net_finance': persen_non_truck_amt + persen_truck_amt,
            },
        },
        'totals_row': {
            'new_nontruck_unit': int(totals['new_nontruck_unit']), 'new_nontruck_amt': totals['new_nontruck_amt'],
            'new_truck_unit': int(totals['new_truck_unit']), 'new_truck_amt': totals['new_truck_amt'],
            'used_nontruck_unit': int(totals['used_nontruck_unit']), 'used_nontruck_amt': totals['used_nontruck_amt'],
            'used_truck_unit': int(totals['used_truck_unit']), 'used_truck_amt': totals['used_truck_amt'],
        },
    }

@cam_by_car_condition_bp.route('/cam-by-car-condition/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    scope = _resolve_scope(requested_branch)
    try:
        result = _fetch_and_process(start_date, end_date, scope)
        return jsonify({
            'data': result['detail'],
            'summary': result['summary'],
            'has_data': result['has_data'],
            'disabled': scope['disabled'],
            'search_branch': scope['search_branch'],
        })
    except Exception as e:
        current_app.logger.exception("CAM by Car Condition list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

def _build_workbook(start_date, end_date, scope, result: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'CamByCarCondition'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = "List Approved Cam By Car's Condition"
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {_branch_name(scope['search_branch'])}"

    ws.merge_cells('A4:D4')
    ws['A4'] = _period_label(start_date, end_date)

    header_row1 = 6
    header_row2 = 7
    header_row3 = 8
    ws.merge_cells(f'A{header_row1}:A{header_row3}')
    ws.merge_cells(f'B{header_row1}:E{header_row1}')
    ws.merge_cells(f'F{header_row1}:I{header_row1}')
    ws.merge_cells(f'B{header_row2}:C{header_row2}')
    ws.merge_cells(f'D{header_row2}:E{header_row2}')
    ws.merge_cells(f'F{header_row2}:G{header_row2}')
    ws.merge_cells(f'H{header_row2}:I{header_row2}')

    ws[f'A{header_row1}'] = 'Approval Date'
    ws[f'B{header_row1}'] = 'NEW'
    ws[f'F{header_row1}'] = 'USED'
    ws[f'B{header_row2}'] = 'Non Truck'
    ws[f'D{header_row2}'] = 'Truck'
    ws[f'F{header_row2}'] = 'Non Truck'
    ws[f'H{header_row2}'] = 'Truck'
    for col in ['B', 'D', 'F', 'H']:
        ws[f'{col}{header_row3}'] = 'Unit'
    for col in ['C', 'E', 'G', 'I']:
        ws[f'{col}{header_row3}'] = 'Net Finance'

    for row_num in (header_row1, header_row2, header_row3):
        for col in 'ABCDEFGHI':
            cell = ws[f'{col}{row_num}']
            cell.border = _ALL_BORDERS
            cell.alignment = _CENTER
            cell.font = _HEADER_FONT

    excel_row = header_row3 + 1
    detail = result['detail']
    if not detail:
        ws.merge_cells(f'A{excel_row}:I{excel_row}')
        cell = ws[f'A{excel_row}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        excel_row += 1
    else:
        for row in detail:
            ws.cell(row=excel_row, column=1, value=row['approval_date']).border = _ALL_BORDERS
            values = [
                row['new_nontruck_unit'], row['new_nontruck_amt'],
                row['new_truck_unit'], row['new_truck_amt'],
                row['used_nontruck_unit'], row['used_nontruck_amt'],
                row['used_truck_unit'], row['used_truck_amt'],
            ]
            for i, val in enumerate(values, start=2):
                cell = ws.cell(row=excel_row, column=i, value=val)
                cell.border = _ALL_BORDERS
                cell.alignment = _RIGHT
            excel_row += 1

        tr = result['summary']['totals_row']
        totals_values = [
            tr['new_nontruck_unit'], tr['new_nontruck_amt'], tr['new_truck_unit'], tr['new_truck_amt'],
            tr['used_nontruck_unit'], tr['used_nontruck_amt'], tr['used_truck_unit'], tr['used_truck_amt'],
        ]
        ws.cell(row=excel_row, column=1).border = _ALL_BORDERS
        for i, val in enumerate(totals_values, start=2):
            cell = ws.cell(row=excel_row, column=i, value=val)
            cell.border = _ALL_BORDERS
            cell.alignment = _RIGHT
            cell.font = _HEADER_FONT
        excel_row += 2

    by_cond = result['summary']['by_condition']
    by_type = result['summary']['by_type']

    def _write_breakdown_table(start_row: int, rows_spec: list) -> int:
        headers = ['', 'Unit', '%', 'Net Finance', '%']
        for col_idx, label in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=label)
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
        r = start_row + 1
        for label, bucket in rows_spec:
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=bucket['units']).alignment = _RIGHT
            ws.cell(row=r, column=3, value=f"{bucket['pct_units']}%").alignment = _RIGHT
            ws.cell(row=r, column=4, value=bucket['net_finance']).alignment = _RIGHT
            ws.cell(row=r, column=5, value=f"{bucket['pct_net_finance']}%").alignment = _RIGHT
            r += 1
        return r + 1

    next_row = _write_breakdown_table(excel_row, [
        ('New', by_cond['new']), ('Used', by_cond['used']), ('Total', by_cond['total']),
    ])
    _write_breakdown_table(next_row, [
        ('Non Truck', by_type['non_truck']), ('Truck', by_type['truck']), ('Total', by_type['total']),
    ])

    ws.column_dimensions['A'].width = 16
    for col in 'BCDEFGHI':
        ws.column_dimensions[col].width = 14

    return wb

@cam_by_car_condition_bp.route('/cam-by-car-condition/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    scope = _resolve_scope(requested_branch)
    try:
        result = _fetch_and_process(start_date, end_date, scope)
        wb = _build_workbook(start_date, end_date, scope, result)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ApprovedCamByCarCond.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("CAM by Car Condition export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Credit Approval Memorandum</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
.b1 {{ border-left:thin solid #000; border-bottom:thin solid #000; border-right:thin solid #000; }}
.b4 {{ border:thin solid #000; }}
.b5 {{ border-left:thin solid #000; border-right:thin solid #000; }}
@media print {{
    body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; size: landscape; }}
}}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();window.close();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%" nowrap="nowrap">Date : {today}<br>Print By : {printed_by}</td>
        </tr>
        <tr>
            <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List Approved Cam By Car`s Condition</strong></td>
        </tr>
        <tr><td colspan="4">Branch : {branch_name}</td></tr>
        <tr><td colspan="4" nowrap="nowrap">{period}</td></tr>
    </table>
    <table border="0" style="border-collapse:collapse;" cellpadding="2" cellspacing="0" width="100%">
        <tr align="center">
            <td width="12%" rowspan="3" style="border:thin solid #000;">Approval Date</td>
            <td colspan="4" style="border:thin solid #000;">NEW</td>
            <td colspan="4" style="border:thin solid #000;">USED</td>
        </tr>
        <tr align="center">
            <td colspan="2" style="border:thin solid #000;">Non Truck</td>
            <td colspan="2" style="border:thin solid #000;">Truck</td>
            <td colspan="2" style="border:thin solid #000;">Non Truck</td>
            <td colspan="2" style="border:thin solid #000;">Truck</td>
        </tr>
        <tr align="center">
            <td style="border:thin solid #000;">Unit</td><td style="border:thin solid #000;">Net Finance</td>
            <td style="border:thin solid #000;">Unit</td><td style="border:thin solid #000;">Net Finance</td>
            <td style="border:thin solid #000;">Unit</td><td style="border:thin solid #000;">Net Finance</td>
            <td style="border:thin solid #000;">Unit</td><td style="border:thin solid #000;">Net Finance</td>
        </tr>
        {rows_html}
    </table>
    <br/>
    <table border="0" cellpadding="2" cellspacing="0" width="56%">
        <tr align="center"><td width="21%" class="b4">&nbsp;</td><td width="13%" class="b4">Unit</td><td width="20%" class="b4">%</td><td width="27%" class="b4">Net Finance</td><td width="19%" class="b4">%</td></tr>
        <tr><td class="b5" align="left">New</td><td class="b5" align="right">{new_units}</td><td class="b5" align="right">{new_pct_units}%</td><td class="b5" align="right">{new_fin}</td><td class="b5" align="right">{new_pct_fin}%</td></tr>
        <tr><td class="b1" align="left">Used</td><td class="b1" align="right">{used_units}</td><td class="b1" align="right">{used_pct_units}%</td><td class="b1" align="right">{used_fin}</td><td class="b1" align="right">{used_pct_fin}%</td></tr>
        <tr><td class="b1" align="left">Total</td><td class="b1" align="right">{cond_total_units}</td><td class="b1" align="right">{cond_total_pct_units}%</td><td class="b1" align="right">{cond_total_fin}</td><td class="b1" align="right">{cond_total_pct_fin}%</td></tr>
    </table>
    <br/>
    <table border="0" cellpadding="2" cellspacing="0" width="56%">
        <tr align="center"><td width="21%" class="b4">&nbsp;</td><td width="13%" class="b4">Unit</td><td width="20%" class="b4">%</td><td width="27%" class="b4">Net Finance</td><td width="19%" class="b4">%</td></tr>
        <tr><td class="b5" align="left">Non Truck</td><td class="b5" align="right">{nontruck_units}</td><td class="b5" align="right">{nontruck_pct_units}%</td><td class="b5" align="right">{nontruck_fin}</td><td class="b5" align="right">{nontruck_pct_fin}%</td></tr>
        <tr><td class="b1" align="left">Truck</td><td class="b1" align="right">{truck_units}</td><td class="b1" align="right">{truck_pct_units}%</td><td class="b1" align="right">{truck_fin}</td><td class="b1" align="right">{truck_pct_fin}%</td></tr>
        <tr><td class="b1" align="left">Total</td><td class="b1" align="right">{type_total_units}</td><td class="b1" align="right">{type_total_pct_units}%</td><td class="b1" align="right">{type_total_fin}</td><td class="b1" align="right">{type_total_pct_fin}</td></tr>
    </table>
</body>
</html>"""

def _render_detail_rows_html(detail: list[dict], totals_row: dict) -> str:
    if not detail:
        return '<tr><td colspan="9" align="center"><font color="red"><b>Data Not Found</b></font></td></tr>'
    parts = []
    for row in detail:
        cells = [f'<td class="b1">{row["approval_date"]}</td>' if row['approval_date'] else '<td></td>']
        for key in ['new_nontruck_unit', 'new_nontruck_amt', 'new_truck_unit', 'new_truck_amt',
                    'used_nontruck_unit', 'used_nontruck_amt', 'used_truck_unit', 'used_truck_amt']:
            val = row[key]
            fmt_val = f"{val:,.0f}" if isinstance(val, float) else f"{val:,}"
            cells.append(f'<td align="right" class="b1">{fmt_val}</td>')
        parts.append(f"<tr>{''.join(cells)}</tr>")
    tr = totals_row
    total_cells = ['<td class="b4"></td>']
    for key in ['new_nontruck_unit', 'new_nontruck_amt', 'new_truck_unit', 'new_truck_amt',
                'used_nontruck_unit', 'used_nontruck_amt', 'used_truck_unit', 'used_truck_amt']:
        val = tr[key]
        fmt_val = f"{val:,.0f}" if isinstance(val, float) else f"{val:,}"
        total_cells.append(f'<td align="right" class="b1">{fmt_val}</td>')
    parts.append(f"<tr>{''.join(total_cells)}</tr>")
    return ''.join(parts)

@cam_by_car_condition_bp.route('/cam-by-car-condition/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    scope = _resolve_scope(requested_branch)
    try:
        result = _fetch_and_process(start_date, end_date, scope)
        by_cond = result['summary']['by_condition']
        by_type = result['summary']['by_type']

        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_branch_name(scope['search_branch']),
            period=_period_label(start_date, end_date),
            rows_html=_render_detail_rows_html(result['detail'], result['summary']['totals_row']),
            new_units=f"{by_cond['new']['units']:,}", new_pct_units=by_cond['new']['pct_units'],
            new_fin=f"{by_cond['new']['net_finance']:,.0f}", new_pct_fin=by_cond['new']['pct_net_finance'],
            used_units=f"{by_cond['used']['units']:,}", used_pct_units=by_cond['used']['pct_units'],
            used_fin=f"{by_cond['used']['net_finance']:,.0f}", used_pct_fin=by_cond['used']['pct_net_finance'],
            cond_total_units=f"{by_cond['total']['units']:,}", cond_total_pct_units=by_cond['total']['pct_units'],
            cond_total_fin=f"{by_cond['total']['net_finance']:,.0f}", cond_total_pct_fin=by_cond['total']['pct_net_finance'],
            nontruck_units=f"{by_type['non_truck']['units']:,}", nontruck_pct_units=by_type['non_truck']['pct_units'],
            nontruck_fin=f"{by_type['non_truck']['net_finance']:,.0f}", nontruck_pct_fin=by_type['non_truck']['pct_net_finance'],
            truck_units=f"{by_type['truck']['units']:,}", truck_pct_units=by_type['truck']['pct_units'],
            truck_fin=f"{by_type['truck']['net_finance']:,.0f}", truck_pct_fin=by_type['truck']['pct_net_finance'],
            type_total_units=f"{by_type['total']['units']:,}", type_total_pct_units=by_type['total']['pct_units'],
            type_total_fin=f"{by_type['total']['net_finance']:,.0f}", type_total_pct_fin=by_type['total']['pct_net_finance'],
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("CAM by Car Condition print error")
        return jsonify({'error': str(e)}), 500