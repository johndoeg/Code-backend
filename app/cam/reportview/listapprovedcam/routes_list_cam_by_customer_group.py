from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

cam_by_cust_group_bp = Blueprint('cam_by_cust_group', __name__, url_prefix='/CAM/Approved')

PAGE_SIZE = 100

_STATUS_LABELS = {
    'A': 'Approved',
    'F': 'Approved & Send to Admin',
    'R': 'Reject',
    'N': 'Send for Approval',
}

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _exec_multi_resultset(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    all_rows: list[dict] = []
    try:
        cursor.execute(sql, params)
        while True:
            if cursor.description is not None:
                columns = [col[0] for col in cursor.description]
                all_rows += [dict(zip(columns, row)) for row in cursor.fetchall()]
            if not cursor.nextset():
                break
        return all_rows
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _own_branch_name() -> str:
    branch_cd = session.get('branchcode')
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''


def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _resolve_scope() -> dict:
    access_level = session.get('accesslevel')
    emp_id = session.get('emp_id')
    own_branch_cd = session.get('branchcode')

    if access_level == 'CMO':
        return {'cmo_filter': [emp_id], 'branch_filter': None}

    if access_level == 'CMH':
        rows = _exec_rows(
            "SELECT EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) WHERE EMPLOYEE_ID = ? OR SUPERIOR = ?",
            [emp_id, emp_id],
        )
        cmo_ids = [r['EMPLOYEE_ID'] for r in rows] or [emp_id]
        return {'cmo_filter': cmo_ids, 'branch_filter': None}

    if access_level == 'BM':
        return {'cmo_filter': None, 'branch_filter': own_branch_cd}

    return {'cmo_filter': None, 'branch_filter': None if own_branch_cd == '100' else own_branch_cd}

def _build_common_filter(scope: dict, start_date: Optional[date], end_date: Optional[date],
                          lessee_nm: str, grp_code: str) -> tuple[str, list]:
    clauses = []
    params: list = []

    if scope['cmo_filter'] is not None:
        placeholders = ','.join('?' for _ in scope['cmo_filter'])
        clauses.append(f"a.CMO IN ({placeholders})")
        params.extend(scope['cmo_filter'])
    elif scope['branch_filter']:
        clauses.append("a.BRANCH_CD = ?")
        params.append(scope['branch_filter'])

    if grp_code:
        clauses.append("b.GrpCode = ?")
        params.append(grp_code)

    if start_date and end_date:
        clauses.append("CAST(a.CREATE_DATE AS DATE) >= ? AND CAST(a.CREATE_DATE AS DATE) <= ?")
        params.append(start_date)
        params.append(end_date)

    if lessee_nm:
        clauses.append("b.LESSEE_NM LIKE ?")
        params.append(f"%{lessee_nm}%")

    if not clauses:
        return "1=1", []
    return " AND ".join(clauses), params

_LIST_SELECT = """
SELECT a.APPL_NO, a.Cons_Leas, b.LESSEE_NM, a.TENOR, a.DECLRATE, j.AdmDate,
       c.USERNAME, h.BRAND_NM, i.TYPE_NM, v.Net_Finance
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    LEFT JOIN APP_USRMST c WITH (NOLOCK) ON a.CMO = c.EMPLOYEE_ID
    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
    LEFT JOIN UMFBRAND h WITH (NOLOCK) ON g.BRAND = h.BRAND
    LEFT JOIN UMFTYPE i WITH (NOLOCK) ON g.BRAND = i.BRAND AND g.TYPE = i.TYPE
    INNER JOIN mcmocamcompletelist j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO
    INNER JOIN V_NetFinance v WITH (NOLOCK) ON a.APPL_NO = v.APPL_NO
"""

def _fetch_count(where_sql: str, params: list) -> int:
    sql = f"""
    SELECT COUNT(*) AS cnt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        INNER JOIN mcmocamcompletelist j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO
        INNER JOIN V_NetFinance v WITH (NOLOCK) ON a.APPL_NO = v.APPL_NO
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    return int(rows[0]['cnt']) if rows else 0

def _fetch_page(where_sql: str, params: list, page: int, page_size: int) -> list[dict]:
    offset = (page - 1) * page_size
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY j.AdmDate OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
    return _exec_rows(sql, params + [offset, page_size])

def _fetch_all(where_sql: str, params: list) -> list[dict]:
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY j.AdmDate"
    return _exec_rows(sql, params)

def _fetch_date_range(where_sql: str, params: list) -> tuple[Optional[date], Optional[date]]:
    sql = f"""
    SELECT MIN(j.AdmDate) AS min_dt, MAX(j.AdmDate) AS max_dt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        INNER JOIN mcmocamcompletelist j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO
        INNER JOIN V_NetFinance v WITH (NOLOCK) ON a.APPL_NO = v.APPL_NO
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    if not rows:
        return None, None
    min_dt, max_dt = rows[0].get('min_dt'), rows[0].get('max_dt')
    return (min_dt.date() if min_dt else None), (max_dt.date() if max_dt else None)

def _fetch_revisions(appl_nos: list[str]) -> dict[str, str]:
    if not appl_nos:
        return {}
    placeholders = ','.join('?' for _ in appl_nos)
    rows = _exec_rows(
        f"SELECT NEWAPPL_NO, REVISI FROM MCMORevisi WITH (NOLOCK) WHERE NEWAPPL_NO IN ({placeholders})",
        appl_nos,
    )
    return {r['NEWAPPL_NO']: f"R{r['REVISI']}" for r in rows}

def _row_to_detail(r: dict, no: int, revision: str) -> dict:
    adm_date = r.get('AdmDate')
    brand = (r.get('BRAND_NM') or '').strip()
    type_nm = (r.get('TYPE_NM') or '').strip()
    return {
        'no': no,
        'cam_no': f"{r.get('Cons_Leas')} - {r.get('APPL_NO')}{revision}",
        'appl_no': r.get('APPL_NO'),
        'customer_name': r.get('LESSEE_NM') or '-',
        'brand_type': f"{brand} / {type_nm}" if (brand or type_nm) else '-',
        'net_finance': _num(r.get('Net_Finance')),
        'interest': round(_num(r.get('DECLRATE')), 4),
        'tenor': int(_num(r.get('TENOR'))),
        'approval_date': adm_date.strftime('%d-%m-%Y') if adm_date else '',
        'cmo': r.get('USERNAME') or '-',
        'is_revision': bool(revision),
    }

_TOTAL_BASE = """
SELECT ISNULL(SUM(v.Net_Finance), 0) AS tot_fin, ISNULL(COUNT(a.APPL_NO), 0) AS tot_unit
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    INNER JOIN mcmocamcompletelist j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO
    INNER JOIN V_NetFinance v WITH (NOLOCK) ON a.APPL_NO = v.APPL_NO
WHERE a.APPL_NO NOT IN (SELECT NewAppl_No FROM MCMORevisi WITH (NOLOCK) WHERE NewAppl_No IS NOT NULL)
  AND {filter}
"""

def _fetch_grand_total(where_sql: str, params: list) -> dict:
    sql = _TOTAL_BASE.format(filter=where_sql)
    rows = _exec_rows(sql, params)
    row = rows[0] if rows else {}
    return {'net_finance': _num(row.get('tot_fin')), 'units': int(_num(row.get('tot_unit')))}

@cam_by_cust_group_bp.route('/cam-by-cust-group/groups')
@login_required
def list_groups():
    try:
        rows = _exec_rows(
            "SELECT RTRIM(GrpCode) AS GrpCode, RTRIM(GrpName) AS GrpName FROM CLFGROUP WITH (NOLOCK) WHERE FLGDEL = '0'",
            [],
        )
        return jsonify([{'grp_code': r.get('GrpCode'), 'grp_name': r.get('GrpName')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Group list error")
        return jsonify({'error': str(e)}), 500

@cam_by_cust_group_bp.route('/cam-by-cust-group/context')
@login_required
def get_context():
    return jsonify({'branch_name': _own_branch_name()})

@cam_by_cust_group_bp.route('/cam-by-cust-group/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    grp_code = request.args.get('grp_code', '').strip()
    try:
        page = max(int(request.args.get('page', 1) or 1), 1)
    except ValueError:
        page = 1

    scope = _resolve_scope()
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm, grp_code)

    try:
        total_count = _fetch_count(where_sql, params)
        total_pages = max((total_count + PAGE_SIZE - 1) // PAGE_SIZE, 1)
        page = min(page, total_pages)

        raw_rows = _fetch_page(where_sql, params, page, PAGE_SIZE)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)

        start_no = (page - 1) * PAGE_SIZE + 1
        details = [_row_to_detail(r, start_no + i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows)]

        grand_total = _fetch_grand_total(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        return jsonify({
            'rows': details,
            'total_count': total_count,
            'page': page,
            'page_size': PAGE_SIZE,
            'total_pages': total_pages,
            'grand_total': grand_total,
            'period_label': period,
            'branch_name': _own_branch_name(),
        })
    except Exception as e:
        current_app.logger.exception("CAM by Customer/Group list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=16)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

_LIST_COLUMNS = ['No.', 'CAM No.', 'Customer Name', 'Brand/Type', 'Net Finance', 'Interest (%)', 'Tenor (Month)',
                  'Approval Date', 'CMO']

def _build_workbook(period_label: str, lessee_nm: str, branch_name: str, rows: list[dict], grand_total: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'CamByCustGroup'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List CAM by Customer / Group'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {branch_name}"

    ws.merge_cells('A4:D4')
    ws['A4'] = period_label or (f"Name : {lessee_nm}" if lessee_nm else '')

    row_num = 6
    for col_idx, label in enumerate(_LIST_COLUMNS, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=label)
        cell.border = _ALL_BORDERS
        cell.alignment = _CENTER
        cell.font = _HEADER_FONT
    row_num += 1

    if not rows:
        ws.merge_cells(f'A{row_num}:I{row_num}')
        cell = ws[f'A{row_num}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        row_num += 1
    else:
        for d in rows:
            values = [d['no'], d['cam_no'], d['customer_name'], d['brand_type'], d['net_finance'],
                      d['interest'], d['tenor'], d['approval_date'], d['cmo']]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = _ALL_BORDERS
                if col_idx in (5, 6):
                    cell.alignment = _RIGHT
            row_num += 1

    row_num += 2
    ws.merge_cells(f'A{row_num}:B{row_num}')
    ws[f'A{row_num}'] = f"Grand Total* ({period_label.replace('Period : ', '') or 'All Dates'})"
    ws[f'A{row_num}'].font = _HEADER_FONT
    row_num += 1
    ws.cell(row=row_num, column=1, value='Net Finance')
    ws.cell(row=row_num, column=2, value=grand_total['net_finance']).alignment = _RIGHT
    row_num += 1
    ws.cell(row=row_num, column=1, value='Units')
    ws.cell(row=row_num, column=2, value=grand_total['units']).alignment = _RIGHT
    row_num += 2
    ws.merge_cells(f'A{row_num}:D{row_num}')
    ws[f'A{row_num}'] = '*Grand Total Amount and Unit Exclude CAM Revision'
    ws[f'A{row_num}'].font = Font(italic=True)

    ws.column_dimensions['A'].width = 6
    for col in 'BCDEFGHI':
        ws.column_dimensions[col].width = 16

    return wb

@cam_by_cust_group_bp.route('/cam-by-cust-group/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    grp_code = request.args.get('grp_code', '').strip()

    scope = _resolve_scope()
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm, grp_code)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_total = _fetch_grand_total(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        wb = _build_workbook(period, lessee_nm, _own_branch_name(), rows, grand_total)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ListCamByCustGroup.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("CAM by Customer/Group export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{company} - List CAM by Customer / Group</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
        </tr>
        <tr>
            <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List CAM by Customer / Group</strong></td>
        </tr>
        <tr><td colspan="4">Branch : {branch_name}</td></tr>
        <tr><td colspan="4" nowrap="nowrap">{period}</td></tr>
    </table>
    <table cellpadding="2" cellspacing="0" width="100%" style="margin-top:6px;">
        <tr align="center">
            <td style="border:thin solid #000" width="3%">No.</td>
            <td style="border:thin solid #000" width="14%">CAM No.</td>
            <td style="border:thin solid #000" width="16%">Customer Name</td>
            <td style="border:thin solid #000" width="17%">Brand/Type</td>
            <td style="border:thin solid #000" width="12%">Net Finance</td>
            <td style="border:thin solid #000" width="6%">Interest (%)</td>
            <td style="border:thin solid #000" width="5%">Tenor (Month)</td>
            <td style="border:thin solid #000" width="11%">Approval Date</td>
            <td style="border:thin solid #000" width="16%">CMO</td>
        </tr>
        {rows_html}
    </table>
    <table border="0" style="border-collapse:collapse; margin-top:10px;" cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="40%" colspan="2">
                <u><strong>Grand Total* ({period_plain}):</strong></u>
            </td>
        </tr>
        <tr>
            <td width="10%">Net Finance</td>
            <td>: {grand_fin}</td>
        </tr>
        <tr>
            <td width="10%">Units</td>
            <td> : {grand_units}</td>
        </tr>
        <tr><td colspan="2">&nbsp;</td></tr>
        <tr><td colspan="2"><i>*Grand Total Amount and Unit Exclude CAM Revision</i></td></tr>
    </table>
</body>
</html>"""

def _render_rows_html(rows: list[dict]) -> str:
    if not rows:
        return ('<tr><td colspan="9" align="center" style="border:thin solid #000">'
                '<font color="red"><b>Data Not Found</b></font></td></tr>')
    parts = []
    for d in rows:
        parts.append(f"""
            <tr>
                <td style="border:thin solid #000" align="center">{d['no']}</td>
                <td style="border:thin solid #000">{d['cam_no']}</td>
                <td style="border:thin solid #000">{d['customer_name']}</td>
                <td style="border:thin solid #000">{d['brand_type']}</td>
                <td style="border:thin solid #000" align="right">{d['net_finance']:,.0f}</td>
                <td style="border:thin solid #000" align="right">{d['interest']:.4f}</td>
                <td style="border:thin solid #000" align="center">{d['tenor']}</td>
                <td style="border:thin solid #000">{d['approval_date']}</td>
                <td style="border:thin solid #000">{d['cmo']}</td>
            </tr>""")
    return ''.join(parts)

@cam_by_cust_group_bp.route('/cam-by-cust-group/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    grp_code = request.args.get('grp_code', '').strip()

    scope = _resolve_scope()
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm, grp_code)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_total = _fetch_grand_total(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end) or (f"Name : {lessee_nm}" if lessee_nm else '')

        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_own_branch_name(),
            period=period,
            period_plain=period.replace('Period : ', '') or 'All Dates',
            rows_html=_render_rows_html(rows),
            grand_fin=f"{grand_total['net_finance']:,.0f}",
            grand_units=grand_total['units'],
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("CAM by Customer/Group print error")
        return jsonify({'error': str(e)}), 500

def _fetch_cam_basic_info(appl_no: str) -> dict:
    rows = _exec_rows(
        """
        SELECT LTRIM(RTRIM(a.CMO)) AS CMO, b.LESSEE_NM, b.apless, c.USERNAME
        FROM CAFCONT a WITH (NOLOCK)
            LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
            LEFT JOIN APP_USRMST c WITH (NOLOCK) ON a.CMO = c.EMPLOYEE_ID
        WHERE a.APPL_NO = ?
        """,
        [appl_no],
    )
    return rows[0] if rows else {}

def _fetch_revision_suffix(appl_no: str) -> str:
    rows = _exec_rows("SELECT REVISI FROM MCMOREVISI WITH (NOLOCK) WHERE NEWAPPL_NO = ?", [appl_no])
    return f"R{rows[0]['REVISI']}" if rows else ''

def _fetch_approval_history(appl_no: str, require_comment: bool) -> list[dict]:
    comment_filter = "AND a.comment IS NOT NULL AND a.comment <> ''" if require_comment else ""
    sql = f"""
        SELECT * FROM (
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMOUTSSTATUS a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
            
            UNION
            
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMStatus a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
        ) a
        WHERE a.APPR_ID IN (SELECT username FROM app_usrmst WHERE ACCESSCAM IN ('CRH', 'CRDH', 'DR', 'PD'))
            {comment_filter}
        ORDER BY a.[No]
    """
    rows = _exec_rows(sql, [appl_no, appl_no])
    history = []
    for r in rows:
        appr_date = r.get('APPR_DATE')
        history.append({
            'date': appr_date.strftime('%d-%m-%Y %H:%M:%S') if appr_date else '',
            'approved_by': r.get('APPR_ID'),
            'status': _STATUS_LABELS.get(r.get('APPR_STS'), ''),
            'comment': r.get('COMMENT') or '',
        })
    return history

def _fetch_related_cams(appl_no: str, apless: Optional[str]) -> list[dict]:
    related_rows = _exec_multi_resultset("EXEC USP_ListCamCommittee ?, ?", [appl_no, apless])
    related = []
    for r in related_rows:
        related_appl_no = r.get('appl_no')
        if not related_appl_no:
            continue
        existence_check = _fetch_approval_history(related_appl_no, require_comment=True)
        if not existence_check:
            continue
        history = _fetch_approval_history(related_appl_no, require_comment=False)
        info = _fetch_cam_basic_info(related_appl_no)
        suffix = _fetch_revision_suffix(related_appl_no)
        related.append({
            'appl_no_display': f"{related_appl_no}{suffix}",
            'customer_name': info.get('LESSEE_NM'),
            'cmo_name': info.get('USERNAME'),
            'history': history,
        })
    return related

@cam_by_cust_group_bp.route('/cam-by-cust-group/comments')
@login_required
def get_comments():
    appl_no = request.args.get('appl_no', '').strip()
    if not appl_no:
        return jsonify({'error': 'appl_no is required'}), 400

    try:
        info = _fetch_cam_basic_info(appl_no)
        suffix = _fetch_revision_suffix(appl_no)
        history = _fetch_approval_history(appl_no, require_comment=False)
        related = _fetch_related_cams(appl_no, info.get('apless'))
        return jsonify({
            'main': {
                'appl_no_display': f"{appl_no}{suffix}",
                'customer_name': info.get('LESSEE_NM'),
                'cmo_name': info.get('USERNAME'),
                'history': history,
            },
            'related': related,
        })
    except Exception as e:
        current_app.logger.exception("Credit committee comments error")
        return jsonify({'error': str(e)}), 500