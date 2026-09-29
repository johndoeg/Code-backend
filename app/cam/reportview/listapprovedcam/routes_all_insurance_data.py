from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, send_file, session

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app import db
from app.auth.decorators import login_required

all_insurance_data_bp = Blueprint('all_insurance_data', __name__, url_prefix='/CAM/Approved')

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

@all_insurance_data_bp.route('/all-insurance-data/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) ORDER BY BRANCH_CD", []
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)

_TITLE_FONT = Font(name='Tahoma', size=12, bold=True)
_TITLE_ALIGN = Alignment(horizontal='center', vertical='center')

_DATA_FONT = Font(name='Tahoma', size=7)

_STYLE_1 = {'font': _DATA_FONT, 'border': _ALL_BORDERS}
_STYLE_2 = {'font': _DATA_FONT, 'border': _ALL_BORDERS, 'alignment': Alignment(horizontal='right')}
_STYLE_3 = {'font': _DATA_FONT, 'border': _ALL_BORDERS, 'alignment': Alignment(horizontal='left')}

def _apply_style(cell, style: dict):
    cell.font = style['font']
    cell.border = style['border']
    if 'alignment' in style:
        cell.alignment = style['alignment']

_COLUMNS = [
    ('A', None, 'No.', 'style1'),
    ('B', 'appl_no', 'CAM No.', 'style1'),
    ('C', 'create_date', 'Created Date', 'style2'),
    ('D', 'last_approved', 'Final Approved Date', 'style2'),
    ('E', 'last_status', 'Last Status', 'style1'),
    ('F', 'LESSEE_NM', 'Customer Name', 'style1'),
    ('G', 'Dealer', 'Dealer Name', 'style1'),
    ('H', 'BRAND_NM', 'Brand', 'style1'),
    ('I', 'tipe_mobil', 'Type', 'style1'),
    ('J', 'model', 'Model/Karoseri', 'style1'),
    ('K', 'WHEELS', 'Wheels', 'style1'),
    ('L', 'tahun', 'Year/Condition', 'style1'),
    ('M', 'tipe_equip', 'Type of Equipment', 'style1'),
    ('N', 'STNK_VALIDDT', 'STNK Valid Date', 'style1'),
    ('O', 'CHASIS', 'Chasis No.', 'style1'),
    ('P', 'ENGINE', 'Engine No.', 'style1'),
    ('Q', 'POLICENO', 'Police No.', 'style1'),
    ('R', 'USERUSAGE', 'User Usage', 'style1'),
    ('S', 'vehicle_purpose', 'Vehicle Purpose', 'style1'),
    ('T', 'tenor', 'Tenor', 'style1'),
    ('U', 'L_AMOUNT', 'Asset Value', 'style1'),
    ('V', 'Ins_co', 'Ins. Company', 'style1'),
    ('W', 'year1', 'Year 1 (%)', 'style3'),
    ('X', 'year2', 'Year 2 (%)', 'style1'),
    ('Y', 'year3', 'Year 3 (%)', 'style1'),
    ('Z', 'year4', 'Year 4 (%)', 'style1'),
    ('AA', 'year5', 'Year 5 (%)', 'style1'),
    ('AB', 'ins_amt_year1', 'Year 1 (Insurance Amount)', 'style1'),
    ('AC', 'ins_amt_year2', 'Year 2 (Insurance Amount)', 'style1'),
    ('AD', 'ins_amt_year3', 'Year 3 (Insurance Amount)', 'style1'),
    ('AE', 'ins_amt_year4', 'Year 4 (Insurance Amount)', 'style1'),
    ('AF', 'ins_amt_year5', 'Year 5 (Insurance Amount)', 'style1'),
    ('AG', 'tpl_amt_year1', 'Year 1 (TPL Amount)', 'style1'),
    ('AH', 'tpl_amt_year2', 'Year 2 (TPL Amount)', 'style1'),
    ('AI', 'tpl_amt_year3', 'Year 3 (TPL Amount)', 'style1'),
    ('AJ', 'tpl_amt_year4', 'Year 4 (TPL Amount)', 'style1'),
    ('AK', 'tpl_amt_year5', 'Year 5 (TPL Amount)', 'style1'),
    ('AL', 'tlo_year1', 'Year 1 (TLO CP)', 'style1'),
    ('AM', 'tlo_year2', 'Year 2 (TLO CP)', 'style1'),
    ('AN', 'tlo_year3', 'Year 3 (TLO CP)', 'style1'),
    ('AO', 'tlo_year4', 'Year 4 (TLO CP)', 'style1'),
    ('AP', 'tlo_year5', 'Year 5 (TLO CP)', 'style1'),
    ('AQ', 'clause_year1', 'Year 1 (CLAUSE)', 'style1'),
    ('AR', 'clause_year2', 'Year 2 (CLAUSE)', 'style1'),
    ('AS', 'clause_year3', 'Year 3 (CLAUSE)', 'style1'),
    ('AT', 'clause_year4', 'Year 4 (CLAUSE)', 'style1'),
    ('AU', 'clause_year5', 'Year 5 (CLAUSE)', 'style1'),
    ('AV', 'prer_year1', 'Year 1 (PREMIUM RECEIVE)', 'style1'),
    ('AW', 'prer_year2', 'Year 2 (PREMIUM RECEIVE)', 'style1'),
    ('AX', 'prer_year3', 'Year 3 (PREMIUM RECEIVE)', 'style1'),
    ('AY', 'prer_year4', 'Year 4 (PREMIUM RECEIVE)', 'style1'),
    ('AZ', 'prer_year5', 'Year 5 (PREMIUM RECEIVE)', 'style1'),
    ('BA', 'CMO', 'CMO Name', 'style1'),
    ('BB', 'lease_no', 'Contract No.', 'style1'),
    ('BC', 'execution', 'Disbursement Date', 'style1'),
    ('BD', 'BRANCH_NAME', 'Branch Name', 'style1'),
]

_STYLES_BY_KEY = {'style1': _STYLE_1, 'style2': _STYLE_2, 'style3': _STYLE_3}

def _build_workbook(search_by: str, date_from: date, date_to: date, branch_cd: str, username: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'AllInsuranceData'

    last_col = _COLUMNS[-1][0]
    
    company_name = current_app.config.get('COMPANY_NAME', '')
    ws.merge_cells(f'A2:BC2')
    ws['A2'] = company_name
    ws['A2'].alignment = Alignment(horizontal='left')

    ws['BD1'] = f"Date : {datetime.now().strftime('%d %b %Y')}"
    ws['BD1'].alignment = Alignment(horizontal='left')
    ws['BD2'] = f"Print By : {username}"
    ws['BD2'].alignment = Alignment(horizontal='left')

    filter_label = 'CAM Create Date ' if search_by == '1' else 'Disbursement Date '

    ws.merge_cells(f'A3:{last_col}3')
    ws['A3'] = 'ALL INSURANCE DATA'
    _apply_style_range(ws, f'A3:{last_col}3', font=_TITLE_FONT, alignment=_TITLE_ALIGN)

    ws.merge_cells(f'A4:{last_col}4')
    ws['A4'] = f"{filter_label}{_fmt_ddmmyyyy(date_from)} - {_fmt_ddmmyyyy(date_to)}"
    ws['A4'].alignment = Alignment(horizontal='center')

    header_row = 5
    for col, _field, label, _style_key in _COLUMNS:
        cell = ws[f'{col}{header_row}']
        cell.value = f' {label} '
        _apply_style(cell, {'font': _DATA_FONT, 'border': _ALL_BORDERS})

    ws.column_dimensions['A'].width = 5
    for col, _field, label, _style_key in _COLUMNS[1:]:
        ws.column_dimensions[col].width = max(len(label) + 2, 8)

    rows = _exec_rows(
        "EXEC USP_ALL_INSURANCE_DATA ?, ?, ?, ?",
        [search_by, date_from, date_to, branch_cd],
    )

    for i, row in enumerate(rows, start=1):
        excel_row = header_row + i
        for col, field, _label, style_key in _COLUMNS:
            cell = ws[f'{col}{excel_row}']
            cell.value = i if field is None else row.get(field)
            _apply_style(cell, _STYLES_BY_KEY[style_key])

    return wb

def _apply_style_range(ws, cell_range: str, font=None, alignment=None):
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if alignment is not None:
                cell.alignment = alignment

@all_insurance_data_bp.route('/all-insurance-data/export')
@login_required
def export_excel():
    search_by = request.args.get('search_by', '1').strip()
    date_from = _parse_date(request.args.get('date_from', ''))
    date_to = _parse_date(request.args.get('date_to', ''))
    branch_cd = request.args.get('branch_cd', '').strip()
    
    if date_from is None or date_to is None:
        return jsonify({'error': 'date_from and date_to are required (YYYY-MM-DD)'}), 400

    try:
        wb = _build_workbook(search_by, date_from, date_to, branch_cd, session.get("username"))
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='AllInsuranceData.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("All Insurance Data export error")
        return jsonify({'error': str(e)}), 500