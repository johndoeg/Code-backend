from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, send_file, session

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

all_cam_data_bp = Blueprint('all_cam_data', __name__, url_prefix='/CAM/Report')

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

@all_cam_data_bp.route('/all-cam-data/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) ORDER BY BRANCH_CD", []
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)

_TITLE_FONT = Font(name='Tahoma', size=12, bold=True)
_TITLE_ALIGN = Alignment(horizontal='center', vertical='center')

_DATA_FONT = Font(name='Tahoma', size=7)

_STYLE_1 = {'font': _DATA_FONT, 'border': _ALL_BORDERS}
_STYLE_2 = {'font': _DATA_FONT, 'border': _ALL_BORDERS, 'alignment': Alignment(horizontal='right')}
_STYLE_3 = {'font': _DATA_FONT, 'border': _ALL_BORDERS, 'alignment': Alignment(horizontal='left')}

def _apply_style(cell, style: dict):
    cell.font = style['font']
    cell.border = style['border']
    if 'alignment' in style:
        cell.alignment = style['alignment']

_COLUMNS = [
    ('A', None, 'No.', 'style1'),
    ('B', 'appl_no', 'CAM No.', 'style1'),
    ('C', 'create_date', 'Created Date', 'style2'),
    ('D', 'last_approved', 'Final Approved Date', 'style2'),
    ('E', 'last_approver', 'Last Approval Name', 'style1'),
    ('F', 'last_status', 'Last Status', 'style1'),
    ('G', 'LESSEE_NM', 'Customer Name', 'style1'),
    ('H', 'LESSEE_TP', 'Type', 'style1'),
    ('I', 'BRAND_NM', 'Brand', 'style1'),
    ('J', 'tipe_mobil', 'Type', 'style1'),
    ('K', 'model', 'Model/Karoseri', 'style1'),
    ('L', 'WHEELS', 'Wheels', 'style1'),
    ('M', 'TAHUN', 'Year', 'style1'),
    ('N', 'CONDITION', 'Condition', 'style1'),
    ('O', 'tipe_equip', 'Type of Equipment', 'style1'),
    ('P', 'STNK_VALIDDT', 'STNK Valid Date', 'style1'),
    ('Q', 'CHASIS', 'Chasis No.', 'style1'),
    ('R', 'ENGINE', 'Engine No.', 'style1'),
    ('S', 'POLICENO', 'Police No.', 'style1'),
    ('T', 'COLOUR', 'Color', 'style1'),
    ('U', 'USERUSAGE', 'User Usage', 'style1'),
    ('V', 'vehicle_purpose', 'Vehicle Purpose', 'style1'),
    ('W', 'cont_type', 'Contract Type', 'style3'),
    ('X', 'fin_purpose', 'Purpose of Finance', 'style1'),
    ('Y', 'tipe_pmt', 'Type of Payment', 'style1'),
    ('Z', 'fix_float', 'Fix/Float', 'style1'),
    ('AA', 'tenor', 'Tenor', 'style1'),
    ('AB', 'L_AMOUNT', 'Asset Value', 'style1'),
    ('AC', 'dp_persen', 'DP %', 'style1'),
    ('AD', 'dp', 'DP Amount', 'style1'),
    ('AE', 'ins_loan', 'Insurance Receive Onloan', 'style1'),
    ('AF', 'ins_cash', 'Insurance Receive Cash', 'style1'),
    ('AG', 'prov_loan', 'Provision Onloan', 'style1'),
    ('AH', 'prov_cash', 'Provision Cash', 'style1'),
    ('AI', 'adm_loan', 'Survey Fee Onloan', 'style1'),
    ('AJ', 'adm_cash', 'Survey Fee Cash', 'style1'),
    ('AK', 'FLATRATE', 'Flat Rate (%)', 'style1'),
    ('AL', 'DECLRATE', 'Eff Rate (%)', 'style1'),
    ('AM', 'RENTAL', 'Rental', 'style1'),
    ('AN', 'disb_dealer', 'Disbursement to Dealer', 'style1'),
    ('AO', 'tot_komisi', 'Total Commission', 'style1'),
    ('AP', 'TPL_AMT', 'TPL Amount (per year)', 'style1'),
    ('AQ', 'TLO_AR', 'TLO CP', 'style1'),
    ('AR', 'LTV', 'LTV', 'style1'),
    ('AS', 'IRR_Actual', 'IRR', 'style1'),
    ('AT', 'Ins_co', 'Ins. Company', 'style1'),
    ('AU', 'RedCode', 'Red Code', 'style1'),
    ('AV', 'CMO', 'CMO Name', 'style1'),
    ('AW', 'EXECUTION', 'Disbursement Date', 'style1'),
    ('AX', 'LEASE_NO', 'Contract No.', 'style1'),
    ('AY', 'UnIncome', 'Interest Income', 'style1'),
    ('AZ', 'ntf', 'Net Finance', 'style1'),
    ('BA', 'AlamatKTP', 'ID Address', 'style1'),
    ('BB', 'AlamatSurat', 'Mail Address', 'style1'),
    ('BC', 'HP', 'Mobile Phone', 'style1'),
    ('BD', 'Dealer', 'Dealer Name', 'style1'),
    ('BE', 'Dealer_Nickname', 'Dealer Nickname', 'style1'),
    ('BF', 'NamaBPKB', 'BPKB Name', 'style1'),
    ('BG', 'AlamatBPKB', 'BPKB Address', 'style1'),
    ('BH', 'BRANCH_NAME', 'Branch Name', 'style1'),
    ('BI', 'LESSEE_NO', 'Customer No.', 'style1'),
    ('BJ', 'IndusCode', 'Industrial Code', 'style1'),
    ('BK', 'BICust', 'BI Customer Type', 'style1'),
    ('BL', 'Rekomendasi', 'Not Recommended', 'style1'),
]

_STYLES_BY_KEY = {'style1': _STYLE_1, 'style2': _STYLE_2, 'style3': _STYLE_3}

def _build_workbook(search_by: str, date_from: date, date_to: date, branch_cd: str, username: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'AllCAMData'

    last_col = _COLUMNS[-1][0]
    second_last_col = _COLUMNS[-2][0]

    company_name = current_app.config.get('COMPANY_NAME', '')
    ws.merge_cells(f'A1:{second_last_col}1')
    ws.merge_cells(f'A2:{second_last_col}2')
    ws['A2'] = company_name
    ws['A2'].alignment = Alignment(horizontal='left')

    ws[f'{last_col}1'] = f"Date : {datetime.now().strftime('%d %b %Y')}"
    ws[f'{last_col}1'].alignment = Alignment(horizontal='left')
    ws[f'{last_col}2'] = f"Print By : {username}"
    ws[f'{last_col}2'].alignment = Alignment(horizontal='left')

    filter_label = 'CAM Create Date ' if search_by == '1' else 'Disbursement Date '

    ws.merge_cells(f'A3:{last_col}3')
    ws['A3'] = 'ALL CAM DATA'
    _apply_style_range(ws, f'A3:{last_col}3', font=_TITLE_FONT, alignment=_TITLE_ALIGN)

    ws.merge_cells(f'A4:{last_col}4')
    ws['A4'] = f"{filter_label}{_fmt_ddmmyyyy(date_from)} - {_fmt_ddmmyyyy(date_to)}"
    ws['A4'].alignment = Alignment(horizontal='center')

    header_row = 5
    for col, _field, label, _style_key in _COLUMNS:
        cell = ws[f'{col}{header_row}']
        cell.value = f' {label} '
        _apply_style(cell, {'font': _DATA_FONT, 'border': _ALL_BORDERS})

    ws.column_dimensions['A'].width = 5
    for col, _field, label, _style_key in _COLUMNS[1:]:
        ws.column_dimensions[col].width = max(len(label) + 2, 8)

    rows = _exec_rows(
        "EXEC USP_ALL_CAM_DATA ?, ?, ?, ?",
        [search_by, date_from, date_to, branch_cd],
    )

    for i, row in enumerate(rows, start=1):
        excel_row = header_row + i
        for col, field, _label, style_key in _COLUMNS:
            cell = ws[f'{col}{excel_row}']
            cell.value = i if field is None else row.get(field)
            _apply_style(cell, _STYLES_BY_KEY[style_key])

    return wb

def _apply_style_range(ws, cell_range: str, font=None, alignment=None):
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if alignment is not None:
                cell.alignment = alignment

@all_cam_data_bp.route('/all-cam-data/export')
@login_required
def export_excel():
    search_by = request.args.get('search_by', '1').strip()
    date_from = _parse_date(request.args.get('date_from', ''))
    date_to = _parse_date(request.args.get('date_to', ''))
    branch_cd = request.args.get('branch_cd', '').strip()

    if date_from is None or date_to is None:
        return jsonify({'error': 'date_from and date_to are required (YYYY-MM-DD)'}), 400

    try:
        wb = _build_workbook(search_by, date_from, date_to, branch_cd, session.get('username'))
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='AllCAMData.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("All CAM Data export error")
        return jsonify({'error': str(e)}), 500