from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

rejected_cam_bp = Blueprint('rejected_cam', __name__, url_prefix='/CAM/Approved')

PAGE_SIZE_DEFAULT = 20

_PASSENGER_MODELS = ['01', '02', '03']
_COMMERCIAL_MODELS = ['04', '05', '06', '07', '08', '09', '10']

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

@rejected_cam_bp.route('/rejected-cam/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD NOT IN ('100') ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@rejected_cam_bp.route('/rejected-cam/reject-detail-options')
@login_required
def reject_detail_options():
    rejected_by = request.args.get('rejected_by', '').strip()
    branch_cd = request.args.get('branch_cd', '').strip()

    try:
        if rejected_by in ('CMH', 'BM'):
            if branch_cd == '000':
                rows = _exec_rows(
                    "SELECT USERNAME, EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) "
                    "WHERE RECSTATUS = 'A' AND ACCESSCAM = ?",
                    [rejected_by],
                )
            else:
                rows = _exec_rows(
                    "SELECT USERNAME, EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) "
                    "WHERE RECSTATUS = 'A' AND ACCESSCAM = ? AND BRANCHCODE = ?",
                    [rejected_by, branch_cd],
                )
            options = [{'value': r.get('USERNAME'), 'label': r.get('USERNAME')} for r in rows]
        elif rejected_by == 'CRH':
            rows = _exec_rows(
                "SELECT USERNAME, EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) "
                "WHERE RECSTATUS = 'A' AND ACCESSCAM = ?",
                [rejected_by],
            )
            options = [{'value': r.get('USERNAME'), 'label': r.get('USERNAME')} for r in rows]
        elif rejected_by:
            options = [{'value': rejected_by, 'label': rejected_by}]
        else:
            options = []
        return jsonify(options)
    except Exception as e:
        current_app.logger.exception("Reject-detail options error")
        return jsonify({'error': str(e)}), 500

def _resolve_scope(requested_branch_cd: str) -> dict:
    access_level = session.get('accesslevel')
    emp_id = session.get('emp_id')
    own_branch_cd = session.get('branchcode')

    if access_level == 'CMO':
        return {
            'disabled': True,
            'search_branch': own_branch_cd,
            'filter_sql': 'b.CMO = ?',
            'filter_params': [emp_id],
        }
    if access_level == 'CMH':
        subordinates = _exec_rows(
            "SELECT EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) WHERE EMPLOYEE_ID = ? OR SUPERIOR = ?",
            [emp_id, emp_id],
        )
        ids = [r['EMPLOYEE_ID'] for r in subordinates] or [emp_id]
        placeholders = ','.join(['?'] * len(ids))
        return {
            'disabled': True,
            'search_branch': own_branch_cd,
            'filter_sql': f'b.CMO IN ({placeholders})',
            'filter_params': ids,
        }
    if access_level == 'BM':
        return {
            'disabled': True,
            'search_branch': own_branch_cd,
            'filter_sql': 'b.BRANCH_CD = ?',
            'filter_params': [own_branch_cd],
        }
    if own_branch_cd == '100':
        search_branch = requested_branch_cd
        return {
            'disabled': False,
            'search_branch': search_branch,
            'filter_sql': '(b.BRANCH_CD = ? OR ? = \'000\')',
            'filter_params': [search_branch, search_branch],
        }
    return {
        'disabled': True,
        'search_branch': own_branch_cd,
        'filter_sql': '(b.BRANCH_CD = ? OR ? = \'000\')',
        'filter_params': [own_branch_cd, own_branch_cd],
    }

@rejected_cam_bp.route('/rejected-cam/context')
@login_required
def get_context():
    resolved = _resolve_scope(request.args.get('branch_cd_search', '').strip())
    return jsonify({'disabled': resolved['disabled'], 'forced_branch_cd': resolved['search_branch']})

def _date_filter_sql(start_date: Optional[date], end_date: Optional[date]) -> tuple[str, list]:
    if start_date and end_date:
        return "AND CONVERT(date, a.APPR_DATE) BETWEEN ? AND ?", [start_date, end_date]
    return "", []

def _fetch_detail_rows(
    start_date: Optional[date], end_date: Optional[date], scope: dict, reject_detail: str,
    page: Optional[int] = None, page_size: int = PAGE_SIZE_DEFAULT,
) -> tuple[list[dict], int]:
    date_sql, date_params = _date_filter_sql(start_date, end_date)

    base_sql = f"""
        SELECT
            a.APPR_DATE, a.APPR_ID, a.APPL_NO,
            b.HAS_GUARANTOR, b.Cons_Leas, b.TENOR, b.APLESS,
            c.LESSEE_NM, c.LESSEE_TP,
            d.SUPP, LTRIM(RTRIM(d.[CONDITION])) AS CONDITION,
            v.Net_Finance,
            ROW_NUMBER() OVER (PARTITION BY a.APPL_NO ORDER BY a.APPR_DATE DESC) AS rn
        FROM MCMOCAMStatus a WITH (NOLOCK)
            LEFT JOIN CAFCONT b WITH (NOLOCK) ON a.APPL_NO = b.APPL_NO
            LEFT JOIN CAFLESS c WITH (NOLOCK) ON b.APLESS = c.APLESS
            LEFT JOIN CAFEQUIP d WITH (NOLOCK) ON a.APPL_NO = d.APPL_NO
            LEFT JOIN V_NetFinance v WITH (NOLOCK) ON b.APPL_NO = v.APPL_NO
        WHERE a.APPR_STS = 'R'
            AND a.APPR_ID = ?
            AND {scope['filter_sql']}
            {date_sql}
    """
    params = [reject_detail] + scope['filter_params'] + date_params

    count_sql = f"SELECT COUNT(*) AS cnt FROM ({base_sql}) ranked WHERE rn = 1"
    count_rows = _exec_rows(count_sql, params)
    total_count = count_rows[0]['cnt'] if count_rows else 0

    if page is None:
        list_sql = f"SELECT * FROM ({base_sql}) ranked WHERE rn = 1 ORDER BY APPL_NO"
        rows = _exec_rows(list_sql, params)
    else:
        offset = max(page - 1, 0) * page_size
        list_sql = (
            f"SELECT * FROM ({base_sql}) ranked WHERE rn = 1 ORDER BY APPL_NO "
            f"OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
        )
        rows = _exec_rows(list_sql, params + [offset, page_size])

    return rows, total_count

def _row_to_detail(row: dict) -> dict:
    appr_date = row.get('APPR_DATE')
    fin_type = 1 if row.get('Cons_Leas') == 'C' else 2 if row.get('Cons_Leas') == 'L' else None
    has_guarantor = row.get('HAS_GUARANTOR')
    guarantor = 2 if has_guarantor == '1' else 1
    ind_cor = 2 if row.get('LESSEE_TP') == 'PT' else 1
    c2c = 1 if row.get('SUPP') == 999999 else 0
    condition = (row.get('CONDITION') or '').strip()
    new_car = 2 if condition == 'NEW' else 1

    return {
        'appl_no': row.get('APPL_NO'),
        'cons_leas': row.get('Cons_Leas'),
        'rejected_date': appr_date.strftime('%d-%m-%Y') if appr_date else '',
        'customer_name': row.get('LESSEE_NM') or '-',
        'net_finance': _num(row.get('Net_Finance')),
        'rejected_by': row.get('APPR_ID'),
        'edit_link_params': {
            'id1': fin_type, 'id2': ind_cor, 'id3': 1, 'id4': guarantor, 'id5': new_car,
            'id6': 'Edit', 'id7': row.get('TENOR'), 'id8': c2c,
            'apless': row.get('APLESS'), 'applno': row.get('APPL_NO'), 'rjt': 1,
        },
    }

def _fetch_grand_total(start_date: Optional[date], end_date: Optional[date], scope: dict, reject_detail: str) -> dict:
   
    date_sql, date_params = _date_filter_sql(start_date, end_date)

    def _bucket(models: list[str], condition: str) -> dict:
        placeholders = ','.join(['?'] * len(models))
        sql = f"""
            SELECT ISNULL(SUM(q.Net_Finance), 0) AS tot_fin, ISNULL(COUNT(q.APPL_NO), 0) AS tot_unit
            FROM (
                SELECT DISTINCT a.APPL_NO, v.Net_Finance
                FROM MCMOCAMStatus a WITH (NOLOCK)
                    LEFT JOIN CAFCONT b WITH (NOLOCK) ON a.APPL_NO = b.APPL_NO
                    LEFT JOIN CAFLESS c WITH (NOLOCK) ON b.APLESS = c.APLESS
                    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
                    LEFT JOIN V_NetFinance v WITH (NOLOCK) ON b.APPL_NO = v.APPL_NO
                WHERE a.APPL_NO NOT IN (
                    SELECT OldAppl_No FROM MCMORevisi WITH (NOLOCK) WHERE NewAppl_No IS NOT NULL
                )
                    AND a.APPR_STS = 'R'
                    AND a.APPR_ID = ?
                    AND {scope['filter_sql']}
                    {date_sql}
                    AND g.MODEL IN ({placeholders})
                    AND g.CONDITION = ?
            ) q
        """
        params = [reject_detail] + scope['filter_params'] + date_params + models + [condition]
        rows = _exec_rows(sql, params)
        r = rows[0] if rows else {'tot_fin': 0, 'tot_unit': 0}
        return {'net_finance': _num(r.get('tot_fin')), 'units': int(_num(r.get('tot_unit')))}

    pass_new = _bucket(_PASSENGER_MODELS, 'NEW')
    pass_used = _bucket(_PASSENGER_MODELS, 'USED')
    comm_new = _bucket(_COMMERCIAL_MODELS, 'NEW')
    comm_used = _bucket(_COMMERCIAL_MODELS, 'USED')

    grand = {
        'net_finance': pass_new['net_finance'] + pass_used['net_finance'] + comm_new['net_finance'] + comm_used['net_finance'],
        'units': pass_new['units'] + pass_used['units'] + comm_new['units'] + comm_used['units'],
    }

    return {
        'passenger_new': pass_new,
        'passenger_used': pass_used,
        'commercial_new': comm_new,
        'commercial_used': comm_used,
        'grand_total': grand,
    }

def _period_label(start_date: Optional[date], end_date: Optional[date], rows: list[dict]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    if rows:
        first = rows[0].get('APPR_DATE')
        last = rows[-1].get('APPR_DATE')
        return f"{first.strftime('%d %B %Y') if first else ''} - {last.strftime('%d %B %Y') if last else ''}"
    return ''

@rejected_cam_bp.route('/rejected-cam/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()
    reject_detail = request.args.get('reject_detail', '').strip()
    page = int(request.args.get('page', 1) or 1)
    page_size = int(request.args.get('page_size', PAGE_SIZE_DEFAULT) or PAGE_SIZE_DEFAULT)

    if not reject_detail:
        return jsonify({'error': 'reject_detail is required'}), 400

    scope = _resolve_scope(requested_branch)
    try:
        rows, total_count = _fetch_detail_rows(start_date, end_date, scope, reject_detail, page, page_size)
        grand_total = _fetch_grand_total(start_date, end_date, scope, reject_detail)
        return jsonify({
            'data': [_row_to_detail(r) for r in rows],
            'total_count': total_count,
            'page': page,
            'page_size': page_size,
            'grand_total': grand_total,
            'disabled': scope['disabled'],
            'search_branch': scope['search_branch'],
        })
    except Exception as e:
        current_app.logger.exception("Rejected CAM list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

def _build_workbook(start_date, end_date, scope, reject_detail, rows, grand_total) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'ListRejectedCam'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List Rejected CAM'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {_branch_name(scope['search_branch'])}"

    ws.merge_cells('A4:D4')
    ws['A4'] = _period_label(start_date, end_date, rows)

    header_row = 6
    headers = ['Rejected Date', 'CAM No.', 'Customer Name', 'Net Finance', 'Rejected By']
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.border = _ALL_BORDERS
        cell.alignment = _CENTER

    excel_row = header_row + 1
    if not rows:
        ws.merge_cells(f'A{excel_row}:E{excel_row}')
        cell = ws[f'A{excel_row}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        excel_row += 1
    else:
        for r in rows:
            detail = _row_to_detail(r)
            values = [
                detail['rejected_date'],
                f"{detail['cons_leas']} - {detail['appl_no']}",
                detail['customer_name'],
                detail['net_finance'],
                detail['rejected_by'],
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = _ALL_BORDERS
                if col_idx == 4:
                    cell.alignment = _RIGHT
            excel_row += 1

    excel_row += 1
    ws.merge_cells(f'A{excel_row}:D{excel_row}')
    ws[f'A{excel_row}'] = 'Grand Total* (based on latest revised CAM, if any)'
    ws[f'A{excel_row}'].font = _HEADER_FONT
    excel_row += 1

    for col_idx, label in enumerate(['Type', 'Condition', 'Net Finance', 'Units'], start=1):
        cell = ws.cell(row=excel_row, column=col_idx, value=label)
        cell.font = _HEADER_FONT
    excel_row += 1

    for type_label, bucket in [
        ('Passenger', 'passenger_new'), ('Passenger', 'passenger_used'),
        ('Commercial', 'commercial_new'), ('Commercial', 'commercial_used'),
    ]:
        condition_label = 'New' if 'new' in bucket else 'Used'
        b = grand_total[bucket]
        ws.cell(row=excel_row, column=1, value=type_label)
        ws.cell(row=excel_row, column=2, value=condition_label)
        ws.cell(row=excel_row, column=3, value=b['net_finance']).alignment = _RIGHT
        ws.cell(row=excel_row, column=4, value=b['units']).alignment = _RIGHT
        excel_row += 1

    ws.merge_cells(f'A{excel_row}:B{excel_row}')
    ws[f'A{excel_row}'] = 'Grand Total'
    ws[f'A{excel_row}'].font = _HEADER_FONT
    ws.cell(row=excel_row, column=3, value=grand_total['grand_total']['net_finance']).alignment = _RIGHT
    ws.cell(row=excel_row, column=4, value=grand_total['grand_total']['units']).alignment = _RIGHT

    for col, width in zip('ABCDE', [15, 20, 30, 15, 20]):
        ws.column_dimensions[col].width = width

    return wb

@rejected_cam_bp.route('/rejected-cam/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()
    reject_detail = request.args.get('reject_detail', '').strip()

    if not reject_detail:
        return jsonify({'error': 'reject_detail is required'}), 400

    scope = _resolve_scope(requested_branch)
    try:
        rows, _total = _fetch_detail_rows(start_date, end_date, scope, reject_detail)
        grand_total = _fetch_grand_total(start_date, end_date, scope, reject_detail)
        wb = _build_workbook(start_date, end_date, scope, reject_detail, rows, grand_total)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ListRejectedCam.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("Rejected CAM export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>List Rejected CAM</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
@media print {{
    .report-header {{ display: none; }}
}}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();" onafterprint="window.close();">
    <div class="report-header">
        <table cellpadding="2" cellspacing="0" width="100%">
            <tr>
                <td width="25%" nowrap="nowrap">{company}</td>
                <td width="50%"></td>
                <td width="10%"></td>
                <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
            </tr>
            <tr>
                <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List Rejected CAM</strong></td>
            </tr>
            <tr><td colspan="4">Branch : {branch_name}</td></tr>
            <tr><td colspan="4">{period}</td></tr>
        </table>
    </div>
    <table style="border:thin solid #000" cellpadding="2" cellspacing="0" width="100%">
        <tr align="center">
            <td style="border:thin solid #000" width="15%">Rejected Date</td>
            <td style="border:thin solid #000" width="15%">CAM No.</td>
            <td style="border:thin solid #000" width="30%">Customer Name</td>
            <td style="border:thin solid #000" width="15%">Net Finance</td>
            <td style="border:thin solid #000" width="25%">Rejected By</td>
        </tr>
    {rows_html}
    </table>
    <br>
    <table style="border:thin solid #000" cellpadding="2" cellspacing="0" width="40%">
        <tr><td colspan="4"><u><strong>Grand Total* ({period}):</strong></u></td></tr>
        <tr style="background-color:#E5E5E5">
        <td width="10%"><strong>Type</strong></td>
        <td width="10%"><strong>Condition</strong></td>
        <td width="15%"><strong>Net Finance</strong></td>
        <td width="15%"><strong>Units</strong></td>
    </tr>
    {grand_rows_html}
    <tr>
        <td colspan="2" align="center"><strong>Grand Total</strong></td>
        <td align="right">{grand_fin}</td>
        <td align="right">{grand_units}</td>
    </tr>
    <tr><td colspan="4"><i>*Grand Total Net Finance and Unit based on latest revised CAM (if any)</i></td></tr>
    </table>
</body>
</html>"""

def _render_detail_rows_html(rows: list[dict]) -> str:
    if not rows:
        return '<tr><td colspan="5" align="center"><font color="red"><b>Data Not Found</b></font></td></tr>'
    parts = []
    for r in rows:
        d = _row_to_detail(r)
        parts.append(f"""
            <tr>
                <td style="border:thin solid #000">{d['rejected_date']}</td>
                <td style="border:thin solid #000">{d['cons_leas']} - {d['appl_no']}</td>
                <td style="border:thin solid #000">{d['customer_name']}</td>
                <td style="border:thin solid #000" align="right">{d['net_finance']:,.0f}</td>
                <td style="border:thin solid #000">{d['rejected_by']}</td>
            </tr>""")
    return ''.join(parts)

def _render_grand_rows_html(grand_total: dict) -> str:
    parts = []
    for type_label, bucket in [
        ('Passenger', 'passenger_new'), ('Passenger', 'passenger_used'),
        ('Commercial', 'commercial_new'), ('Commercial', 'commercial_used'),
    ]:
        condition_label = 'New' if 'new' in bucket else 'Used'
        b = grand_total[bucket]
        parts.append(f"""
            <tr>
                <td>{type_label}</td>
                <td>{condition_label}</td>
                <td align="right">{b['net_finance']:,.0f}</td>
                <td align="right">{b['units']}</td>
            </tr>""")
    return ''.join(parts)

@rejected_cam_bp.route('/rejected-cam/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()
    reject_detail = request.args.get('reject_detail', '').strip()

    if not reject_detail:
        return jsonify({'error': 'reject_detail is required'}), 400

    scope = _resolve_scope(requested_branch)
    try:
        rows, _total = _fetch_detail_rows(start_date, end_date, scope, reject_detail)
        grand_total = _fetch_grand_total(start_date, end_date, scope, reject_detail)
        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_branch_name(scope['search_branch']),
            period=_period_label(start_date, end_date, rows),
            rows_html=_render_detail_rows_html(rows),
            grand_rows_html=_render_grand_rows_html(grand_total),
            grand_fin=f"{grand_total['grand_total']['net_finance']:,.0f}",
            grand_units=grand_total['grand_total']['units'],
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("Rejected CAM print error")
        return jsonify({'error': str(e)}), 500