from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

cam_by_app_date_bp = Blueprint('cam_by_app_date', __name__, url_prefix='/CAM/Approved')

PAGE_SIZE = 100

_PASSENGER_MODELS = ['01', '02', '03']
_COMMERCIAL_MODELS = ['04', '05', '06', '07', '08', '09', '10']

_STATUS_LABELS = {
    'A': 'Approved',
    'F': 'Approved & Send to Admin',
    'R': 'Reject',
    'N': 'Send for Approval',
}

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _exec_multi_resultset(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    all_rows: list[dict] = []
    try:
        cursor.execute(sql, params)
        while True:
            if cursor.description is not None:
                columns = [col[0] for col in cursor.description]
                all_rows += [dict(zip(columns, row)) for row in cursor.fetchall()]
            if not cursor.nextset():
                break
        return all_rows
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''


def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _resolve_scope(requested_branch_cd: str) -> dict:
    access_level = session.get('accesslevel')
    emp_id = session.get('emp_id')
    own_branch_cd = session.get('branchcode')

    if access_level == 'CMO':
        return {
            'disabled': True,
            'branch_cd_search': own_branch_cd,
            'cmo_filter': [emp_id],
            'show_all_branches': False,
        }

    if access_level == 'CMH':
        rows = _exec_rows(
            "SELECT EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) WHERE EMPLOYEE_ID = ? OR SUPERIOR = ?",
            [emp_id, emp_id],
        )
        cmo_ids = [r['EMPLOYEE_ID'] for r in rows] or [emp_id]
        return {
            'disabled': True,
            'branch_cd_search': own_branch_cd,
            'cmo_filter': cmo_ids,
            'show_all_branches': False,
        }

    if access_level == 'BM':
        return {
            'disabled': True,
            'branch_cd_search': own_branch_cd,
            'cmo_filter': None,
            'show_all_branches': False,
        }

    branch_cd_search = requested_branch_cd or own_branch_cd or ''
    return {
        'disabled': False,
        'branch_cd_search': branch_cd_search,
        'cmo_filter': None,
        'show_all_branches': branch_cd_search == '000',
    }

def _build_common_filter(scope: dict, start_date: Optional[date], end_date: Optional[date],
                          lessee_nm: str) -> tuple[str, list]:
    clauses = ["c.AdmSts = '1'", "k.APPR_ID_From IS NOT NULL"]
    params: list = []

    if start_date and end_date:
        clauses.append("c.AdmDate >= ? AND c.AdmDate <= ?")
        params.append(datetime.combine(start_date, datetime.min.time()))
        params.append(datetime.combine(end_date, datetime.max.time()))

    if lessee_nm:
        clauses.append("b.LESSEE_NM LIKE ?")
        params.append(f"%{lessee_nm}%")

    if scope['cmo_filter'] is not None:
        placeholders = ','.join('?' for _ in scope['cmo_filter'])
        clauses.append(f"a.CMO IN ({placeholders})")
        params.extend(scope['cmo_filter'])

    if not scope['show_all_branches'] and scope['branch_cd_search']:
        clauses.append("a.BRANCH_CD = ?")
        params.append(scope['branch_cd_search'])

    return " AND ".join(clauses), params

_LIST_SELECT = """
SELECT k.APPR_ID_From, a.APPL_NO, a.Cons_Leas, b.LESSEE_NM, a.TENOR, c.AdmDate,
       e.USERNAME, h.BRAND_NM, i.TYPE_NM, v.Net_Finance, l.IRR_Actual AS IRR,
       a.APLESS, a.HAS_GUARANTOR, b.LESSEE_TP, g.SUPP, LTRIM(RTRIM(g.[CONDITION])) AS CONDITION
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    LEFT JOIN MCMOCAMCOMPLETELIST c WITH (NOLOCK) ON a.APPL_NO = c.APPL_NO
    LEFT JOIN APP_USRMST e WITH (NOLOCK) ON a.CMO = e.EMPLOYEE_ID
    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
    LEFT JOIN UMFBRAND h WITH (NOLOCK) ON g.BRAND = h.BRAND
    LEFT JOIN UMFTYPE i WITH (NOLOCK) ON g.BRAND = i.BRAND AND g.TYPE = i.TYPE
    LEFT JOIN (
        SELECT APPR_ID_From, APPL_NO FROM mcmocamstatus WITH (NOLOCK)
        WHERE appr_sts = 'E' GROUP BY APPR_ID_From, APPL_NO, APPR_STS
    ) k ON a.APPL_NO = k.APPL_NO
    LEFT JOIN V_NetFinance v WITH (NOLOCK) ON c.APPL_NO = v.APPL_NO
    LEFT JOIN CAFCONT_IRR l WITH (NOLOCK) ON a.APPL_NO = l.APPL_NO
"""

def _fetch_count(where_sql: str, params: list) -> int:
    sql = f"""
    SELECT COUNT(*) AS cnt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        LEFT JOIN MCMOCAMCOMPLETELIST c WITH (NOLOCK) ON a.APPL_NO = c.APPL_NO
        LEFT JOIN (
            SELECT APPR_ID_From, APPL_NO FROM mcmocamstatus WITH (NOLOCK)
            WHERE appr_sts = 'E' GROUP BY APPR_ID_From, APPL_NO, APPR_STS
        ) k ON a.APPL_NO = k.APPL_NO
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    return int(rows[0]['cnt']) if rows else 0

def _fetch_page(where_sql: str, params: list, page: int, page_size: int) -> list[dict]:
    offset = (page - 1) * page_size
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY c.AdmDate OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
    return _exec_rows(sql, params + [offset, page_size])

def _fetch_all(where_sql: str, params: list) -> list[dict]:
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY c.AdmDate"
    return _exec_rows(sql, params)

def _fetch_date_range(where_sql: str, params: list) -> tuple[Optional[date], Optional[date]]:
    sql = f"""
    SELECT MIN(c.AdmDate) AS min_dt, MAX(c.AdmDate) AS max_dt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        LEFT JOIN MCMOCAMCOMPLETELIST c WITH (NOLOCK) ON a.APPL_NO = c.APPL_NO
        LEFT JOIN (
            SELECT APPR_ID_From, APPL_NO FROM mcmocamstatus WITH (NOLOCK)
            WHERE appr_sts = 'E' GROUP BY APPR_ID_From, APPL_NO, APPR_STS
        ) k ON a.APPL_NO = k.APPL_NO
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    if not rows:
        return None, None
    min_dt, max_dt = rows[0].get('min_dt'), rows[0].get('max_dt')
    return (min_dt.date() if min_dt else None), (max_dt.date() if max_dt else None)

def _fetch_revisions(appl_nos: list[str]) -> dict[str, str]:
    if not appl_nos:
        return {}
    placeholders = ','.join('?' for _ in appl_nos)
    rows = _exec_rows(
        f"SELECT NEWAPPL_NO, REVISI FROM MCMORevisi WITH (NOLOCK) WHERE NEWAPPL_NO IN ({placeholders})",
        appl_nos,
    )
    return {r['NEWAPPL_NO']: f"R{r['REVISI']}" for r in rows}

def _edit_link_params(r: dict) -> dict:
    cons_leas = r.get('Cons_Leas')
    fin_type = 1 if cons_leas == 'C' else 2 if cons_leas == 'L' else None
    guarantor = 2 if r.get('HAS_GUARANTOR') == '1' else 1
    ind_cor = 2 if r.get('LESSEE_TP') == 'PT' else 1
    c2c = 1 if str(r.get('SUPP')) == '999999' else 0
    condition = (r.get('CONDITION') or '').strip()
    new_car = 2 if condition == 'NEW' else 1

    return {
        'id1': fin_type, 'id2': ind_cor, 'id3': 1, 'id4': guarantor, 'id5': new_car,
        'id6': 'Edit', 'id7': r.get('TENOR'), 'id8': c2c,
        'apless': r.get('APLESS'), 'applno': r.get('APPL_NO'),
    }

def _row_to_detail(r: dict, no: int, revision: str) -> dict:
    adm_date = r.get('AdmDate')
    brand = (r.get('BRAND_NM') or '').strip()
    type_nm = (r.get('TYPE_NM') or '').strip()
    return {
        'no': no,
        'cam_no': f"{r.get('Cons_Leas')} - {r.get('APPL_NO')}{revision}",
        'appl_no': r.get('APPL_NO'),
        'customer_name': r.get('LESSEE_NM') or '-',
        'brand_type': f"{brand} / {type_nm}" if (brand or type_nm) else '-',
        'net_finance': _num(r.get('Net_Finance')),
        'irr': round(_num(r.get('IRR')), 2),
        'tenor': int(_num(r.get('TENOR'))),
        'approval_date': adm_date.strftime('%d-%m-%Y') if adm_date else '',
        'cmo': r.get('USERNAME') or '-',
        'approved_by': r.get('APPR_ID_From') or '-',
        'is_revision': bool(revision),
        'edit_link_params': _edit_link_params(r),
    }

_GRANDTOTAL_BASE = """
SELECT ISNULL(SUM(v.Net_Finance), 0) AS tot_fin,
       ISNULL(COUNT(a.APPL_NO), 0) AS tot_unit,
       ISNULL(SUM(a.TENOR), 0) AS tot_tenor,
       ISNULL(SUM(d.tot_income), 0) AS tot_income
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    LEFT JOIN MCMOCAMCOMPLETELIST c WITH (NOLOCK) ON a.APPL_NO = c.APPL_NO
    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
    LEFT JOIN (
        SELECT APPR_ID_From, APPL_NO FROM mcmocamstatus WITH (NOLOCK)
        WHERE appr_sts = 'E' GROUP BY APPR_ID_From, APPL_NO, APPR_STS
    ) k ON a.APPL_NO = k.APPL_NO
    LEFT JOIN V_NetFinance v WITH (NOLOCK) ON c.APPL_NO = v.APPL_NO
    LEFT JOIN V_IncomeIRR d WITH (NOLOCK) ON a.APPL_NO = d.APPL_NO
WHERE a.APPL_NO NOT IN (SELECT oldappl_no FROM mcmorevisi WITH (NOLOCK) WHERE newappl_no IS NOT NULL)
    AND {filter}
    AND g.MODEL IN ({models})
    AND g.CONDITION = ?
    AND a.cons_leas = ?
"""

def _fetch_segment_total(where_sql: str, params: list, models: list[str], condition: str, cons_leas: str) -> dict:
    placeholders = ','.join('?' for _ in models)
    sql = _GRANDTOTAL_BASE.format(filter=where_sql, models=placeholders)
    rows = _exec_rows(sql, params + list(models) + [condition, cons_leas])
    row = rows[0] if rows else {}
    return {
        'fin': _num(row.get('tot_fin')),
        'unit': int(_num(row.get('tot_unit'))),
        'tenor': _num(row.get('tot_tenor')),
        'irr_income': _num(row.get('tot_income')),
    }

def _compute_segment_metrics(seg: dict) -> dict:
    fin, unit, tenor, irr_income = seg['fin'], seg['unit'], seg['tenor'], seg['irr_income']
    return {
        'net_finance': fin,
        'units': unit,
        'avg_irr': round(irr_income / fin * 100, 2) if fin > 0 else 0,
        'avg_tenor': round(tenor / unit) if unit > 0 else 0,
        'avg_net_finance': round(fin / unit) if unit > 0 else 0,
    }

def _sum_segments(group: dict) -> dict:
    return {
        'fin': sum(s['fin'] for s in group.values()),
        'unit': sum(s['unit'] for s in group.values()),
        'tenor': sum(s['tenor'] for s in group.values()),
        'irr_income': sum(s['irr_income'] for s in group.values()),
    }

def _fetch_breakdown(where_sql: str, params: list, cons_leas: str) -> dict:
    return {
        'passenger_new': _fetch_segment_total(where_sql, params, _PASSENGER_MODELS, 'NEW', cons_leas),
        'passenger_used': _fetch_segment_total(where_sql, params, _PASSENGER_MODELS, 'USED', cons_leas),
        'commercial_new': _fetch_segment_total(where_sql, params, _COMMERCIAL_MODELS, 'NEW', cons_leas),
        'commercial_used': _fetch_segment_total(where_sql, params, _COMMERCIAL_MODELS, 'USED', cons_leas),
    }

def _fetch_grand_totals(where_sql: str, params: list) -> dict:
    sl = _fetch_breakdown(where_sql, params, 'S')
    if_ = _fetch_breakdown(where_sql, params, 'I')
    sl_total = _sum_segments(sl)
    if_total = _sum_segments(if_)
    grand = {
        'fin': sl_total['fin'] + if_total['fin'],
        'unit': sl_total['unit'] + if_total['unit'],
        'tenor': sl_total['tenor'] + if_total['tenor'],
        'irr_income': sl_total['irr_income'] + if_total['irr_income'],
    }

    def _finalize(group: dict, total: dict) -> dict:
        return {**{k: _compute_segment_metrics(v) for k, v in group.items()}, 'total': _compute_segment_metrics(total)}

    return {
        'sl': _finalize(sl, sl_total),
        'if': _finalize(if_, if_total),
        'grand_total': _compute_segment_metrics(grand),
    }

@cam_by_app_date_bp.route('/cam-by-app-date/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@cam_by_app_date_bp.route('/cam-by-app-date/context')
@login_required
def get_context():
    scope = _resolve_scope(request.args.get('branch_cd', '').strip())
    return jsonify({'disabled': scope['disabled'], 'forced_branch_cd': scope['branch_cd_search']})

@cam_by_app_date_bp.route('/cam-by-app-date/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch_cd', '').strip()
    try:
        page = max(int(request.args.get('page', 1) or 1), 1)
    except ValueError:
        page = 1

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        total_count = _fetch_count(where_sql, params)
        total_pages = max((total_count + PAGE_SIZE - 1) // PAGE_SIZE, 1)
        page = min(page, total_pages)

        raw_rows = _fetch_page(where_sql, params, page, PAGE_SIZE)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)

        start_no = (page - 1) * PAGE_SIZE + 1
        details = [_row_to_detail(r, start_no + i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows)]

        grand_totals = _fetch_grand_totals(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        return jsonify({
            'rows': details,
            'total_count': total_count,
            'page': page,
            'page_size': PAGE_SIZE,
            'total_pages': total_pages,
            'grand_totals': grand_totals,
            'period_label': period,
            'disabled': scope['disabled'],
            'search_branch': scope['branch_cd_search'],
        })
    except Exception as e:
        current_app.logger.exception("CAM by Approval Date list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=16)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

_LIST_COLUMNS = ['No.', 'CAM No.', 'Customer Name', 'Brand/Type', 'Net Finance', '%IRR', 'Tenor (Month)', 'Approval Date', 'CMO', 'Approved by']
_GT_COLUMNS = ['Contract Type', 'Type', 'Condition', 'Net Finance', 'Units', 'Avg. IRR', 'Avg. Tenor', 'Avg. Net Finance']

_SEGMENT_LABELS = [
    ('passenger_new', 'Passenger', 'New'),
    ('passenger_used', 'Passenger', 'Used'),
    ('commercial_new', 'Commercial', 'New'),
    ('commercial_used', 'Commercial', 'Used'),
]

def _build_workbook(period_label: str, branch_cd: str, lessee_nm: str, rows: list[dict], grand_totals: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'CamByApprovalDate'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List CAM by Approval Date'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {_branch_name(branch_cd) or 'All Branches'}"

    ws.merge_cells('A4:D4')
    ws['A4'] = period_label.replace('Period : ', 'Period : ') or (f"Name : {lessee_nm}" if lessee_nm else '')

    row_num = 6
    for col_idx, label in enumerate(_LIST_COLUMNS, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=label)
        cell.border = _ALL_BORDERS
        cell.alignment = _CENTER
        cell.font = _HEADER_FONT
    row_num += 1

    if not rows:
        ws.merge_cells(f'A{row_num}:J{row_num}')
        cell = ws[f'A{row_num}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        row_num += 1
    else:
        for d in rows:
            values = [d['no'], d['cam_no'], d['customer_name'], d['brand_type'], d['net_finance'], d['irr'],
                      d['tenor'], d['approval_date'], d['cmo'], d['approved_by']]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = _ALL_BORDERS
                if col_idx in (5, 6):
                    cell.alignment = _RIGHT
            row_num += 1

    row_num += 2
    ws.merge_cells(f'A{row_num}:H{row_num}')
    ws[f'A{row_num}'] = f"Grand Total* ({period_label.replace('Period : ', '') or 'All Dates'})"
    ws[f'A{row_num}'].font = _HEADER_FONT
    row_num += 1

    for col_idx, label in enumerate(_GT_COLUMNS, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=label)
        cell.border = _ALL_BORDERS
        cell.font = _HEADER_FONT
    row_num += 1

    def _write_group(label: str, group: dict):
        nonlocal row_num
        start_row = row_num
        for key, type_label, cond_label in _SEGMENT_LABELS:
            seg = group[key]
            values = [type_label, cond_label, seg['net_finance'], seg['units'], seg['avg_irr'],
                      seg['avg_tenor'], seg['avg_net_finance']]
            for offset, val in enumerate(values, start=2):
                cell = ws.cell(row=row_num, column=offset, value=val)
                cell.border = _ALL_BORDERS
                if offset >= 4:
                    cell.alignment = _RIGHT
            row_num += 1
        ws.merge_cells(f'A{start_row}:A{row_num - 1}')
        label_cell = ws.cell(row=start_row, column=1, value=label)
        label_cell.border = _ALL_BORDERS
        label_cell.alignment = _CENTER

        total = group['total']
        ws.merge_cells(f'A{row_num}:B{row_num}')
        total_label_cell = ws.cell(row=row_num, column=1, value=f'Grand Total {label}')
        total_label_cell.font = _HEADER_FONT
        total_label_cell.alignment = _RIGHT
        for col_idx, val in ((4, total['net_finance']), (5, total['units']), (6, total['avg_irr']),
                              (7, total['avg_tenor']), (8, total['avg_net_finance'])):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = _ALL_BORDERS
            cell.alignment = _RIGHT
        row_num += 1

    _write_group('SL', grand_totals['sl'])
    _write_group('IF', grand_totals['if'])

    gt = grand_totals['grand_total']
    ws.merge_cells(f'A{row_num}:C{row_num}')
    gt_label_cell = ws.cell(row=row_num, column=1, value='Grand Total')
    gt_label_cell.font = _HEADER_FONT
    gt_label_cell.alignment = _RIGHT
    for col_idx, val in ((4, gt['net_finance']), (5, gt['units']), (6, gt['avg_irr']),
                          (7, gt['avg_tenor']), (8, gt['avg_net_finance'])):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.border = _ALL_BORDERS
        cell.alignment = _RIGHT
    row_num += 2

    ws.merge_cells(f'A{row_num}:D{row_num}')
    ws[f'A{row_num}'] = '*Grand Total Net Finance and Unit based on latest revised CAM (if any)'
    ws[f'A{row_num}'].font = Font(italic=True)

    ws.column_dimensions['A'].width = 14
    for col in 'BCDEFGHIJ':
        ws.column_dimensions[col].width = 16

    return wb

@cam_by_app_date_bp.route('/cam-by-app-date/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch_cd', '').strip()

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_totals = _fetch_grand_totals(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        wb = _build_workbook(period, scope['branch_cd_search'], lessee_nm, rows, grand_totals)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ListCamByApprovalDate.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("CAM by Approval Date export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{company} - List CAM by Approval Date</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
        </tr>
        <tr>
            <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List CAM by Approval Date</strong></td>
        </tr>
        <tr><td colspan="4">Branch : {branch_name}</td></tr>
        <tr><td colspan="4" nowrap="nowrap">{period}</td></tr>
    </table>
    <table cellpadding="2" cellspacing="0" width="100%" style="margin-top:6px;">
        <tr align="center">
            <td style="border:thin solid #000" width="3%">No.</td>
            <td style="border:thin solid #000" width="14%">CAM No.</td>
            <td style="border:thin solid #000" width="16%">Customer Name</td>
            <td style="border:thin solid #000" width="17%">Brand/Type</td>
            <td style="border:thin solid #000" width="12%">Net Finance</td>
            <td style="border:thin solid #000" width="6%">%IRR</td>
            <td style="border:thin solid #000" width="5%">Tenor (Month)</td>
            <td style="border:thin solid #000" width="11%">Approval Date</td>
            <td style="border:thin solid #000" width="8%">CMO</td>
            <td style="border:thin solid #000" width="8%">Approved by</td>
        </tr>
        {rows_html}
    </table>
    <table border="1" style="border-collapse:collapse; margin-top:10px;" cellpadding="2" cellspacing="0" width="70%">
        <tr><td colspan="8"><u><strong>Grand Total* ({period_plain}):</strong></u></td></tr>
        <tr style="background-color:#E5E5E5">
            <td width="10%"><strong>Contract Type</strong></td>
            <td width="10%"><strong>Type</strong></td>
            <td width="10%"><strong>Condition</strong></td>
            <td width="15%"><strong>Net Finance</strong></td>
            <td width="15%"><strong>Units</strong></td>
            <td width="10%"><strong>Avg. IRR</strong></td>
            <td width="15%"><strong>Avg. Tenor</strong></td>
            <td width="15%"><strong>Avg. Net Finance</strong></td>
        </tr>
        {sl_rows}
        {if_rows}
        <tr>
            <td colspan="3" align="right"><strong>Grand Total</strong></td>
            <td align="right">{grand_fin}</td>
            <td align="right">{grand_units}</td>
            <td align="right">{grand_irr}</td>
            <td align="right">{grand_tenor}</td>
            <td align="right">{grand_avg_fin}</td>
        </tr>
        <tr><td colspan="8"><i>*Grand Total Net Finance and Unit based on latest revised CAM (if any)</i></td></tr>
    </table>
</body>
</html>"""

def _render_rows_html(rows: list[dict]) -> str:
    if not rows:
        return ('<tr><td colspan="10" align="center" style="border:thin solid #000">'
                '<font color="red"><b>Data Not Found</b></font></td></tr>')
    parts = []
    for d in rows:
        parts.append(f"""
            <tr>
                <td style="border:thin solid #000" align="center">{d['no']}</td>
                <td style="border:thin solid #000">{d['cam_no']}</td>
                <td style="border:thin solid #000">{d['customer_name']}</td>
                <td style="border:thin solid #000">{d['brand_type']}</td>
                <td style="border:thin solid #000" align="right">{d['net_finance']:,.0f}</td>
                <td style="border:thin solid #000" align="right">{d['irr']:.2f}</td>
                <td style="border:thin solid #000" align="center">{d['tenor']}</td>
                <td style="border:thin solid #000">{d['approval_date']}</td>
                <td style="border:thin solid #000">{d['cmo']}</td>
                <td style="border:thin solid #000">{d['approved_by']}</td>
            </tr>""")
    return ''.join(parts)

def _render_segment_rows_html(label: str, group: dict) -> str:
    rows_html = []
    for idx, (key, type_label, cond_label) in enumerate(_SEGMENT_LABELS):
        seg = group[key]
        contract_cell = f'<td rowspan="5">{label}</td>' if idx == 0 else ''
        rows_html.append(f"""
            <tr>
                {contract_cell}
                <td>{type_label}</td>
                <td>{cond_label}</td>
                <td align="right">{seg['net_finance']:,.0f}</td>
                <td align="right">{seg['units']}</td>
                <td align="right">{seg['avg_irr']}</td>
                <td align="right">{seg['avg_tenor']}</td>
                <td align="right">{seg['avg_net_finance']}</td>
            </tr>""")
    total = group['total']
    rows_html.append(f"""
        <tr>
            <td colspan="2" align="right"><strong>Grand Total {label}</strong></td>
            <td align="right">{total['net_finance']:,.0f}</td>
            <td align="right">{total['units']}</td>
            <td align="right">{total['avg_irr']}</td>
            <td align="right">{total['avg_tenor']}</td>
            <td align="right">{total['avg_net_finance']}</td>
        </tr>""")
    return ''.join(rows_html)

@cam_by_app_date_bp.route('/cam-by-app-date/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch_cd', '').strip()

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_totals = _fetch_grand_totals(where_sql, params)
        gt = grand_totals['grand_total']

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end) or (f"Name : {lessee_nm}" if lessee_nm else '')

        html = _PRINT_TEMPLATE.format(
            company = current_app.config.get('COMPANY_NAME', ''),
            today = datetime.now().strftime('%d-%m-%Y'),
            printed_by = session.get('username', ''),
            branch_name = _branch_name(scope['branch_cd_search']) or 'All Branches',
            period = period,
            period_plain = period.replace('Period : ', '') or 'All Dates',
            rows_html = _render_rows_html(rows),
            sl_rows = _render_segment_rows_html('SL', grand_totals['sl']),
            if_rows = _render_segment_rows_html('IF', grand_totals['if']),
            grand_fin = f"{gt['net_finance']:,.0f}",
            grand_units = gt['units'],
            grand_irr = gt['avg_irr'],
            grand_tenor = gt['avg_tenor'],
            grand_avg_fin = gt['avg_net_finance'],
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("CAM by Approval Date print error")
        return jsonify({'error': str(e)}), 500

def _fetch_cam_basic_info(appl_no: str) -> dict:
    rows = _exec_rows(
        """
        SELECT LTRIM(RTRIM(a.CMO)) AS CMO, b.LESSEE_NM, b.apless, c.USERNAME
        FROM CAFCONT a WITH (NOLOCK)
            LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
            LEFT JOIN APP_USRMST c WITH (NOLOCK) ON a.CMO = c.EMPLOYEE_ID
        WHERE a.APPL_NO = ?
        """,
        [appl_no],
    )
    return rows[0] if rows else {}

def _fetch_revision_suffix(appl_no: str) -> str:
    rows = _exec_rows("SELECT REVISI FROM MCMOREVISI WITH (NOLOCK) WHERE NEWAPPL_NO = ?", [appl_no])
    return f"R{rows[0]['REVISI']}" if rows else ''

def _fetch_approval_history(appl_no: str, require_comment: bool) -> list[dict]:
    comment_filter = "AND a.comment IS NOT NULL AND a.comment <> ''" if require_comment else ""
    sql = f"""
        SELECT * FROM (
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMOUTSSTATUS a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
            
            UNION
            
            SELECT a.[No], a.APPR_ID, a.APPR_DATE, a.COMMENT, SUBSTRING(a.APPR_STS, 1, 1) AS APPR_STS
            FROM MCMOCAMStatus a WITH (NOLOCK)
            WHERE a.APPL_NO = ? AND (a.APPR_STS IN ('N', 'A', 'F', 'R') OR a.APPR_STS LIKE '%R')
        ) a
        WHERE a.APPR_ID IN (SELECT username FROM app_usrmst WHERE ACCESSCAM IN ('CRH', 'CRDH', 'DR', 'PD'))
            {comment_filter}
        ORDER BY a.[No]
    """
    rows = _exec_rows(sql, [appl_no, appl_no])
    history = []
    for r in rows:
        appr_date = r.get('APPR_DATE')
        history.append({
            'date': appr_date.strftime('%d-%m-%Y %H:%M:%S') if appr_date else '',
            'approved_by': r.get('APPR_ID'),
            'status': _STATUS_LABELS.get(r.get('APPR_STS'), ''),
            'comment': r.get('COMMENT') or '',
        })
    return history

def _fetch_related_cams(appl_no: str, apless: Optional[str]) -> list[dict]:
    related_rows = _exec_multi_resultset("EXEC USP_ListCamCommittee ?, ?", [appl_no, apless])
    related = []
    for r in related_rows:
        related_appl_no = r.get('appl_no')
        if not related_appl_no:
            continue
        existence_check = _fetch_approval_history(related_appl_no, require_comment=True)
        if not existence_check:
            continue
        history = _fetch_approval_history(related_appl_no, require_comment=False)
        info = _fetch_cam_basic_info(related_appl_no)
        suffix = _fetch_revision_suffix(related_appl_no)
        related.append({
            'appl_no_display': f"{related_appl_no}{suffix}",
            'customer_name': info.get('LESSEE_NM'),
            'cmo_name': info.get('USERNAME'),
            'history': history,
        })
    return related

@cam_by_app_date_bp.route('/cam-by-app-date/comments')
@login_required
def get_comments():
    appl_no = request.args.get('appl_no', '').strip()
    if not appl_no:
        return jsonify({'error': 'appl_no is required'}), 400

    try:
        info = _fetch_cam_basic_info(appl_no)
        suffix = _fetch_revision_suffix(appl_no)
        history = _fetch_approval_history(appl_no, require_comment=False)
        related = _fetch_related_cams(appl_no, info.get('apless'))
        return jsonify({
            'main': {
                'appl_no_display': f"{appl_no}{suffix}",
                'customer_name': info.get('LESSEE_NM'),
                'cmo_name': info.get('USERNAME'),
                'history': history,
            },
            'related': related,
        })
    except Exception as e:
        current_app.logger.exception("Credit committee comments error")
        return jsonify({'error': str(e)}), 500