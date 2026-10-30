from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app import db
from app.auth.decorators import login_required

in_progress_cam_bp = Blueprint('in_progress_cam', __name__, url_prefix='/CAM/Approved')

PAGE_SIZE = 100

_PASSENGER_MODELS = ['01', '02', '03']
_COMMERCIAL_MODELS = ['04', '05', '06', '07', '08', '09', '10']

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd or branch_cd == '000':
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

def _period_label(start_date: Optional[date], end_date: Optional[date]) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _resolve_scope(requested_branch: str) -> dict:
    access_level = session.get('accesslevel')
    emp_id = session.get('emp_id')
    own_branch_cd = session.get('branchcode')
    is_ho = (own_branch_cd == '100')

    branch = (requested_branch or '000') if is_ho else own_branch_cd

    if access_level == 'CMO':
        cmo_filter = [emp_id]
    elif access_level == 'CMH':
        rows = _exec_rows(
            "SELECT EMPLOYEE_ID FROM APP_USRMST WITH (NOLOCK) WHERE EMPLOYEE_ID = ? OR SUPERIOR = ?",
            [emp_id, emp_id],
        )
        cmo_filter = [r['EMPLOYEE_ID'] for r in rows] or [emp_id]
    else:
        cmo_filter = None

    return {
        'disabled': not is_ho,
        'branch': branch,
        'cmo_filter': cmo_filter,
        'show_all_branches': branch == '000',
    }

def _build_common_filter(scope: dict, start_date: Optional[date], end_date: Optional[date],
                          lessee_nm: str) -> tuple[str, list]:
    clauses = []
    params: list = []

    if scope['cmo_filter'] is not None:
        placeholders = ','.join('?' for _ in scope['cmo_filter'])
        clauses.append(f"a.CMO IN ({placeholders})")
        params.extend(scope['cmo_filter'])

    if not scope['show_all_branches'] and scope['branch']:
        clauses.append("a.BRANCH_CD = ?")
        params.append(scope['branch'])

    if start_date and end_date:
        clauses.append("CAST(a.CREATE_DATE AS DATE) >= ? AND CAST(a.CREATE_DATE AS DATE) <= ?")
        params.append(start_date)
        params.append(end_date)

    if lessee_nm:
        clauses.append("b.LESSEE_NM LIKE ?")
        params.append(f"%{lessee_nm}%")

    if not clauses:
        return "1=1", []
    return " AND ".join(clauses), params

_OUTS_SUBQUERY = """
INNER JOIN (
    SELECT SUM(OUTS) AS tot_outs, APPL_NO, ASOF FROM CAFCUST_OUTS WITH (NOLOCK) GROUP BY APPL_NO, ASOF
) d ON a.APPL_NO = d.APPL_NO AND d.ASOF = (
    SELECT MAX(ASOF) FROM CAFCUST_OUTS WITH (NOLOCK) WHERE APPL_NO = a.APPL_NO
)
"""

_LIST_SELECT = f"""
SELECT a.APPL_NO, a.Cons_Leas, b.LESSEE_NM, a.TENOR, a.DECLRATE, a.CREATE_DATE,
       c.fullname AS cmo, h.BRAND_NM, i.TYPE_NM, v.Net_Finance, k.fullname AS approver, d.tot_outs,
       a.APLESS, a.HAS_GUARANTOR, b.LESSEE_TP, g.SUPP, LTRIM(RTRIM(g.[CONDITION])) AS CONDITION
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    LEFT JOIN APP_USRMST c WITH (NOLOCK) ON a.CMO = c.EMPLOYEE_ID
    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
    LEFT JOIN UMFBRAND h WITH (NOLOCK) ON g.BRAND = h.BRAND
    LEFT JOIN UMFTYPE i WITH (NOLOCK) ON g.BRAND = i.BRAND AND g.TYPE = i.TYPE
    INNER JOIN V_LastCAMStatusApproval j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO AND j.APPR_STS = 'P'
    LEFT JOIN APP_USRMST k WITH (NOLOCK) ON j.APPR_ID = k.username
    INNER JOIN V_NetFinance v WITH (NOLOCK) ON j.APPL_NO = v.APPL_NO
{_OUTS_SUBQUERY}
"""

def _fetch_count(where_sql: str, params: list) -> int:
    sql = f"""
    SELECT COUNT(*) AS cnt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        INNER JOIN V_LastCAMStatusApproval j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO AND j.APPR_STS = 'P'
        INNER JOIN V_NetFinance v WITH (NOLOCK) ON j.APPL_NO = v.APPL_NO
    {_OUTS_SUBQUERY}
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    return int(rows[0]['cnt']) if rows else 0

def _fetch_page(where_sql: str, params: list, page: int, page_size: int) -> list[dict]:
    offset = (page - 1) * page_size
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY a.CREATE_DATE OFFSET ? ROWS FETCH NEXT ? ROWS ONLY"
    return _exec_rows(sql, params + [offset, page_size])

def _fetch_all(where_sql: str, params: list) -> list[dict]:
    sql = _LIST_SELECT + f"WHERE {where_sql} ORDER BY a.CREATE_DATE"
    return _exec_rows(sql, params)

def _fetch_date_range(where_sql: str, params: list) -> tuple[Optional[date], Optional[date]]:
    sql = f"""
    SELECT MIN(a.CREATE_DATE) AS min_dt, MAX(a.CREATE_DATE) AS max_dt
    FROM CAFCONT a WITH (NOLOCK)
        LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
        INNER JOIN V_LastCAMStatusApproval j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO AND j.APPR_STS = 'P'
        INNER JOIN V_NetFinance v WITH (NOLOCK) ON j.APPL_NO = v.APPL_NO
    {_OUTS_SUBQUERY}
    WHERE {where_sql}
    """
    rows = _exec_rows(sql, params)
    if not rows:
        return None, None
    min_dt, max_dt = rows[0].get('min_dt'), rows[0].get('max_dt')
    return (min_dt.date() if min_dt else None), (max_dt.date() if max_dt else None)

def _fetch_revisions(appl_nos: list[str]) -> dict[str, str]:
    if not appl_nos:
        return {}
    placeholders = ','.join('?' for _ in appl_nos)
    rows = _exec_rows(
        f"SELECT NEWAPPL_NO, REVISI FROM MCMORevisi WITH (NOLOCK) WHERE NEWAPPL_NO IN ({placeholders})",
        appl_nos,
    )
    return {r['NEWAPPL_NO']: f"R{r['REVISI']}" for r in rows}

def _edit_link_params(r: dict) -> dict:
    cons_leas = r.get('Cons_Leas')
    fin_type = 1 if cons_leas == 'C' else 2 if cons_leas == 'L' else None
    guarantor = 2 if r.get('HAS_GUARANTOR') == '1' else 1
    ind_cor = 2 if r.get('LESSEE_TP') == 'PT' else 1
    c2c = 1 if str(r.get('SUPP')) == '999999' else 0
    condition = (r.get('CONDITION') or '').strip()
    new_car = 2 if condition == 'NEW' else 1

    return {
        'id1': fin_type, 'id2': ind_cor, 'id3': 1, 'id4': guarantor, 'id5': new_car,
        'id6': 'Edit', 'id7': r.get('TENOR'), 'id8': c2c,
        'apless': r.get('APLESS'), 'applno': r.get('APPL_NO'),
    }

def _row_to_detail(r: dict, no: int, revision: str) -> dict:
    create_date = r.get('CREATE_DATE')
    brand = (r.get('BRAND_NM') or '').strip()
    type_nm = (r.get('TYPE_NM') or '').strip()
    return {
        'no': no,
        'cam_no': f"{r.get('Cons_Leas')} - {r.get('APPL_NO')}{revision}",
        'appl_no': r.get('APPL_NO'),
        'customer_name': r.get('LESSEE_NM') or '-',
        'brand_type': f"{brand} / {type_nm}" if (brand or type_nm) else '-',
        'net_finance': _num(r.get('Net_Finance')),
        'total_outstanding': _num(r.get('tot_outs')),
        'interest': round(_num(r.get('DECLRATE')), 4),
        'tenor': int(_num(r.get('TENOR'))),
        'create_date': create_date.strftime('%d-%m-%Y') if create_date else '',
        'cmo': r.get('cmo') or '-',
        'status': r.get('approver') or '-',
        'is_revision': bool(revision),
        'edit_link_params': _edit_link_params(r),
    }

_GRANDTOTAL_BASE = """
SELECT ISNULL(SUM(v.Net_Finance), 0) AS tot_fin, ISNULL(COUNT(a.APPL_NO), 0) AS tot_unit
FROM CAFCONT a WITH (NOLOCK)
    LEFT JOIN CAFLESS b WITH (NOLOCK) ON a.APLESS = b.APLESS
    LEFT JOIN CAFEQUIP g WITH (NOLOCK) ON a.APPL_NO = g.APPL_NO
    INNER JOIN V_LastCAMStatusApproval j WITH (NOLOCK) ON a.APPL_NO = j.APPL_NO AND j.APPR_STS = 'P'
    INNER JOIN V_NetFinance v WITH (NOLOCK) ON a.APPL_NO = v.APPL_NO
WHERE {filter}
  AND g.MODEL IN ({models})
  AND g.CONDITION = ?
"""

def _fetch_segment_total(where_sql: str, params: list, models: list[str], condition: str) -> dict:
    placeholders = ','.join('?' for _ in models)
    sql = _GRANDTOTAL_BASE.format(filter=where_sql, models=placeholders)
    rows = _exec_rows(sql, params + list(models) + [condition])
    row = rows[0] if rows else {}
    return {'net_finance': _num(row.get('tot_fin')), 'units': int(_num(row.get('tot_unit')))}

def _fetch_grand_totals(where_sql: str, params: list) -> dict:
    breakdown = {
        'passenger_new': _fetch_segment_total(where_sql, params, _PASSENGER_MODELS, 'NEW'),
        'passenger_used': _fetch_segment_total(where_sql, params, _PASSENGER_MODELS, 'USED'),
        'commercial_new': _fetch_segment_total(where_sql, params, _COMMERCIAL_MODELS, 'NEW'),
        'commercial_used': _fetch_segment_total(where_sql, params, _COMMERCIAL_MODELS, 'USED'),
    }
    grand_total = {
        'net_finance': sum(s['net_finance'] for s in breakdown.values()),
        'units': sum(s['units'] for s in breakdown.values()),
    }
    return {**breakdown, 'grand_total': grand_total}

@in_progress_cam_bp.route('/in-progress-cam/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

@in_progress_cam_bp.route('/in-progress-cam/context')
@login_required
def get_context():
    scope = _resolve_scope(request.args.get('branch', '').strip())
    return jsonify({'disabled': scope['disabled'], 'forced_branch_cd': scope['branch']})

@in_progress_cam_bp.route('/in-progress-cam/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch', '').strip()
    try:
        page = max(int(request.args.get('page', 1) or 1), 1)
    except ValueError:
        page = 1

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        total_count = _fetch_count(where_sql, params)
        total_pages = max((total_count + PAGE_SIZE - 1) // PAGE_SIZE, 1)
        page = min(page, total_pages)

        raw_rows = _fetch_page(where_sql, params, page, PAGE_SIZE)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)

        start_no = (page - 1) * PAGE_SIZE + 1
        details = [_row_to_detail(r, start_no + i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows)]

        grand_totals = _fetch_grand_totals(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        return jsonify({
            'rows': details,
            'total_count': total_count,
            'page': page,
            'page_size': PAGE_SIZE,
            'total_pages': total_pages,
            'grand_totals': grand_totals,
            'period_label': period,
            'disabled': scope['disabled'],
            'search_branch': scope['branch'],
            'branch_name': _branch_name(scope['branch']),
        })
    except Exception as e:
        current_app.logger.exception("In Progress CAM list error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_TITLE_FONT = Font(bold=True, size=16)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

_LIST_COLUMNS = ['No.', 'CAM No.', 'Customer Name', 'Brand/Type', 'Net Finance', 'Total Outstanding', 'Interest (%)', 'Tenor (Month)', 'Create Date', 'CMO', 'Status']

_SEGMENT_LABELS = [
    ('passenger_new', 'Passenger', 'New'),
    ('passenger_used', 'Passenger', 'Used'),
    ('commercial_new', 'Commercial', 'New'),
    ('commercial_used', 'Commercial', 'Used'),
]

def _build_workbook(period_label: str, lessee_nm: str, branch_name: str, rows: list[dict], grand_totals: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'InProgressCAM'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"
    ws['D1'].alignment = Alignment(wrap_text=True)

    ws.merge_cells('A2:D2')
    ws['A2'] = 'List In Progress CAM'
    ws['A2'].font = _TITLE_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {branch_name or 'All Branches'}"

    ws.merge_cells('A4:D4')
    ws['A4'] = period_label or (f"Name : {lessee_nm}" if lessee_nm else '')

    row_num = 6
    for col_idx, label in enumerate(_LIST_COLUMNS, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=label)
        cell.border = _ALL_BORDERS
        cell.alignment = _CENTER
        cell.font = _HEADER_FONT
    row_num += 1

    if not rows:
        ws.merge_cells(f'A{row_num}:K{row_num}')
        cell = ws[f'A{row_num}']
        cell.value = 'Data Not Found'
        cell.font = Font(color='FFFF0000', bold=True)
        cell.alignment = _CENTER
        row_num += 1
    else:
        for d in rows:
            values = [d['no'], d['cam_no'], d['customer_name'], d['brand_type'], d['net_finance'],
                      d['total_outstanding'], d['interest'], d['tenor'], d['create_date'], d['cmo'], d['status']]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = _ALL_BORDERS
                if col_idx in (5, 6, 7):
                    cell.alignment = _RIGHT
            row_num += 1

    row_num += 2
    ws.merge_cells(f'A{row_num}:D{row_num}')
    ws[f'A{row_num}'] = f"Grand Total* ({period_label.replace('Period : ', '') or 'All Dates'})"
    ws[f'A{row_num}'].font = _HEADER_FONT
    row_num += 1

    gt_headers = ['Type', 'Condition', 'Net Finance', 'Units']
    for col_idx, label in enumerate(gt_headers, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=label)
        cell.border = _ALL_BORDERS
        cell.font = _HEADER_FONT
    row_num += 1

    for key, type_label, cond_label in _SEGMENT_LABELS:
        seg = grand_totals[key]
        values = [type_label, cond_label, seg['net_finance'], seg['units']]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = _ALL_BORDERS
            if col_idx >= 3:
                cell.alignment = _RIGHT
        row_num += 1

    total = grand_totals['grand_total']
    ws.merge_cells(f'A{row_num}:B{row_num}')
    gt_label_cell = ws.cell(row=row_num, column=1, value='Grand Total')
    gt_label_cell.font = _HEADER_FONT
    gt_label_cell.alignment = _CENTER
    for col_idx, val in ((3, total['net_finance']), (4, total['units'])):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.border = _ALL_BORDERS
        cell.alignment = _RIGHT
    row_num += 2

    ws.merge_cells(f'A{row_num}:D{row_num}')
    ws[f'A{row_num}'] = '*Grand Total Net Finance and Unit based on latest revised CAM (if any)'
    ws[f'A{row_num}'].font = Font(italic=True)

    ws.column_dimensions['A'].width = 14
    for col in 'BCDEFGHIJK':
        ws.column_dimensions[col].width = 16

    return wb

@in_progress_cam_bp.route('/in-progress-cam/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch', '').strip()

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_totals = _fetch_grand_totals(where_sql, params)

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end)

        wb = _build_workbook(period, lessee_nm, _branch_name(scope['branch']), rows, grand_totals)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='ListInProgressCAM.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("In Progress CAM export error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{company} - List In Progress CAM</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();">
    <table cellpadding="2" cellspacing="0" width="100%">
        <tr>
            <td width="25%" nowrap="nowrap">{company}</td>
            <td width="50%"></td>
            <td width="10%"></td>
            <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
        </tr>
        <tr>
            <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;"><strong>List In Progress CAM</strong></td>
        </tr>
        <tr><td colspan="4">Branch : {branch_name}</td></tr>
        <tr><td colspan="4" nowrap="nowrap">{period}</td></tr>
    </table>
    <table cellpadding="2" cellspacing="0" width="100%" style="margin-top:6px;">
        <tr align="center">
            <td style="border:thin solid #000" width="3%">No.</td>
            <td style="border:thin solid #000" width="14%">CAM No.</td>
            <td style="border:thin solid #000" width="16%">Customer Name</td>
            <td style="border:thin solid #000" width="17%">Brand/Type</td>
            <td style="border:thin solid #000" width="12%">Net Finance</td>
            <td style="border:thin solid #000" width="12%">Total Outstanding</td>
            <td style="border:thin solid #000" width="6%">Interest (%)</td>
            <td style="border:thin solid #000" width="5%">Tenor (Month)</td>
            <td style="border:thin solid #000" width="11%">Create Date</td>
            <td style="border:thin solid #000" width="8%">CMO</td>
            <td style="border:thin solid #000" width="8%">Status</td>
        </tr>
        {rows_html}
    </table>
    <table border="1" style="border-collapse:collapse; margin-top:10px;" cellpadding="2" cellspacing="0" width="40%">
        <tr><td colspan="4"><u><strong>Grand Total* ({period_plain}):</strong></u></td></tr>
        <tr style="background-color:#E5E5E5">
            <td width="10%"><strong>Type</strong></td>
            <td width="10%"><strong>Condition</strong></td>
            <td width="15%"><strong>Net Finance</strong></td>
            <td width="15%"><strong>Units</strong></td>
        </tr>
        {segment_rows}
        <tr>
            <td colspan="2" align="center"><strong>Grand Total</strong></td>
            <td align="right">{grand_fin}</td>
            <td align="right">{grand_units}</td>
        </tr>
        <tr><td colspan="4"><i>*Grand Total Net Finance and Unit based on latest revised CAM (if any)</i></td></tr>
    </table>
</body>
</html>"""

def _render_rows_html(rows: list[dict]) -> str:
    if not rows:
        return ('<tr><td colspan="11" align="center" style="border:thin solid #000">'
                '<font color="red"><b>Data Not Found</b></font></td></tr>')
    parts = []
    for d in rows:
        parts.append(f"""
            <tr>
                <td style="border:thin solid #000" align="center">{d['no']}</td>
                <td style="border:thin solid #000">{d['cam_no']}</td>
                <td style="border:thin solid #000">{d['customer_name']}</td>
                <td style="border:thin solid #000">{d['brand_type']}</td>
                <td style="border:thin solid #000" align="right">{d['net_finance']:,.0f}</td>
                <td style="border:thin solid #000" align="right">{d['total_outstanding']:,.0f}</td>
                <td style="border:thin solid #000" align="right">{d['interest']:.4f}</td>
                <td style="border:thin solid #000" align="center">{d['tenor']}</td>
                <td style="border:thin solid #000">{d['create_date']}</td>
                <td style="border:thin solid #000">{d['cmo']}</td>
                <td style="border:thin solid #000">{d['status']}</td>
            </tr>""")
    return ''.join(parts)

def _render_segment_rows_html(grand_totals: dict) -> str:
    parts = []
    for key, type_label, cond_label in _SEGMENT_LABELS:
        seg = grand_totals[key]
        parts.append(f"""
            <tr>
                <td>{type_label}</td>
                <td>{cond_label}</td>
                <td align="right">{seg['net_finance']:,.0f}</td>
                <td align="right">{seg['units']}</td>
            </tr>""")
    return ''.join(parts)

@in_progress_cam_bp.route('/in-progress-cam/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    lessee_nm = request.args.get('lessee_nm', '').strip()
    requested_branch = request.args.get('branch', '').strip()

    scope = _resolve_scope(requested_branch)
    where_sql, params = _build_common_filter(scope, start_date, end_date, lessee_nm)

    try:
        raw_rows = _fetch_all(where_sql, params)
        appl_nos = [r['APPL_NO'] for r in raw_rows if r.get('APPL_NO')]
        revisions = _fetch_revisions(appl_nos)
        rows = [_row_to_detail(r, i, revisions.get(r['APPL_NO'], '')) for i, r in enumerate(raw_rows, start=1)]
        grand_totals = _fetch_grand_totals(where_sql, params)
        gt = grand_totals['grand_total']

        if start_date and end_date:
            period = _period_label(start_date, end_date)
        else:
            data_start, data_end = _fetch_date_range(where_sql, params)
            period = _period_label(data_start, data_end) or (f"Name : {lessee_nm}" if lessee_nm else '')

        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_branch_name(scope['branch']) or 'All Branches',
            period=period,
            period_plain=period.replace('Period : ', '') or 'All Dates',
            rows_html=_render_rows_html(rows),
            segment_rows=_render_segment_rows_html(grand_totals),
            grand_fin=f"{gt['net_finance']:,.0f}",
            grand_units=gt['units'],
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("In Progress CAM print error")
        return jsonify({'error': str(e)}), 500