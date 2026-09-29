from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from flask import Blueprint, jsonify, request, current_app, session, send_file, Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app import db
from app.auth.decorators import login_required

summary_approved_cam_bp = Blueprint('summary_approved_cam', __name__, url_prefix='/CAM/Approved')

def _exec_rows(sql: str, params: list) -> list[dict]:
    raw_conn = db.session.connection().connection
    cursor = raw_conn.cursor()
    try:
        cursor.execute(sql, params)
        if cursor.description is None:
            return []
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        cursor.close()

def _parse_date(raw: str) -> Optional[date]:
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        return None

def _num(val) -> float:
    if val is None:
        return 0.0
    return float(val)

def _fmt_ddmmyyyy(d: Optional[date]) -> str:
    return d.strftime('%d-%m-%Y') if d else ''

@summary_approved_cam_bp.route('/summary-approved-cam/branches')
@login_required
def list_branches():
    try:
        rows = _exec_rows(
            "SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD <> '100' ORDER BY BRANCH_CD",
            [],
        )
        return jsonify([{'branch_cd': r.get('BRANCH_CD'), 'branch_name': r.get('BRANCH_NAME')} for r in rows])
    except Exception as e:
        current_app.logger.exception("Branch list error")
        return jsonify({'error': str(e)}), 500

def _resolve_branch(requested_branch_cd: str, cmo_branch_param):
    access_level = session.get('accesslevel')
    own_branch_cd = session.get('branchcode')

    if access_level in ('CMO', 'CMH', 'BM'):
        return {'search_branch': own_branch_cd, 'disabled': True, 'sp_branch_param': cmo_branch_param}
    if own_branch_cd == '100':
        return {'search_branch': requested_branch_cd, 'disabled': False, 'sp_branch_param': requested_branch_cd}
    return {'search_branch': own_branch_cd, 'disabled': True, 'sp_branch_param': own_branch_cd}

@summary_approved_cam_bp.route('/summary-approved-cam/context')
@login_required
def get_context():
    resolved = _resolve_branch(request.args.get('branch_cd_search', '').strip(), '')
    return jsonify({
        'disabled': resolved['disabled'],
        'forced_branch_cd': resolved['search_branch'],
    })

def _fetch_rows(start_date: Optional[date], end_date: Optional[date], sp_branch_param) -> list[dict]:
    sql = "EXEC USP_S_SUMMAPPRCAM ?, ?, ?, ?, ?, ?"
    params = [
        start_date.isoformat() if start_date else '',
        end_date.isoformat() if end_date else '',
        session.get('emp_id'),
        session.get('branchcode'),
        session.get('accesslevel'),
        sp_branch_param,
    ]
    return _exec_rows(sql, params)

def _process_rows(rows: list[dict]) -> list[dict]:
    out: list[dict] = []
    grand = {'cam_unit': 0.0, 'cam_amt': 0.0, 'revisi_unit': 0.0, 'revisi_amt': 0.0, 'baris_unit': 0.0, 'baris_amt': 0.0}
    sub = dict(grand)
    iii = 0
    temp_fullname = None

    for r in rows:
        baris = _num(r.get('baris'))
        fullname = r.get('FULLNAME')
        unit_new = _num(r.get('unit_new'))
        unit_rev = _num(r.get('unit_rev'))
        net_new = _num(r.get('net_finance_new'))
        net_rev = _num(r.get('net_finance_rev'))
        total_unit = unit_new + unit_rev
        total_amt = net_new + net_rev

        out.append({
            'type': 'detail',
            'cmo_name': fullname if fullname != temp_fullname else '',
            'approval_date': r.get('AdmDate').strftime('%d-%m-%Y') if r.get('AdmDate') else '',
            'unit_new': int(unit_new),
            'net_finance_new': net_new,
            'unit_rev': int(unit_rev),
            'net_finance_rev': net_rev,
            'total_unit': int(total_unit),
            'total_net_finance': total_amt,
        })
        temp_fullname = fullname

        iii += 1
        if iii <= baris:
            deltas = {
                'cam_unit': unit_new, 'cam_amt': net_new,
                'revisi_unit': unit_rev, 'revisi_amt': net_rev,
                'baris_unit': total_unit, 'baris_amt': total_amt,
            }
            for k, v in deltas.items():
                sub[k] += v
                grand[k] += v

        if iii == baris:
            out.append({
                'type': 'subtotal',
                'unit_new': int(sub['cam_unit']),
                'net_finance_new': sub['cam_amt'],
                'unit_rev': int(sub['revisi_unit']),
                'net_finance_rev': sub['revisi_amt'],
                'total_unit': int(sub['baris_unit']),
                'total_net_finance': sub['baris_amt'],
            })
            iii = 0
            sub = {k: 0.0 for k in sub}

    out.append({
        'type': 'grand_total',
        'unit_new': int(grand['cam_unit']),
        'net_finance_new': grand['cam_amt'],
        'unit_rev': int(grand['revisi_unit']),
        'net_finance_rev': grand['revisi_amt'],
        'total_unit': int(grand['baris_unit']),
        'total_net_finance': grand['baris_amt'],
    })
    return out

def _period_label(start_date: Optional[date], end_date: Optional[date], branch_cd_search: str) -> str:
    if start_date and end_date:
        return f"Period : {_fmt_ddmmyyyy(start_date)} to {_fmt_ddmmyyyy(end_date)}"
    return ''

def _branch_name(branch_cd: str) -> str:
    if not branch_cd:
        return ''
    rows = _exec_rows("SELECT BRANCH_NAME FROM UMFBRANCH WITH (NOLOCK) WHERE BRANCH_CD = ?", [branch_cd])
    return rows[0]['BRANCH_NAME'] if rows else ''

@summary_approved_cam_bp.route('/summary-approved-cam/list')
@login_required
def get_list():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    resolved = _resolve_branch(requested_branch, '')
    try:
        rows = _fetch_rows(start_date, end_date, resolved['sp_branch_param'])
        data = _process_rows(rows)
        return jsonify({
            'data': data,
            'disabled': resolved['disabled'],
            'search_branch': resolved['search_branch'],
        })
    except Exception as e:
        current_app.logger.exception("Summary Approved CAM list error")
        return jsonify({'error': str(e)}), 500

_PRINT_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Credit Approval Memorandum</title>
<style>
body {{ font-family: Tahoma, Geneva, sans-serif; font-size: 12px; }}
td {{ font-size: 12px; }}
table {{ border-collapse: collapse; }}
@media print {{
    .report-header {{ display: none; }}
}}
</style>
</head>
<body bgcolor="#FFFFFF" onload="window.print();" onafterprint="window.close();">
    <div class="report-header">
        <table style="border-collapse:collapse; margin-left:1px;" cellpadding="2" cellspacing="0" width="100%">
            <tr>
                <td width="25%" nowrap="nowrap">{company}</td>
                <td width="50%"></td>
                <td width="10%"></td>
                <td width="15%">Date : {today}<br>Print By : {printed_by}</td>
            </tr>
            <tr>
                <td colspan="4" align="center" style="font-size:16px; letter-spacing:0.2em;" nowrap="nowrap"><strong>Summary Approved CAM</strong></td>
            </tr>
            <tr>
                <td colspan="4">Branch : {branch_name}</td>
            </tr>
            <tr>
                <td colspan="4" nowrap="nowrap">{period}</td>
            </tr>
        </table>
    </div>
    <table style="border:thin solid #000000;" cellpadding="2" cellspacing="0" width="100%">
        <tr align="center">
            <td style="border:thin solid #000;" rowspan="2">CMO</td>
            <td style="border:thin solid #000;" rowspan="2">Approval Date</td>
            <td style="border:thin solid #000" colspan="2">New CAM</td>
            <td style="border:thin solid #000" colspan="2">Revisi</td>
            <td style="border:thin solid #000" colspan="2">Total</td>
        </tr>
        <tr align="center">
            <td style="border:thin solid #000;">Unit</td>
            <td style="border:thin solid #000;">Net Finance</td>
            <td style="border:thin solid #000;">Unit</td>
            <td style="border:thin solid #000;">Net Finance</td>
            <td style="border:thin solid #000;">Unit</td>
            <td style="border:thin solid #000;">Net Finance</td>
        </tr>
    {rows_html}
    </table>
</body>
</html>"""

def _render_rows_html(data: list[dict]) -> str:
    parts = []
    for row in data:
        if row['type'] == 'detail':
            parts.append(f"""
                <tr>
                    <td style="border:thin solid #000000;">{row['cmo_name']}</td>
                    <td align="right" style="border:thin solid #000000;">{row['approval_date']}</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_new']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_new']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_rev']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_rev']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_unit']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_net_finance']:,.0f}</td>
                </tr>""")
        elif row['type'] == 'subtotal':
            parts.append(f"""
                <tr>
                    <td style="border:thin solid #000000;"></td>
                    <td align="right" style="border:thin solid #000000;">Total</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_new']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_new']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_rev']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_rev']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_unit']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_net_finance']:,.0f}</td>
                </tr>""")
        else:  # grand_total
            parts.append(f"""
                <tr>
                    <td colspan="2" style="border:thin solid #000000;">Grand Total</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_new']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_new']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['unit_rev']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['net_finance_rev']:,.0f}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_unit']:,}</td>
                    <td align="right" style="border:thin solid #000000;">{row['total_net_finance']:,.0f}</td>
                </tr>""")
    return ''.join(parts)

@summary_approved_cam_bp.route('/summary-approved-cam/print')
@login_required
def print_view():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    resolved = _resolve_branch(requested_branch, None)
    try:
        rows = _fetch_rows(start_date, end_date, resolved['sp_branch_param'])
        data = _process_rows(rows)
        html = _PRINT_TEMPLATE.format(
            company=current_app.config.get('COMPANY_NAME', ''),
            today=datetime.now().strftime('%d-%m-%Y'),
            printed_by=session.get('username', ''),
            branch_name=_branch_name(resolved['search_branch']),
            period=_period_label(start_date, end_date, resolved['search_branch']),
            rows_html=_render_rows_html(data),
        )
        return Response(html, mimetype='text/html')
    except Exception as e:
        current_app.logger.exception("Summary Approved CAM print error")
        return jsonify({'error': str(e)}), 500

_THIN_SIDE = Side(style='thin', color='FF000000')
_ALL_BORDERS = Border(top=_THIN_SIDE, bottom=_THIN_SIDE, left=_THIN_SIDE, right=_THIN_SIDE)
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')

def _build_workbook(start_date: Optional[date], end_date: Optional[date], resolved: dict, data: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'SummaryApprovedCam'

    ws['A1'] = current_app.config.get('COMPANY_NAME', '')
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}"
    ws['D1'].alignment = Alignment(wrap_text=True)
    ws['D1'] = f"Date : {datetime.now().strftime('%d-%m-%Y')}\nPrint By : {session.get('username', '')}"

    ws.merge_cells('A2:D2')
    ws['A2'] = 'Summary Approved CAM'
    ws['A2'].font = _HEADER_FONT
    ws['A2'].alignment = _CENTER

    ws.merge_cells('A3:D3')
    ws['A3'] = f"Branch : {_branch_name(resolved['search_branch'])}"

    ws.merge_cells('A4:D4')
    ws['A4'] = _period_label(start_date, end_date, resolved['search_branch'])

    header_row1 = 6
    header_row2 = 7
    ws.merge_cells(f'A{header_row1}:A{header_row2}')
    ws.merge_cells(f'B{header_row1}:B{header_row2}')
    ws.merge_cells(f'C{header_row1}:D{header_row1}')
    ws.merge_cells(f'E{header_row1}:F{header_row1}')
    ws.merge_cells(f'G{header_row1}:H{header_row1}')

    ws[f'A{header_row1}'] = 'CMO'
    ws[f'B{header_row1}'] = 'Approval Date'
    ws[f'C{header_row1}'] = 'New CAM'
    ws[f'E{header_row1}'] = 'Revisi'
    ws[f'G{header_row1}'] = 'Total'
    ws[f'C{header_row2}'] = 'Unit'
    ws[f'D{header_row2}'] = 'Net Finance'
    ws[f'E{header_row2}'] = 'Unit'
    ws[f'F{header_row2}'] = 'Net Finance'
    ws[f'G{header_row2}'] = 'Unit'
    ws[f'H{header_row2}'] = 'Net Finance'

    for row_num in (header_row1, header_row2):
        for col in 'ABCDEFGH':
            cell = ws[f'{col}{row_num}']
            cell.border = _ALL_BORDERS
            cell.alignment = _CENTER

    excel_row = header_row2 + 1
    for row in data:
        if row['type'] == 'detail':
            ws[f'A{excel_row}'] = row['cmo_name']
            ws[f'B{excel_row}'] = row['approval_date']
        elif row['type'] == 'subtotal':
            ws[f'B{excel_row}'] = 'Total'
        else:
            ws.merge_cells(f'A{excel_row}:B{excel_row}')
            ws[f'A{excel_row}'] = 'Grand Total'

        ws[f'C{excel_row}'] = row['unit_new']
        ws[f'D{excel_row}'] = row['net_finance_new']
        ws[f'E{excel_row}'] = row['unit_rev']
        ws[f'F{excel_row}'] = row['net_finance_rev']
        ws[f'G{excel_row}'] = row['total_unit']
        ws[f'H{excel_row}'] = row['total_net_finance']

        for col in 'ABCDEFGH':
            cell = ws[f'{col}{excel_row}']
            cell.border = _ALL_BORDERS
            if col not in ('A', 'B'):
                cell.alignment = _RIGHT

        excel_row += 1

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 14
    for col in 'CDEFGH':
        ws.column_dimensions[col].width = 14

    return wb

@summary_approved_cam_bp.route('/summary-approved-cam/export')
@login_required
def export_excel():
    start_date = _parse_date(request.args.get('start_date', ''))
    end_date = _parse_date(request.args.get('end_date', ''))
    requested_branch = request.args.get('branch_cd_search', '').strip()

    resolved = _resolve_branch(requested_branch, None)
    try:
        rows = _fetch_rows(start_date, end_date, resolved['sp_branch_param'])
        data = _process_rows(rows)
        wb = _build_workbook(start_date, end_date, resolved, data)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='SummaryApprovedCam.xlsx',
        )
    except Exception as e:
        current_app.logger.exception("Summary Approved CAM export error")
        return jsonify({'error': str(e)}), 500