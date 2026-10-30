from flask import Blueprint, jsonify, request, current_app, session, make_response
from app import db
from app.auth.decorators import login_required

from sqlalchemy import text
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

pending_contract_delete_bp = Blueprint("pending_contract_delete", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")
_FONT_TITLE  = Font(name="Tahoma", size=14, bold=True)
_FONT_HEADER = Font(name="Tahoma", size=10, bold=True)
_FONT_DATA   = Font(name="Tahoma", size=10)
_FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")

MAX_PAGE_SIZE = 200

def _apply(cell, font=None, alignment=None, border=None, fill=None):
    if font:      cell.font      = font
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border
    if fill:      cell.fill      = fill

def _is_admin() -> bool:
    return session.get("branchcode", "") == "100"

def _serialize_rows(rows) -> list[dict]:
    result = []
    for r in rows:
        appr = r["Appr_Date"]
        del_ = r["Del_Date"]
        result.append({
            "appl_no":   r["Appl_No"]   or "",
            "lease_no":  r["Lease_no"]  or "",
            "lessee_nm": r["Lessee_nm"] or "",
            "appr_date": appr.strftime("%Y-%m-%d") if appr else None,
            "del_date":  del_.strftime("%Y-%m-%d") if del_ else None,
            "cmo":       r["Cmo"]       or "",
        })
    return result

def _fetch_all(branch_cd: str) -> list[dict]:
    today = date.today().strftime("%Y-%m-%d")
    rows = db.session.execute(
        text("EXEC USP_AutoDeleteContractList :today_date, :branch"),
        {"today_date": today, "branch": branch_cd},
    ).mappings().all()
    return _serialize_rows(rows)

def _build_excel(rows: list[dict], branch_name: str) -> bytes:
    today_str = datetime.now().strftime("%d %B %Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "PendingContractDelete"

    ws.merge_cells("A1:G1")
    ws["A1"].value = "PENDING CONTRACT WILL BE DELETED BY SYSTEM"
    _apply(ws["A1"], font=_FONT_TITLE, alignment=_CENTER)
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:G2")
    ws["A2"].value = f"As of {today_str}"
    _apply(ws["A2"], font=Font(name="Tahoma", size=12, bold=True), alignment=_CENTER)

    ws.merge_cells("A3:G3")
    ws["A3"].value = f"Branch : {branch_name}"
    _apply(ws["A3"], font=_FONT_DATA, alignment=_LEFT)

    headers = [
        "No.", "CAM No.", "Contract No.", "Customer Name",
        "Last Approved Date", "Deleted Date By System", "CMO",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=col, value=h)
        _apply(c, font=_FONT_HEADER, fill=_FILL_HEADER, border=_BORDER, alignment=_CENTER)

    for idx, r in enumerate(rows):
        r_num = idx + 6

        def fmt_date(d_str):
            if not d_str: return ""
            try:
                return datetime.strptime(d_str, "%Y-%m-%d").strftime("%d-%m-%Y")
            except ValueError:
                return d_str

        row_vals = [
            idx + 1,
            r["appl_no"],
            r["lease_no"],
            r["lessee_nm"],
            fmt_date(r["appr_date"]),
            fmt_date(r["del_date"]),
            r["cmo"],
        ]
        for col, val in enumerate(row_vals, 1):
            c = ws.cell(row=r_num, column=col, value=val)
            _apply(c, font=_FONT_DATA, border=_BORDER,
                   alignment=_CENTER if col != 4 else _LEFT)

    total_row = len(rows) + 6
    ws.merge_cells(f"A{total_row}:G{total_row}")
    c = ws[f"A{total_row}"]
    c.value = f"Total Unit : {len(rows)}"
    _apply(c, font=_FONT_HEADER, alignment=_LEFT)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return buf.read()

@pending_contract_delete_bp.route("/PendingContractDelete/branches")
@login_required
def get_branches():

    try:
        user_branch_cd = session.get("branchcode", "")
        admin = _is_admin()

        if admin:
            branch_rows = db.session.execute(
                text("""
                    SELECT BRANCH_CD, BRANCH_NAME
                    FROM UMFBRANCH WITH(NOLOCK)
                    WHERE BRANCH_CD <> '100'
                    ORDER BY BRANCH_CD
                """)
            ).mappings().all()
        else:
            branch_rows = db.session.execute(
                text("""
                    SELECT BRANCH_CD, BRANCH_NAME
                    FROM UMFBRANCH WITH(NOLOCK)
                    WHERE BRANCH_CD = :cd
                """),
                {"cd": user_branch_cd},
            ).mappings().all()

        user_branch_name = ""
        for r in branch_rows:
            if r["BRANCH_CD"] == user_branch_cd:
                user_branch_name = r["BRANCH_NAME"]
                break
        if not user_branch_name and not admin:
            row = db.session.execute(
                text("SELECT BRANCH_NAME FROM UMFBRANCH WITH(NOLOCK) WHERE BRANCH_CD = :cd"),
                {"cd": user_branch_cd},
            ).mappings().first()
            if row:
                user_branch_name = row["BRANCH_NAME"]

        return jsonify({
            "is_admin":         admin,
            "user_branch_cd":   user_branch_cd,
            "user_branch_name": user_branch_name,
            "branches": [
                {"branch_cd": r["BRANCH_CD"], "branch_name": r["BRANCH_NAME"]}
                for r in branch_rows
            ],
        })

    except Exception:
        current_app.logger.exception("Branches fetch error")
        return jsonify({"error": "Failed to load branches"}), 500

@pending_contract_delete_bp.route("/PendingContractDelete/data")
@login_required
def get_data():

    branch_cd = request.args.get("branch_cd", "").strip()
    if not branch_cd:
        return jsonify({"error": "branch_cd is required"}), 400

    want_all = request.args.get("all", "").strip().lower() in ("1", "true", "yes")

    page  = max(request.args.get("page", 1, type=int) or 1, 1)
    limit = request.args.get("limit", 25, type=int) or 25
    limit = min(max(limit, 1), MAX_PAGE_SIZE)

    try:
        all_rows = _fetch_all(branch_cd)
        total    = len(all_rows)

        if want_all:
            page_rows = all_rows
        else:
            start = (page - 1) * limit
            page_rows = all_rows[start:start + limit]

        return jsonify({
            "rows":  page_rows,
            "total": total,
            "page":  1 if want_all else page,
            "limit": total if want_all else limit,
        })

    except Exception:
        current_app.logger.exception("Pending Contract Delete data fetch error")
        return jsonify({"error": "Failed to load data"}), 500

@pending_contract_delete_bp.route("/PendingContractDelete/excel")
@login_required
def export_excel():

    branch_cd = request.args.get("branch_cd", "").strip()
    if not branch_cd:
        return jsonify({"error": "branch_cd is required"}), 400

    try:
        branch_row = db.session.execute(
            text("SELECT BRANCH_NAME FROM UMFBRANCH WITH(NOLOCK) WHERE BRANCH_CD = :cd"),
            {"cd": branch_cd},
        ).mappings().first()
        branch_name = branch_row["BRANCH_NAME"] if branch_row else branch_cd

        xlsx_bytes = _build_excel(_fetch_all(branch_cd), branch_name)

        resp = make_response(xlsx_bytes)
        resp.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp.headers["Content-Disposition"] = (
            'attachment; filename="ListPendingContractDeletedBySystem.xlsx"'
        )
        return resp

    except Exception:
        current_app.logger.exception("Pending Contract Delete Excel export error")
        return jsonify({"error": "Failed to generate Excel"}), 500