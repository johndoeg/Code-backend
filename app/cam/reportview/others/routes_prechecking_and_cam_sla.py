from flask import Blueprint, request, current_app, make_response, jsonify
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text
import io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

prechecking_and_cam_sla_bp = Blueprint("prechecking_and_cam_sla", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

_FONT_HEADER = Font(name="Tahoma", size=12, bold=True)
_FONT_DATA   = Font(name="Tahoma", size=11)
_NUM_FMT     = "#,##0"

def _h(ws, ref: str, value: str):
    c = ws[ref]
    c.value     = value
    c.font      = _FONT_HEADER
    c.alignment = _CENTER
    c.border    = _BORDER

def _d(ws, ref: str, value, number_format: str | None = None):
    c = ws[ref]
    c.value     = value if value is not None else ""
    c.font      = _FONT_DATA
    c.alignment = _LEFT
    c.border    = _BORDER
    if number_format:
        c.number_format = number_format

def _fmt_dt(dt) -> str:
    if dt is None:
        return ""
    try:
        return dt.strftime("%d %B %Y %I:%M:%S")
    except Exception:
        return str(dt)

def _build_excel(date_from: str, date_to: str) -> bytes:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Prechecking & CAM SLA"

    s1_headers = [
        ("A", 20, "Prechecking ID"),
        ("B", 20, "CARRO Non CAI"),
        ("C", 50, "Lessee Name"),
        ("D", 30, "Leads In"),
        ("E", 35, "Document Status Code"),
        ("F", 35, "Last Approval Action Code"),
        ("G", 35, "Last Action Date"),
        ("H", 20, "Last Action By"),
        ("I", 20, "CMO"),
        ("J", 20, "CMH"),
        ("K", 20, "Total CAM in Period"),
        ("L", 20, "Total CAM All"),
    ]
    for col, width, label in s1_headers:
        ws1.column_dimensions[col].width = width
        _h(ws1, f"{col}1", label)

    rows1 = db.session.execute(
        text("EXEC USP_PrecheckingSLA :df, :dt"), {"df": date_from, "dt": date_to}
    ).mappings().all()

    for i, row in enumerate(rows1, start=2):
        _d(ws1, f"A{i}", row.get("Prechecking_Id"))
        _d(ws1, f"B{i}", row.get("Carro"))
        _d(ws1, f"C{i}", row.get("LesseeName"))
        _d(ws1, f"D{i}", _fmt_dt(row.get("CreatedDate")))
        _d(ws1, f"E{i}", row.get("DocumentStatusCode"))
        _d(ws1, f"F{i}", row.get("LastApprovalActionCode"))
        _d(ws1, f"G{i}", _fmt_dt(row.get("LastActionDate")))
        _d(ws1, f"H{i}", row.get("LastActionBy"))
        _d(ws1, f"I{i}", row.get("CMO"))
        _d(ws1, f"J{i}", row.get("CMH"))
        _d(ws1, f"K{i}", row.get("totcam_periode"))
        _d(ws1, f"L{i}", row.get("totcam_all"))

    ws2 = wb.create_sheet("CAM Proses")

    s2_headers = [
        ("A",20,"CAMNo"),       ("B",20,"Prechecking ID"),  ("C",35,"CARRO Non CARRO"),
        ("D",50,"Lessee Name"), ("E",30,"Created Date"),     ("F",30,"Execution"),
        ("G",35,"Dealer Name"), ("H",20,"TAHUN"),            ("I",20,"COLOUR"),
        ("J",20,"CONDITION"),   ("K",35,"BRAND NM"),         ("L",35,"MODEL NM"),
        ("M",50,"TYPE NM"),     ("N",20,"NTF"),              ("O",20,"L AMOUNT"),
        ("P",20,"TENOR"),       ("Q",35,"Approved1By"),      ("R",35,"Approved1Date"),
        ("S",35,"Approved2By"), ("T",35,"Approved2Date"),    ("U",35,"Approved3By"),
        ("V",35,"Approved3Date"),("W",35,"Approved4By"),     ("X",35,"Approved4Date"),
        ("Y",35,"Approved5By"), ("Z",35,"Approved5Date"),    ("AA",35,"Approved6By"),
        ("AB",35,"Approved6Date"),("AC",35,"Approved7By"),   ("AD",35,"Approved7Date"),
        ("AE",35,"Last Action"),("AF",35,"Last Action By"),  ("AG",35,"Last Action Date"),
    ]
    for col, width, label in s2_headers:
        ws2.column_dimensions[col].width = width
        _h(ws2, f"{col}1", label)

    rows2 = db.session.execute(
        text("EXEC USP_CAMSLA :df, :dt"), {"df": date_from, "dt": date_to}
    ).mappings().all()

    for j, row in enumerate(rows2, start=2):
        _d(ws2, f"A{j}",  row.get("appl_no"))
        _d(ws2, f"B{j}",  row.get("Prechecking_id"))
        _d(ws2, f"C{j}",  row.get("Carro"))
        _d(ws2, f"D{j}",  row.get("LESSEE_NM"))
        _d(ws2, f"E{j}",  _fmt_dt(row.get("create_date")))
        _d(ws2, f"F{j}",  row.get("EXECUTION"))
        _d(ws2, f"G{j}",  row.get("Dealer"))
        _d(ws2, f"H{j}",  row.get("TAHUN"))
        _d(ws2, f"I{j}",  row.get("COLOUR"))
        _d(ws2, f"J{j}",  row.get("CONDITION"))
        _d(ws2, f"K{j}",  row.get("BRAND_NM"))
        _d(ws2, f"L{j}",  row.get("model"))
        _d(ws2, f"M{j}",  row.get("tipe_mobil"))
        _d(ws2, f"N{j}",  row.get("ntf"),      number_format=_NUM_FMT)
        _d(ws2, f"O{j}",  row.get("L_AMOUNT"), number_format=_NUM_FMT)
        _d(ws2, f"P{j}",  row.get("tenor"))
        _d(ws2, f"Q{j}",  row.get("Approved1By"))
        _d(ws2, f"R{j}",  row.get("Approved1Date"))
        _d(ws2, f"S{j}",  row.get("Approved2By"))
        _d(ws2, f"T{j}",  row.get("Approved2Date"))
        _d(ws2, f"U{j}",  row.get("Approved3By"))
        _d(ws2, f"V{j}",  row.get("Approved3Date"))
        _d(ws2, f"W{j}",  row.get("Approved4By"))
        _d(ws2, f"X{j}",  row.get("Approved4Date"))
        _d(ws2, f"Y{j}",  row.get("Approved5By"))
        _d(ws2, f"Z{j}",  row.get("Approved5Date"))
        _d(ws2, f"AA{j}", row.get("Approved6By"))
        _d(ws2, f"AB{j}", row.get("Approved6Date"))
        _d(ws2, f"AC{j}", row.get("Approved7By"))
        _d(ws2, f"AD{j}", row.get("Approved7Date"))
        _d(ws2, f"AE{j}", row.get("LastAction"))
        _d(ws2, f"AF{j}", row.get("LastActionBy"))
        _d(ws2, f"AG{j}", row.get("LastActionDate"))

    wb.active = ws1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    return buf.read()

@prechecking_and_cam_sla_bp.route("/PrecheckingAndCAMSLA/excel")
@login_required
def export_excel():
   
    date_from = request.args.get("dateFrom", "").strip()
    date_to   = request.args.get("dateTo",   "").strip()

    if not date_from or not date_to:
        return jsonify({"error": "dateFrom and dateTo are required"}), 400

    try:
        xlsx = _build_excel(date_from, date_to)
        resp = make_response(xlsx)
        resp.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp.headers["Content-Disposition"] = (
            'attachment; filename="Prechecking_CAM_SLA.xlsx"'
        )
        return resp
    except Exception:
        current_app.logger.exception("Prechecking CAM SLA Excel error")
        return jsonify({"error": "Failed to generate Excel"}), 500