from flask import Blueprint, jsonify, request, current_app, make_response
from app import db
from app.auth.decorators import login_required

from sqlalchemy import text
from datetime import datetime
import calendar, io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

f6pd_bp = Blueprint("f6pd", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="bottom")
_LEFT   = Alignment(horizontal="left",   vertical="bottom")

_FONT_TITLE    = Font(name="Tahoma", size=12, bold=True)
_FONT_SUBTITLE = Font(name="Tahoma", size=8,  bold=True)
_FONT_DATA     = Font(name="Tahoma", size=7)
_FONT_CMH      = Font(name="Tahoma", size=7, bold=True)

_NUM_FMT = '#,##0'

def _apply(cell, font=None, alignment=None, border=None, fill=None, number_format=None):
    if font:          cell.font          = font
    if alignment:     cell.alignment     = alignment
    if border:        cell.border        = border
    if fill:          cell.fill          = fill
    if number_format: cell.number_format = number_format

def _subtitle(ws, ref: str, value, merge_to: str | None = None):
    cell = ws[ref]
    cell.value = value
    _apply(cell, font=_FONT_SUBTITLE, alignment=_CENTER, border=_BORDER)
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")

def _period(bulan: str, tahun: str) -> tuple[str, str]:
    last = calendar.monthrange(int(tahun), int(bulan))[1]
    p1   = f"{tahun}-{bulan.zfill(2)}-{last:02d}"
    fmt  = f"{last:02d}-{bulan.zfill(2)}-{tahun}"
    return p1, fmt

def _pct(num: float, denom: float) -> float:
    return 0.0 if denom == 0 else round(num / denom * 100, 2)

def _exec_f6pd(emp_id: str, period1: str, mode: str) -> dict | None:
    try:
        row = db.session.execute(
            text("EXEC USP_F6PD_REPORT :eid, :p, :m"),
            {"eid": emp_id, "p": period1, "m": mode},
        ).mappings().first()
        return dict(row) if row else None
    except Exception:
        return None

def _build_summary_row(r: dict) -> dict:
    u1 = int(r.get("unit1") or 0)
    u2 = int(r.get("unit2") or 0)
    o1 = float(r.get("op1")  or 0)
    o2 = float(r.get("op2")  or 0)
    tu = u1 + u2
    to = o1 + o2
    return dict(unit1=u1, unit2=u2, tot_unit=tu, pd_unit=_pct(u2, tu),
                op1=o1,   op2=o2,   tot_op=to,   pd_op=_pct(o2, to))

def _fetch_all(bulan: str, tahun: str) -> dict:
    period1, formatted_date = _period(bulan, tahun)

    cmh_rows = db.session.execute(text("""
        SELECT DISTINCT a.SUPERIOR AS employee_id,
               ISNULL(b.FULLNAME, 'No CMH') AS fullname
        FROM app_usrmst a WITH(NOLOCK)
            LEFT JOIN APP_USRMST b WITH(NOLOCK) ON a.SUPERIOR = b.EMPLOYEE_ID
        WHERE a.ACCESSCAM = 'CMO' AND a.RECSTATUS = 'A'
    """)).mappings().all()

    summary_cmo = []
    for cmh in cmh_rows:
        cmh_id, cmh_name = cmh["employee_id"] or "", cmh["fullname"]

        r = _exec_f6pd(cmh_id, period1, "CMH")
        if r:
            row = _build_summary_row(r)
            summary_cmo.append({"fullname": cmh_name, "is_cmh": True, **row})

        cmo_rows = db.session.execute(text("""
            SELECT EMPLOYEE_ID, FULLNAME
            FROM app_usrmst WITH(NOLOCK)
            WHERE ACCESSCAM = 'CMO' AND SUPERIOR = :sup AND RECSTATUS = 'A'
        """), {"sup": cmh_id}).mappings().all()

        for cmo in cmo_rows:
            r = _exec_f6pd(cmo["EMPLOYEE_ID"], period1, "CMO")
            if r:
                row = _build_summary_row(r)
                summary_cmo.append({"fullname": cmo["FULLNAME"] or "", "is_cmh": False, **row})

    bulan_rows = db.session.execute(
        text("EXEC USP_F6PD_REPORT '', :p, 'BULAN'"), {"p": period1}
    ).mappings().all()
    summary_monthly = []
    for b in bulan_rows:
        row = _build_summary_row(b)
        summary_monthly.append({"bulan": b.get("bulan") or "", **row})

    raw_rows = db.session.execute(
        text("EXEC USP_F6PD_REPORT '', :p, 'RAW'"), {"p": period1}
    ).mappings().all()
    raw_data = [{
        "lease_no":   r.get("lease_no")   or "",
        "lessee_nm":  r.get("LESSEE_NM")  or "",
        "od":         int(r.get("od")     or 0),
        "outs_princ": float(r.get("outs_princ") or 0),
        "cmo":        r.get("CMO")        or "",
    } for r in raw_rows]

    return dict(formatted_date=formatted_date, period1=period1,
                summary_cmo=summary_cmo, summary_monthly=summary_monthly,
                raw_data=raw_data)

def _build_excel(d: dict, company: str = "") -> bytes:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Summary"

    ws1["A1"].value = company
    _apply(ws1["A1"], font=_FONT_TITLE, alignment=_LEFT)
    ws1["A2"].value = f"F6PD REPORT AS OF {d['formatted_date']}"
    _apply(ws1["A2"], font=_FONT_TITLE, alignment=_LEFT)

    ws1.merge_cells("A4:A5"); _subtitle(ws1, "A4", "CMO")
    ws1.merge_cells("B4:E4"); _subtitle(ws1, "B4", "By Unit")
    ws1.merge_cells("F4:I4"); _subtitle(ws1, "F4", "By Amount")
    for ref, val in [("B5","<= 30 DAYS"),("C5","> 30 DAYS"),("D5","GRAND TOTAL"),("E5","F6PD"),
                     ("F5","<= 30 DAYS"),("G5","> 30 DAYS"),("H5","GRAND TOTAL"),("I5","F6PD")]:
        _subtitle(ws1, ref, val)

    col_widths = {"A":25,"B":15,"C":15,"D":15,"E":15,"F":15,"G":15,"H":15,"I":15}
    for col, w in col_widths.items():
        ws1.column_dimensions[col].width = w

    row_i = 6
    for emp in d["summary_cmo"]:
        font = _FONT_CMH if emp["is_cmh"] else _FONT_DATA
        ws1[f"A{row_i}"].value = emp["fullname"]
        _apply(ws1[f"A{row_i}"], font=font, border=_BORDER)
        for col, key in [("B","unit1"),("C","unit2"),("D","tot_unit"),("E","pd_unit")]:
            c = ws1[f"{col}{row_i}"]
            c.value = emp[key];  _apply(c, font=font, border=_BORDER, alignment=_CENTER)
        for col, key in [("F","op1"),("G","op2"),("H","tot_op")]:
            c = ws1[f"{col}{row_i}"]
            c.value = emp[key];  _apply(c, font=font, border=_BORDER,
                                        alignment=_CENTER, number_format=_NUM_FMT)
        c = ws1[f"I{row_i}"]
        c.value = emp["pd_op"]; _apply(c, font=font, border=_BORDER, alignment=_CENTER)
        row_i += 1

    ws1.getattr = getattr
    
    j, k = row_i + 1, row_i + 2
    ws1.merge_cells(f"A{j}:A{k}"); _subtitle(ws1, f"A{j}", "MONTH-YEAR")
    ws1.merge_cells(f"B{j}:E{j}"); _subtitle(ws1, f"B{j}", "By Unit")
    ws1.merge_cells(f"F{j}:I{j}"); _subtitle(ws1, f"F{j}", "By Amount")
    for ref, val in [(f"B{k}","<= 30 DAYS"),(f"C{k}","> 30 DAYS"),(f"D{k}","GRAND TOTAL"),(f"E{k}","F6PD"),
                     (f"F{k}","<= 30 DAYS"),(f"G{k}","> 30 DAYS"),(f"H{k}","GRAND TOTAL"),(f"I{k}","F6PD")]:
        _subtitle(ws1, ref, val)

    row_l = k + 1
    for m in d["summary_monthly"]:
        ws1[f"A{row_l}"].value = m["bulan"]
        _apply(ws1[f"A{row_l}"], font=_FONT_DATA, border=_BORDER)
        for col, key in [("B","unit1"),("C","unit2"),("D","tot_unit"),("E","pd_unit")]:
            c = ws1[f"{col}{row_l}"]
            c.value = m[key]; _apply(c, font=_FONT_DATA, border=_BORDER, alignment=_CENTER)
        for col, key in [("F","op1"),("G","op2"),("H","tot_op")]:
            c = ws1[f"{col}{row_l}"]
            c.value = m[key]; _apply(c, font=_FONT_DATA, border=_BORDER,
                                     alignment=_CENTER, number_format=_NUM_FMT)
        c = ws1[f"I{row_l}"]
        c.value = m["pd_op"]; _apply(c, font=_FONT_DATA, border=_BORDER, alignment=_CENTER)
        row_l += 1

    ws2 = wb.create_sheet("Raw Data")

    ws2["A1"].value = company
    _apply(ws2["A1"], font=_FONT_TITLE, alignment=_LEFT)
    ws2["A2"].value = f"F6PD REPORT AS OF {d['formatted_date']}"
    _apply(ws2["A2"], font=_FONT_TITLE, alignment=_LEFT)

    for ref, val in [("A4","Contract No."),("B4","Customer Name"),
                     ("C4","OD Days"),("D4","Outstanding Principal"),("E4","CMO")]:
        _subtitle(ws2, ref, val)

    for col, w in {"A":25,"B":30,"C":12,"D":20,"E":20}.items():
        ws2.column_dimensions[col].width = w

    row_m = 5
    for r in d["raw_data"]:
        ws2[f"A{row_m}"].value = r["lease_no"]
        ws2[f"B{row_m}"].value = r["lessee_nm"]
        ws2[f"C{row_m}"].value = r["od"]
        ws2[f"D{row_m}"].value = r["outs_princ"]
        ws2[f"E{row_m}"].value = r["cmo"]
        _apply(ws2[f"A{row_m}"], font=_FONT_DATA, border=_BORDER)
        _apply(ws2[f"B{row_m}"], font=_FONT_DATA, border=_BORDER)
        _apply(ws2[f"C{row_m}"], font=_FONT_DATA, border=_BORDER, alignment=_CENTER)
        _apply(ws2[f"D{row_m}"], font=_FONT_DATA, border=_BORDER,
               alignment=_CENTER, number_format=_NUM_FMT)
        _apply(ws2[f"E{row_m}"], font=_FONT_DATA, border=_BORDER)
        row_m += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    return buf.read()

@f6pd_bp.route("/F6PD/data")
@login_required
def get_data():
    bulan = request.args.get("bulan_awal", "").strip()
    tahun = request.args.get("tahun_awal", "").strip()
    if not bulan or not tahun:
        return jsonify({"error": "bulan_awal and tahun_awal required"}), 400
    try:
        return jsonify(_fetch_all(bulan, tahun))
    except Exception:
        current_app.logger.exception("F6PD data error")
        return jsonify({"error": "Failed to load data"}), 500

@f6pd_bp.route("/F6PD/excel")
@login_required
def export_excel():
    bulan = request.args.get("bulan_awal", "").strip()
    tahun = request.args.get("tahun_awal", "").strip()
    if not bulan or not tahun:
        return jsonify({"error": "bulan_awal and tahun_awal required"}), 400
    try:
        d = _fetch_all(bulan, tahun)
        xlsx = _build_excel(d)
        resp = make_response(xlsx)
        resp.headers["Content-Type"] = \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        resp.headers["Content-Disposition"] = 'attachment; filename="F6PDAging.xlsx"'
        return resp
    except Exception:
        current_app.logger.exception("F6PD Excel error")
        return jsonify({"error": "Failed to generate Excel"}), 500