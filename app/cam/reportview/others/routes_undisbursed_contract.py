from flask import Blueprint, jsonify, request, current_app, make_response
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

undisburse_contract_bp = Blueprint("undisburse_contract", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT  = Alignment(horizontal="right",  vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")
_FONT_TITLE  = Font(name="Tahoma", size=14, bold=True)
_FONT_HEADER = Font(name="Tahoma", size=10, bold=True)
_FONT_DATA   = Font(name="Tahoma", size=10)
_FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")

def _apply(cell, font=None, alignment=None, border=None, fill=None):
    if font:      cell.font      = font
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border
    if fill:      cell.fill      = fill

_QUERY = text("""
    SELECT
        d.LESSEE_NM            AS cust_name,
        e.employee_fullname    AS emp_name,
        j.BRANCH_NAME          AS branch,
        h.BRAND_NM             AS brand,
        i.TYPE_NM              AS type_nm,
        b.LEASE_NO             AS lease_no,
        a.APPL_NO              AS appl_no,
        k.Net_Finance,
        f.AdmDate
    FROM cafcont a WITH(NOLOCK)
        INNER JOIN clfcont              b WITH(NOLOCK) ON a.APPL_NO   = b.APPL_NO
        LEFT  JOIN clfless              d WITH(NOLOCK) ON b.LESSEE_NO = d.LESSEE_NO
        LEFT  JOIN EMPLOYEE_DATA        e WITH(NOLOCK) ON a.cmo       = e.employee_ID
        LEFT  JOIN MCMOCAMCOMPLETELIST  f WITH(NOLOCK) ON a.APPL_NO   = f.APPL_NO
        LEFT  JOIN CLFEQUIP             g WITH(NOLOCK) ON b.LEASE_NO  = g.LEASE_NO AND g.FLGDEL = '0'
        LEFT  JOIN UMFBRAND             h WITH(NOLOCK) ON g.BRAND     = h.BRAND
        LEFT  JOIN UMFTYPE              i WITH(NOLOCK) ON g.TYPE      = i.TYPE AND g.BRAND = i.BRAND
        LEFT  JOIN UMFBRANCH            j WITH(NOLOCK) ON b.BRANCH_CD = j.BRANCH_CD
        LEFT  JOIN V_NetFinance         k WITH(NOLOCK) ON a.APPL_NO   = k.APPL_NO
    WHERE b.EXECUTION IS NULL
        AND b.FLGDEL = '0'
        AND a.DateTrans IS NOT NULL
        AND (a.branch_cd = :branch_cd OR :branch_cd = '000')

    UNION

    SELECT
        d.LESSEE_NM            AS cust_name,
        e.employee_fullname    AS emp_name,
        j.BRANCH_NAME          AS branch,
        h.BRAND_NM             AS brand,
        i.TYPE_NM              AS type_nm,
        ''                     AS lease_no,
        a.APPL_NO              AS appl_no,
        k.Net_Finance,
        b.AdmDate
    FROM cafcont a WITH(NOLOCK)
        INNER JOIN MCMOCAMCOMPLETELIST  b WITH(NOLOCK) ON a.APPL_NO = b.APPL_NO
        LEFT  JOIN cafless               d WITH(NOLOCK) ON a.APLESS  = d.APLESS
        LEFT  JOIN EMPLOYEE_DATA         e WITH(NOLOCK) ON a.cmo     = e.employee_ID
        LEFT  JOIN CaFEQUIP              g WITH(NOLOCK) ON a.APPL_NO = g.APPL_NO
        LEFT  JOIN UMFBRAND              h WITH(NOLOCK) ON g.BRAND   = h.BRAND
        LEFT  JOIN UMFTYPE               i WITH(NOLOCK) ON g.TYPE    = i.TYPE AND g.BRAND = i.BRAND
        LEFT  JOIN UMFBRANCH             j WITH(NOLOCK) ON a.BRANCH_CD = j.BRANCH_CD
        LEFT  JOIN V_NetFinance          k WITH(NOLOCK) ON a.APPL_NO = k.APPL_NO
    WHERE b.AdmSts = '1'
        AND a.DateTrans IS NULL
        AND (a.branch_cd = :branch_cd OR :branch_cd = '000')
        
    ORDER BY AdmDate
""")

def _serialize_rows(rows) -> list[dict]:
    result = []
    for r in rows:
        adm = r["AdmDate"]
        result.append({
            "appl_no":    r["appl_no"]   or "",
            "cust_name":  r["cust_name"] or "",
            "brand":      r["brand"]     or "",
            "type_nm":    r["type_nm"]   or "",
            "net_finance": float(r["Net_Finance"]) if r["Net_Finance"] is not None else None,
            "emp_name":   r["emp_name"]  or "",
            "branch":     r["branch"]    or "",
            "lease_no":   r["lease_no"]  or "",
            "adm_date":   adm.strftime("%Y-%m-%d") if adm else None,
        })
    return result

def _build_excel(rows: list[dict], branch_name: str) -> bytes:
    now = datetime.now().strftime("%d-%m-%Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "UndisburseContract"

    ws.merge_cells("A1:H1")
    ws["A1"].value = "List of Undisburse Contract"
    _apply(ws["A1"], font=_FONT_TITLE, alignment=_CENTER)
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:H2")
    ws["A2"].value = f"Branch : {branch_name}"
    _apply(ws["A2"], font=_FONT_DATA, alignment=_LEFT)

    ws.merge_cells("A3:H3")
    ws["A3"].value = f"Date   : {now}"
    _apply(ws["A3"], font=_FONT_DATA, alignment=_LEFT)

    headers = [
        "No.", "CAM No.", "Customer Name", "Collateral",
        "Total Net Finance", "CMO (Branch)", "Contract No.", "Last Approval Date",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=col, value=h)
        _apply(c, font=_FONT_HEADER, fill=_FILL_HEADER, border=_BORDER, alignment=_CENTER)

    tot = 0.0
    for idx, r in enumerate(rows):
        r_num = idx + 6
        adm   = r["adm_date"]
        adm_display = (
            datetime.strptime(adm, "%Y-%m-%d").strftime("%d %B %Y") if adm else ""
        )
        nf         = r["net_finance"] or 0.0
        tot       += nf
        collateral = f"{r['brand']} {r['type_nm']}".strip()
        cmo_branch = f"{r['emp_name']} ({r['branch']})"

        row_vals = [
            idx + 1, r["appl_no"], r["cust_name"], collateral,
            nf, cmo_branch, r["lease_no"], adm_display,
        ]
        alignments = [
            _CENTER, _LEFT, _LEFT, _LEFT,
            _RIGHT,  _LEFT, _LEFT, _LEFT,
        ]
        for col, (val, aln) in enumerate(zip(row_vals, alignments), 1):
            c = ws.cell(row=r_num, column=col, value=val)
            _apply(c, font=_FONT_DATA, border=_BORDER, alignment=aln)

    total_row = len(rows) + 6
    ws.merge_cells(f"A{total_row}:D{total_row}")
    c = ws[f"A{total_row}"]
    c.value = "TOTAL"
    _apply(c, font=_FONT_HEADER, border=_BORDER, alignment=_RIGHT)

    c = ws[f"E{total_row}"]
    c.value = tot
    _apply(c, font=_FONT_HEADER, border=_BORDER, alignment=_RIGHT)

    for col in range(6, 9):
        _apply(ws.cell(row=total_row, column=col), border=_BORDER)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    return buf.read()

@undisburse_contract_bp.route("/UndisbursedContract/branches")
@login_required
def get_branches():
   
    try:
        rows = db.session.execute(
            text("SELECT BRANCH_CD, BRANCH_NAME FROM umfbranch WITH(NOLOCK) ORDER BY BRANCH_NAME")
        ).mappings().all()
        return jsonify([
            {"branch_cd": r["BRANCH_CD"], "branch_name": r["BRANCH_NAME"]}
            for r in rows
        ])
    except Exception:
        current_app.logger.exception("Branches fetch error")
        return jsonify({"error": "Failed to load branches"}), 500

@undisburse_contract_bp.route("/UndisbursedContract/data")
@login_required
def get_data():
   
    branch_cd = request.args.get("branch_cd", "").strip()
    if not branch_cd:
        return jsonify({"error": "branch_cd is required"}), 400

    try:
        rows = db.session.execute(_QUERY, {"branch_cd": branch_cd}).mappings().all()
        return jsonify(_serialize_rows(rows))
    except Exception:
        current_app.logger.exception("Undisburse Contract data fetch error")
        return jsonify({"error": "Failed to load data"}), 500

@undisburse_contract_bp.route("/UndisbursedContract/excel")
@login_required
def export_excel():
    
    branch_cd = request.args.get("branch_cd", "").strip()
    if not branch_cd:
        return jsonify({"error": "branch_cd is required"}), 400

    try:
        branch_row = db.session.execute(
            text("SELECT BRANCH_NAME FROM UMFBRANCH WITH(NOLOCK) WHERE BRANCH_CD = :cd"),
            {"cd": branch_cd},
        ).mappings().first()
        branch_name = branch_row["BRANCH_NAME"] if branch_row else branch_cd

        rows = db.session.execute(_QUERY, {"branch_cd": branch_cd}).mappings().all()
        xlsx_bytes = _build_excel(_serialize_rows(rows), branch_name)

        resp = make_response(xlsx_bytes)
        resp.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp.headers["Content-Disposition"] = (
            'attachment; filename="UndisburseContract.xlsx"'
        )
        return resp

    except Exception:
        current_app.logger.exception("Undisburse Contract Excel export error")
        return jsonify({"error": "Failed to generate Excel"}), 500