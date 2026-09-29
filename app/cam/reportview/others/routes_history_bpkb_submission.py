from flask import Blueprint, jsonify, request, current_app, make_response
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text
import io, calendar
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

history_bpkb_submission_bp = Blueprint("history_bpkb_submission", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_FONT_TITLE  = Font(name="Tahoma", size=14, bold=True)
_FONT_HEADER = Font(name="Tahoma", size=10, bold=True)
_FONT_DATA   = Font(name="Tahoma", size=10)
_FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")

def _apply(cell, font=None, alignment=None, border=None, fill=None):
    if font:      cell.font      = font
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border
    if fill:      cell.fill      = fill

_MONTHS_ID = ["","Januari","Februari","Maret","April","Mei","Juni",
               "Juli","Agustus","September","Oktober","November","Desember"]

_HEADERS = {
    "0": ["No.","Dealer/ Showroom Name","Customer Name","Contract No.","Branch","Disbursement Date","Received Date (by CA)","Status","Remark"],
    "1": ["No.","Dealer/ Showroom Name","Customer Name","Contract No.","Disbursement Date","Received Date (by CA)","Status","Remark"],
    "2": ["No.","Customer Name","Contract No.","Branch","Disbursement Date","Received Date (by CA)","Status","Remark"],
}

_CENTER_COLS = {
    "0": {0, 3, 5, 6},
    "1": {0, 3, 4, 5},
    "2": {0, 2, 4, 5},
}

def _row_values(r: dict, mode: str) -> list:
    exec_  = r["execution"]  or ""
    recv_  = r["status_dt"]  or ""
    if mode == "0":
        return [r["dealer_name"], r["lessee_nm"], r["lease_no"], r["branch_name"], exec_, recv_, r["status"], r["remark"]]
    if mode == "1":
        return [r["dealer_name"], r["lessee_nm"], r["lease_no"], exec_, recv_, r["status"], r["remark"]]
    return [r["lessee_nm"], r["lease_no"], r["branch_name"], exec_, recv_, r["status"], r["remark"]]

_BASE_SQL = """
    SELECT d.NAME           AS dealer_name,
           c.LESSEE_NM,
           a.LEASE_NO,
           e.BRANCH_NAME,
           a.EXECUTION,
           g.Status_Dt,
           CASE WHEN DATEDIFF(d, a.EXECUTION, g.Status_Dt) <= 90
                THEN 'Not Due' ELSE 'Due' END AS Status,
           CASE WHEN DATEDIFF(d, a.EXECUTION, g.Status_Dt) > 90
                THEN CONCAT(DATEDIFF(d, a.EXECUTION, g.Status_Dt)-90,' ','Days')
                ELSE '' END AS Remark
    FROM clfcont a WITH(NOLOCK)
        INNER JOIN clfequip b WITH(NOLOCK) ON a.LEASE_NO = b.LEASE_NO AND b.FLGDEL = 0
        INNER JOIN clfless c WITH(NOLOCK) ON a.LESSEE_NO = c.LESSEE_NO
        INNER JOIN clfsupp d WITH(NOLOCK) ON b.supp = d.SUPP
        INNER JOIN UMFBRANCH e WITH(NOLOCK) ON a.BRANCH_CD = e.BRANCH_CD
        LEFT JOIN (
            SELECT cafequip_id, MIN(movement_id) AS id
            FROM Collateral_Movement WITH(NOLOCK)
            WHERE status_id = 'RC' AND flag_del = 0
            GROUP BY cafequip_id
        ) f ON b.CLFEQUIP_ID = f.Cafequip_Id
        LEFT JOIN Collateral_Movement g WITH(NOLOCK) ON f.id = g.Movement_Id
    WHERE g.Status_Dt BETWEEN :from1 AND :to1
        AND a.EXECUTION IS NOT NULL
"""

_EXTRA_WHERE = {
    "0": "",
    "1": " AND a.BRANCH_CD = :branch_cd",
    "2": " AND b.SUPP = :supp",
}

def _serialize(rows) -> list[dict]:
    def fmt(d) -> str | None:
        if d is None: return None
        return d.strftime("%d-%b-%y")

    result = []
    for r in rows:
        result.append({
            "dealer_name": r["dealer_name"] or "",
            "lessee_nm":   r["LESSEE_NM"]   or "",
            "lease_no":    r["LEASE_NO"]     or "",
            "branch_name": r["BRANCH_NAME"]  or "",
            "execution":   fmt(r["EXECUTION"]),
            "status_dt":   fmt(r["Status_Dt"]),
            "status":      r["Status"]       or "",
            "remark":      r["Remark"]       or "",
        })
    return result

def _date_range(bulan_awal, tahun_awal, bulan_akhir, tahun_akhir) -> tuple[str, str]:
    from1 = f"{tahun_awal}-{str(bulan_awal).zfill(2)}-01"
    last  = calendar.monthrange(int(tahun_akhir), int(bulan_akhir))[1]
    to1   = f"{tahun_akhir}-{str(bulan_akhir).zfill(2)}-{last:02d}"
    return from1, to1

def _build_excel(rows: list[dict], mode: str, period: str, subjudul: str) -> bytes:
    headers = _HEADERS[mode]
    center  = _CENTER_COLS[mode]
    now_str = datetime.now().strftime("%d-%m-%Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "HistoryBPKBSubmission"

    n_cols = len(headers)
    last_col = get_column_letter(n_cols)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"].value = "HISTORY OF BPKB SUBMISSION BY DEALER/ SHOWROOM"
    _apply(ws["A1"], font=_FONT_TITLE, alignment=_CENTER)
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"].value = f"Period Received Date (by CA) : {period}"
    _apply(ws["A2"], font=_FONT_DATA, alignment=_CENTER)

    ws.merge_cells(f"A3:{last_col}3")
    ws["A3"].value = f"Print by : {subjudul}"
    _apply(ws["A3"], font=_FONT_DATA, alignment=Alignment(horizontal="left", vertical="center"))

    ws.merge_cells(f"A4:{last_col}4")
    ws["A4"].value = f"Date : {now_str}"
    _apply(ws["A4"], font=_FONT_DATA, alignment=Alignment(horizontal="left", vertical="center"))

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=6, column=col, value=h)
        _apply(c, font=_FONT_HEADER, fill=_FILL_HEADER, border=_BORDER, alignment=_CENTER)

    for idx, r in enumerate(rows):
        r_num = idx + 7
        vals  = [idx + 1] + _row_values(r, mode)
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r_num, column=col, value=val)
            aln = _CENTER if (col - 1) in center else _LEFT
            _apply(c, font=_FONT_DATA, border=_BORDER, alignment=aln)

    col_widths = {"No.": 5, "Dealer/ Showroom Name": 28, "Customer Name": 28,
                  "Contract No.": 14, "Branch": 16,
                  "Disbursement Date": 16, "Received Date (by CA)": 18,
                  "Status": 12, "Remark": 16}
    for col, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = col_widths.get(h, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

@history_bpkb_submission_bp.route("/HistoryBPKBSubmission/branches")
@login_required
def get_branches():

    try:
        rows = db.session.execute(
            text("""
                SELECT BRANCH_CD, BRANCH_NAME FROM UMFBRANCH WITH(NOLOCK)
                WHERE BRANCH_CD NOT IN ('000','100')
                ORDER BY BRANCH_CD
            """)
        ).mappings().all()
        return jsonify([{"branch_cd": r["BRANCH_CD"], "branch_name": r["BRANCH_NAME"]} for r in rows])
    except Exception:
        current_app.logger.exception("Branches fetch error")
        return jsonify({"error": "Failed to load branches"}), 500

@history_bpkb_submission_bp.route("/HistoryBPKBSubmission/dealers")
@login_required
def get_dealers():
    
    q = request.args.get("q", "").strip()
    try:
        rows = db.session.execute(
            text("""
                SELECT TOP 200 SUPP, NAME, NICKNAME, ADDRESS, PHONE, CONTACT
                FROM CLFSUPP WITH(NOLOCK)
                WHERE Supp_Type <> 'KR'
                    AND (NICKNAME LIKE :q OR NAME LIKE :q)
                ORDER BY NAME
            """),
            {"q": f"%{q}%"},
        ).mappings().all()
        return jsonify([{
            "supp":     r["SUPP"]     or "",
            "name":     r["NAME"]     or "",
            "nickname": r["NICKNAME"] or "",
            "address":  r["ADDRESS"]  or "",
            "phone":    r["PHONE"]    or "",
            "contact":  r["CONTACT"]  or "",
        } for r in rows])
    except Exception:
        current_app.logger.exception("Dealers fetch error")
        return jsonify({"error": "Failed to load dealers"}), 500

@history_bpkb_submission_bp.route("/HistoryBPKBSubmission/data")
@login_required
def get_data():
   
    bulan_awal  = request.args.get("bulan_awal",  "").strip()
    tahun_awal  = request.args.get("tahun_awal",  "").strip()
    bulan_akhir = request.args.get("bulan_akhir", "").strip()
    tahun_akhir = request.args.get("tahun_akhir", "").strip()
    print_by    = request.args.get("print_by",    "0").strip()
    branch_cd   = request.args.get("branch_cd",   "").strip()
    supp        = request.args.get("supp",         "").strip()

    if not all([bulan_awal, tahun_awal, bulan_akhir, tahun_akhir]):
        return jsonify({"error": "Period params required"}), 400

    from1, to1 = _date_range(bulan_awal, tahun_awal, bulan_akhir, tahun_akhir)
    params = {"from1": from1, "to1": to1}

    if print_by == "1":
        params["branch_cd"] = branch_cd
    elif print_by == "2":
        params["supp"] = supp

    sql = _BASE_SQL + _EXTRA_WHERE.get(print_by, "") + " ORDER BY a.EXECUTION DESC"

    try:
        rows = db.session.execute(text(sql), params).mappings().all()
        return jsonify(_serialize(rows))
    except Exception:
        current_app.logger.exception("BPKB Submission data fetch error")
        return jsonify({"error": "Failed to load data"}), 500

@history_bpkb_submission_bp.route("/HistoryBPKBSubmission/excel")
@login_required
def export_excel():
   
    bulan_awal  = request.args.get("bulan_awal",  "").strip()
    tahun_awal  = request.args.get("tahun_awal",  "").strip()
    bulan_akhir = request.args.get("bulan_akhir", "").strip()
    tahun_akhir = request.args.get("tahun_akhir", "").strip()
    print_by    = request.args.get("print_by",    "0").strip()
    branch_cd   = request.args.get("branch_cd",   "").strip()
    supp        = request.args.get("supp",         "").strip()
    branch_nm   = request.args.get("branch_nm",   "").strip()
    dealer_nm   = request.args.get("dealer_nm",   "").strip()

    from1, to1 = _date_range(bulan_awal, tahun_awal, bulan_akhir, tahun_akhir)
    params = {"from1": from1, "to1": to1}
    if print_by == "1": params["branch_cd"] = branch_cd
    if print_by == "2": params["supp"] = supp

    period = f"{_MONTHS_ID[int(bulan_awal)]} {tahun_awal} - {_MONTHS_ID[int(bulan_akhir)]} {tahun_akhir}"
    subjudul = (
        "Consolidate"                if print_by == "0" else
        f"Branch - {branch_nm}"      if print_by == "1" else
        f"Dealer - {dealer_nm}"
    )

    sql = _BASE_SQL + _EXTRA_WHERE.get(print_by, "") + " ORDER BY a.EXECUTION DESC"

    try:
        rows = db.session.execute(text(sql), params).mappings().all()
        xlsx = _build_excel(_serialize(rows), print_by, period, subjudul)

        resp = make_response(xlsx)
        resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        resp.headers["Content-Disposition"] = 'attachment; filename="HistoryOfBPKBSubmission.xlsx"'
        return resp
    except Exception:
        current_app.logger.exception("BPKB Submission Excel export error")
        return jsonify({"error": "Failed to generate Excel"}), 500