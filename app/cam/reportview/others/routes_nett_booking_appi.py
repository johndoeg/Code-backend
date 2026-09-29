from flask import Blueprint, jsonify, request, current_app, make_response
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

nett_booking_appi_bp = Blueprint("nett_booking_appi", __name__, url_prefix="/CAM/Others")

_THIN   = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT  = Alignment(horizontal="right",  vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")
_FONT_TITLE  = Font(name="Tahoma", size=12, bold=True)
_FONT_HEADER = Font(name="Tahoma", size=10, bold=True)
_FONT_DATA   = Font(name="Tahoma", size=10)
_FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")

def _apply(cell, font=None, alignment=None, border=None, fill=None):
    if font:      cell.font      = font
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border
    if fill:      cell.fill      = fill

def _build_excel(row: dict, date_from: str) -> bytes:
    d     = datetime.strptime(date_from, "%Y-%m-%d")
    month = d.strftime("%B %Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "NettBookingAPPI"

    ws.merge_cells("A1:C1")
    ws["A1"].value = "NETT BOOKING APPI"
    _apply(ws["A1"], font=_FONT_TITLE, alignment=_CENTER)
    ws.row_dimensions[1].height = 22

    headers = [
        ("A2", "C2",  " FORMULIR1 ", " NETT BOOKING KONVENSIONAL ",         month),
        ("A7", "C7",  " FORMULIR3 ", " NETT BOOKING BERDASARKAN KEGIATAN USAHA ", " JAWABAN DALAM JUTAAN RP "),
        ("A11","C11", " FORMULIR9 ", " FORMULIR EXISTING CUSTOMER ",         " JAWABAN "),
        ("A14","D14", " KBLBB sesuai Periode Laporan Bulanan ", " UNIT ", " NOMINAL "),
        ("A18","D18", " KBLBB sesuai Outstanding ",             " UNIT ", " NOMINAL "),
    ]
    for a_ref, _, *vals in headers:
        row_num = int("".join(filter(str.isdigit, a_ref)))
        for col_idx, val in enumerate(vals, start=1):
            col_letter = get_column_letter(col_idx)
            c = ws[f"{col_letter}{row_num}"]
            c.value = val
            _apply(c, font=_FONT_HEADER, border=_BORDER, fill=_FILL_HEADER, alignment=_CENTER)

    for ref in ("D14", "D18"):
        ws[ref].value = " KONTRAK "
        _apply(ws[ref], font=_FONT_HEADER, border=_BORDER, fill=_FILL_HEADER, alignment=_CENTER)

    data_rows = [
        (3,  "1A", "Total Nett Booking Bulanan MOBIL BENSIN BARU KONVENSIONAL (tanpa bunga)",    row.get("tot_fin1a")),
        (4,  "1B", "Total Nett Booking Bulanan MOBIL BENSIN BEKAS KONVENSIONAL (tanpa bunga)",   row.get("tot_fin1b")),
        (5,  "1C", "Total Nett Booking Bulanan MOBIL LISTRIK KONVENSIONAL (tanpa bunga)",        row.get("tot_fin1c")),
        (6,  "1D", "Total Nett Booking Bulanan MOBIL HYBRID KONVENSIONAL (tanpa bunga)",         row.get("tot_fin1d")),
        (8,  "3A", "Total Nett Booking Bulanan berdasarkan PEMBIAYAAN INVESTASI",                row.get("tot_fin3a")),
        (9,  "3B", "Total Nett Booking Bulanan berdasarkan PEMBIAYAAN MODAL KERJA",              row.get("tot_fin3b")),
        (10, "3C", "Total Nett Booking Bulanan berdasarkan PEMBIAYAAN MULTIGUNA",                row.get("tot_fin3c")),
        (12, "9A", "Existing Customer",                                                          row.get("tot_cust")),
    ]
    
    for r_num, code, label, val in data_rows:
        ws[f"A{r_num}"].value = code
        _apply(ws[f"A{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_CENTER)
        ws[f"B{r_num}"].value = label
        _apply(ws[f"B{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_LEFT)
        ws[f"C{r_num}"].value = val
        _apply(ws[f"C{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_RIGHT)

    kblbb_rows = [
        (15, "Murni Listrik", row.get("tot_unit_listrik_period"), row.get("tot_fin_listrik_period"), row.get("tot_unit_listrik_period")),
        (16, "Hybrid",        row.get("tot_unit_hybrid_period"),  row.get("tot_fin_hybrid_period"),  row.get("tot_unit_hybrid_period")),
        (19, "Murni Listrik", row.get("tot_unit_listrik_out"),    row.get("tot_fin_listrik_out"),    row.get("tot_unit_listrik_out")),
        (20, "Hybrid",        row.get("tot_unit_hybrid_out"),     row.get("tot_fin_hybrid_out"),     row.get("tot_unit_hybrid_out")),
    ]
    
    for r_num, label, unit, nominal, kontrak in kblbb_rows:
        ws[f"A{r_num}"].value = label;   _apply(ws[f"A{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_LEFT)
        ws[f"B{r_num}"].value = unit;    _apply(ws[f"B{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_RIGHT)
        ws[f"C{r_num}"].value = nominal; _apply(ws[f"C{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_RIGHT)
        ws[f"D{r_num}"].value = kontrak; _apply(ws[f"D{r_num}"], font=_FONT_DATA, border=_BORDER, alignment=_RIGHT)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    return buf.read()

@nett_booking_appi_bp.route("/NetBookingAPPI/nett-booking")
@login_required
def get_nett_booking():
   
    date_from = request.args.get("date_from", "").strip()
    if not date_from:
        return jsonify({"error": "date_from is required"}), 400

    try:
        datetime.strptime(date_from, "%Y-%m-%d")
    except ValueError:
        return jsonify({"error": "date_from must be YYYY-MM-DD"}), 400

    try:
        rows = db.session.execute(
            text("EXEC USP_Nett_Booking_APPI :d"), {"d": date_from}
        ).mappings().all()

        if not rows:
            return jsonify({}), 200

        r = dict(rows[0])
        return jsonify({
            "tot_fin1a":               r.get("tot_fin1a"),
            "tot_fin1b":               r.get("tot_fin1b"),
            "tot_fin1c":               r.get("tot_fin1c"),
            "tot_fin1d":               r.get("tot_fin1d"),
            "tot_fin3a":               r.get("tot_fin3a"),
            "tot_fin3b":               r.get("tot_fin3b"),
            "tot_fin3c":               r.get("tot_fin3c"),
            "tot_cust":                r.get("tot_cust"),
            "tot_unit_listrik_period": r.get("tot_unit_listrik_period"),
            "tot_fin_listrik_period":  r.get("tot_fin_listrik_period"),
            "tot_unit_hybrid_period":  r.get("tot_unit_hybrid_period"),
            "tot_fin_hybrid_period":   r.get("tot_fin_hybrid_period"),
            "tot_unit_listrik_out":    r.get("tot_unit_listrik_out"),
            "tot_fin_listrik_out":     r.get("tot_fin_listrik_out"),
            "tot_unit_hybrid_out":     r.get("tot_unit_hybrid_out"),
            "tot_fin_hybrid_out":      r.get("tot_fin_hybrid_out"),
        })

    except Exception:
        current_app.logger.exception("Nett Booking APPI fetch error")
        return jsonify({"error": "Failed to load data"}), 500

@nett_booking_appi_bp.route("/NetBookingAPPI/nett-booking/excel")
@login_required
def export_excel():
  
    date_from = request.args.get("date_from", "").strip()
    if not date_from:
        return jsonify({"error": "date_from is required"}), 400

    try:
        datetime.strptime(date_from, "%Y-%m-%d")
    except ValueError:
        return jsonify({"error": "date_from must be YYYY-MM-DD"}), 400

    try:
        rows = db.session.execute(
            text("EXEC USP_Nett_Booking_APPI :d"), {"d": date_from}
        ).mappings().all()

        row = dict(rows[0]) if rows else {}
        xlsx_bytes = _build_excel(row, date_from)

        resp = make_response(xlsx_bytes)
        resp.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp.headers["Content-Disposition"] = 'attachment; filename="NettBookingAPPI.xlsx"'
        
        return resp

    except Exception:
        current_app.logger.exception("Nett Booking APPI Excel export error")
        return jsonify({"error": "Failed to generate Excel"}), 500