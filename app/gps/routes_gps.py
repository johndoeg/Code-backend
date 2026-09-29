from flask import Blueprint, jsonify, request, current_app, send_file
from app import db
from app.auth.decorators import login_required
from sqlalchemy import text

from werkzeug.utils import secure_filename
import openpyxl
import io

gps_bp = Blueprint('gps', __name__, url_prefix='/GPS')

ALLOWED_EXTENSIONS = {'xls', 'xlsx'}
MAX_FILE_SIZE = 25 * 1024 * 1024

VENDOR_CONFIG = {
    '2400001': {'start_row': 2, 'chasis_col': 10},
}
DEFAULT_VENDOR_CONFIG = {'start_row': 8, 'chasis_col': 5}

def _allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@gps_bp.route('/invoice/vendors')
@login_required
def get_gps_vendors():
    try:
        result = db.session.execute(text("""
            SELECT SUPP, NICKNAME
            FROM VENDOR
            WHERE Supp_Category = 'GPS' AND Flgdel = '0'
            ORDER BY SUPP
        """))
        vendors = [{"SUPP": row[0], "NICKNAME": row[1]} for row in result]
        return jsonify(vendors)
    except Exception:
        current_app.logger.exception("GPS vendors error")
        return jsonify([])

@gps_bp.route('/invoice')
@login_required
def get_invoice_summary():
    vendor = request.args.get('vendor')
    page  = request.args.get('page',  1,  type=int)
    limit = request.args.get('limit', 25, type=int)

    if not vendor:
        return jsonify({"error": "vendor is required"}), 400

    try:
        result = db.session.execute(
            text("EXEC USP_INVOICE_GPS :mode, :asof, :vendor"),
            {"mode": "Summary", "asof": "", "vendor": vendor},
        )
        all_rows = [
            {
                "asof":           str(row.tanggal)        if row.tanggal        is not None else "",
                "disb_month":     row.disb_month          or "",
                "tot_kontrak_inv": str(row.tot_kontrak_inv) if row.tot_kontrak_inv is not None else "0",
                "tot_kontrak_sys": str(row.tot_kontrak_sys) if row.tot_kontrak_sys is not None else "0",
                "vendor":         vendor,
            }
            for row in result
        ]

        total     = len(all_rows)
        start     = (page - 1) * limit
        paginated = all_rows[start:start + limit]

        return jsonify({"data": paginated, "total": total})

    except Exception:
        current_app.logger.exception("Invoice summary error")
        return jsonify({"error": "Failed to load invoice summary"}), 500

@gps_bp.route('/invoice/detail')
@login_required
def get_invoice_detail():
    asof   = request.args.get('asof')
    vendor = request.args.get('vendor')

    if not asof or not vendor:
        return jsonify({"error": "asof and vendor are required"}), 400

    try:
        result = db.session.execute(
            text("EXEC USP_INVOICE_GPS :mode, :asof, :vendor"),
            {"mode": "Detail", "asof": asof, "vendor": vendor},
        )
        details = [
            {
                "lease_no":   row.lease_no  or "",
                "lessee_nm":  row.LESSEE_NM or "",
                "execution":  str(row.EXECUTION) if row.EXECUTION else "",
                "invoice":    row.Invoice   or "",
                "system":     row.System    or "",
                "remark":     "No. kontrak" if row.Invoice != row.System else "",
            }
            for row in result
        ]
        return jsonify({"data": details})

    except Exception:
        current_app.logger.exception("Invoice detail error")
        return jsonify({"error": "Failed to load invoice details"}), 500

@gps_bp.route('/invoice/upload', methods=['POST'])
@login_required
def upload_invoice():
    try:
        file   = request.files.get('invoiceGPS')
        vendor = request.form.get('gps_vendor_upload', '').strip()

        if not file or file.filename == '':
            return jsonify({"success": False, "error": "No file uploaded"}), 400
        if not vendor:
            return jsonify({"success": False, "error": "Vendor is required"}), 400
        if not _allowed_file(file.filename):
            return jsonify({"success": False, "error": "Only .xls and .xlsx files are allowed"}), 400

        raw = file.read()
        if len(raw) > MAX_FILE_SIZE:
            return jsonify({"success": False, "error": "File size exceeds the 25 MB limit"}), 400

        cfg        = VENDOR_CONFIG.get(vendor, DEFAULT_VENDOR_CONFIG)
        start_row  = cfg['start_row']
        chasis_col = cfg['chasis_col']

        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active

        last_row = 0
        for row_idx, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
            if row[0] not in (None, ''):
                last_row = row_idx

        chasis_list = []
        for row_idx, row in enumerate(
            ws.iter_rows(min_col=chasis_col, max_col=chasis_col, values_only=True), start=1
        ):
            if row_idx < start_row or row_idx > last_row:
                continue
            val = row[0]
            if val not in (None, ''):
                chasis_list.append(str(val))

        wb.close()

        if not chasis_list:
            return jsonify({"success": False, "error": "No valid chassis data found in file"}), 400

        asof = request.form.get('asof', '').strip()
        if not asof:
            return jsonify({"success": False, "error": "asof date is required"}), 400

        db.session.execute(
            text("DELETE FROM INVOICE_GPS WHERE AsOf = :asof AND GPS_Vendor = :vendor"),
            {"asof": asof, "vendor": vendor},
        )
        for chasis in chasis_list:
            db.session.execute(
                text("INSERT INTO INVOICE_GPS (AsOf, GPS_Vendor, Chasis) VALUES (:asof, :vendor, :chasis)"),
                {"asof": asof, "vendor": vendor, "chasis": chasis},
            )
        db.session.commit()

        return jsonify({
            "success": True,
            "message": f"Uploaded {len(chasis_list)} records successfully.",
        })

    except Exception:
        db.session.rollback()
        current_app.logger.exception("Invoice upload error")
        return jsonify({"success": False, "error": "Upload failed. Please try again."}), 500