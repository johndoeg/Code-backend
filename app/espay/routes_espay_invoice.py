import io
import json
from datetime import datetime

import openpyxl
import requests
from openpyxl.utils import get_column_letter
from flask import Blueprint, current_app, jsonify, request, send_file, session
from sqlalchemy import text

from app import db
from app.auth.decorators import login_required

espay_bp = Blueprint('espay', __name__, url_prefix='/Espay')

ESPAY_BASE_URL = "https://ifc.sfintech.id:28888"
ESPAY_TIMEOUT = 30
ESPAY_OK_CODE = "0000"

_ELIGIBLE = """
    AND a.lease_no IN (SELECT lease_no FROM ESPAY_Member_Contract)
    AND a.LEASE_NO NOT IN (
        SELECT LEASE_NO FROM ESPAY_FromTagih
        WHERE BatchNo = '2' AND DATEDIFF(d, BatchDate, GETDATE()) <= 0
    )
"""

SQL_INSTALLMENT = f"""
    SELECT b.LESSEE_NM, d.period, CAST(d.nominal AS DECIMAL(20,0)) AS nominal,
           FORMAT(d.GIRODT, 'dd-MM-yyyy') AS GIRODT, '' AS overdue
    FROM CLFCONT_DATA a WITH(NOLOCK)
        INNER JOIN CLFLESS b WITH(NOLOCK) ON a.LESSEE_NO = b.LESSEE_NO
        LEFT JOIN (
            SELECT lease_no, MIN(period) AS jth_tmp
            FROM BGIRO WITH(NOLOCK)
            WHERE PAYMENT IS NULL AND NOMINAL > 0
            GROUP BY LEASE_NO
        ) c ON a.lease_no = c.LEASE_NO
        LEFT JOIN BGIRO d WITH(NOLOCK) ON a.lease_no = d.lease_no AND c.jth_tmp = d.period
    WHERE a.lease_no = :lease_no {_ELIGIBLE}
"""

SQL_PENALTY = f"""
    SELECT b.LESSEE_NM, '' AS period, CAST(c.penalty AS DECIMAL(20,0)) AS nominal,
           '' AS GIRODT, '' AS overdue
    FROM CLFCONT_DATA a WITH(NOLOCK)
        INNER JOIN CLFLESS b WITH(NOLOCK) ON a.LESSEE_NO = b.LESSEE_NO
        LEFT JOIN (
            SELECT lease_no, SUM(PENALTY) AS penalty FROM BGIRO WITH(NOLOCK) GROUP BY LEASE_NO
        ) c ON a.lease_no = c.LEASE_NO
    WHERE a.lease_no = :lease_no {_ELIGIBLE}
"""

SQL_OVERDUE = f"""
    SELECT b.LESSEE_NM, d.period, CAST(c.nominal AS DECIMAL(20,0)) AS nominal,
           FORMAT(d.GIRODT, 'dd-MM-yyyy') AS GIRODT, c.overdue
    FROM CLFCONT_DATA a WITH(NOLOCK)
        INNER JOIN CLFLESS b WITH(NOLOCK) ON a.LESSEE_NO = b.LESSEE_NO
        LEFT JOIN (
            SELECT lease_no, MIN(period) AS jth_tmp, COUNT(period) AS overdue,
                   SUM(NOMINAL) AS nominal
            FROM BGIRO WITH(NOLOCK)
            WHERE PAYMENT IS NULL AND STATUS = 'R' AND NOMINAL > 0
            GROUP BY LEASE_NO
        ) c ON a.lease_no = c.LEASE_NO
        LEFT JOIN BGIRO d WITH(NOLOCK) ON a.lease_no = d.lease_no AND c.jth_tmp = d.period
    WHERE a.lease_no = :lease_no {_ELIGIBLE}
"""

SQL_ET = f"""
    SELECT b.LESSEE_NM, '' AS period,
           FORMAT(c.ET_DT, 'dd-MM-yyyy') AS GIRODT,
           CAST(c.ET_PAY AS DECIMAL(20,0)) AS overdue,
           CAST(ISNULL(d.et_paid, 0) AS DECIMAL(20,0)) AS et_paid,
           CAST((c.ET_PAY - ISNULL(d.et_paid, 0)) AS DECIMAL(20,0)) AS nominal
    FROM CLFCONT_DATA a WITH(NOLOCK)
        INNER JOIN CLFLESS b WITH(NOLOCK) ON a.LESSEE_NO = b.LESSEE_NO
        INNER JOIN CLFET c WITH(NOLOCK) ON a.LEASE_NO = c.LEASE_NO
        LEFT JOIN (
            SELECT SUM(ISNULL(total_amount, 0)) AS et_paid, lease_no
            FROM ESPAY_Invoice WITH(NOLOCK)
            WHERE trans_type = 'ET' AND flagdel = '0'
            GROUP BY lease_no
        ) d ON a.LEASE_NO = d.lease_no
    WHERE a.lease_no = :lease_no {_ELIGIBLE}
"""

SQL_OTHERS = f"""
    SELECT b.LESSEE_NM, '' AS period, '' AS nominal, '' AS GIRODT, '' AS overdue
    FROM CLFCONT_DATA a WITH(NOLOCK)
        INNER JOIN CLFLESS b WITH(NOLOCK) ON a.LESSEE_NO = b.LESSEE_NO
    WHERE a.lease_no = :lease_no {_ELIGIBLE}
"""

CONTRACT_SQL = {
    'INSTALLMENT': SQL_INSTALLMENT,
    'MULTIPLE': SQL_INSTALLMENT,
    'PENALTY': SQL_PENALTY,
    'OVERDUE': SQL_OVERDUE,
    'ET': SQL_ET,
}

def _col(row, name, default=''):
    value = row._mapping.get(name, default)
    return default if value is None else value

def _num(value):
    return str(value) if value not in (None, '') else "0"

def _to_iso_date(value):
    value = (value or '').strip()
    if not value:
        return ''
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d-%b-%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return value

def _espay_post(path, payload):
    try:
        response = requests.post(
            f"{ESPAY_BASE_URL}{path}", data=payload, timeout=ESPAY_TIMEOUT
        )
    except requests.RequestException:
        current_app.logger.exception("ESPAY gateway unreachable: %s", path)
        return False, "Could not reach the ESPAY service. Please try again."

    try:
        body = response.json()
    except ValueError:
        current_app.logger.error("ESPAY non-JSON reply from %s: %s", path, response.text[:500])
        return False, response.text.strip() or "Unexpected response from the ESPAY service."

    if str(body.get('error_code')) == ESPAY_OK_CODE:
        return True, body.get('responseMessage') or "Successful"
    return False, body.get('responseMessage') or body.get('error_message') or "ESPAY request failed."

def _period_label(row):
    payment = (_col(row, 'Payment') or '').strip()
    period = _col(row, 'period')
    if payment == 'Installment':
        return str(period or '')
    if payment == 'Installment Overdue':
        return f"{period or ''} - {_col(row, 'period_to') or ''}"
    return ''

def _fetch_invoices(search_by, search_val, search_val_to):
    result = db.session.execute(
        text("EXEC USP_ESPAY_Get_Invoice :search_by, :search_val, :search_val_to"),
        {
            "search_by": search_by,
            "search_val": search_val,
            "search_val_to": search_val_to,
        },
    )
    return [
        {
            "id_data": _col(row, 'id'),
            "invoice_date": _col(row, 'InvoiceDate'),
            "contract_no": _col(row, 'lease_no'),
            "customer_name": _col(row, 'lessee_nm'),
            "tot_amt": _num(_col(row, 'tot_amt', None)),
            "expired_date": _col(row, 'ExpiredDate'),
            "payment_type": _col(row, 'Payment'),
            "period": _period_label(row),
            "invoice_status": _col(row, 'inv_status'),
            "payment_date": _col(row, 'PaymentDate'),
        }
        for row in result
    ]

def _combo_options(cat_value, exclude=None):
    exclude = exclude or []
    query = "SELECT VALUE, DESC_VALUE FROM MSTCOMBO WITH (NOLOCK) WHERE CAT_VALUE = :cat_value"
    params = {"cat_value": cat_value}
    if exclude:
        placeholders = ", ".join(f":ex{i}" for i in range(len(exclude)))
        query += f" AND VALUE NOT IN ({placeholders})"
        for i, val in enumerate(exclude):
            params[f"ex{i}"] = val

    result = db.session.execute(text(query), params)
    return [{"value": (row.VALUE or '').strip(), "label": row.DESC_VALUE} for row in result]

@espay_bp.route('/search-options')
@login_required
def get_search_options():
    try:
        return jsonify({
            "search_by": _combo_options("SEARCH_INVOICE", exclude=["4", "5"]),
            "payment_for": _combo_options("PAYMENT_FOR"),
            "invoice_status": _combo_options("INVOICE_STATUS"),
        })
    except Exception:
        current_app.logger.exception("Espay search-options error")
        return jsonify({"search_by": [], "payment_for": [], "invoice_status": []})

@espay_bp.route('/invoice')
@login_required
def search_invoices():
    search_by = request.args.get('search_by', 'All')
    search_val = request.args.get('search_val')
    search_val_to = request.args.get('search_val_to') or None

    try:
        return jsonify({"data": _fetch_invoices(search_by, search_val, search_val_to)})
    except Exception:
        current_app.logger.exception("Espay invoice search error")
        return jsonify({"data": [], "message": "Failed to load invoice data."}), 500

@espay_bp.route('/invoice/contract')
@login_required
def get_contract_data():
    lease_no = (request.args.get('lease_no') or '').strip()
    payment_for = (request.args.get('payment_for') or '').strip().upper()
    ref_lease_no = (request.args.get('ref_lease_no') or '').strip()

    if not lease_no:
        return jsonify({"error": "not_found"}), 404

    try:
        if payment_for == 'MULTIPLE' and ref_lease_no and ref_lease_no != lease_no:
            rows = db.session.execute(
                text("""
                    SELECT LEASE_NO, LESSEE_NO FROM CLFCONT_DATA WITH(NOLOCK)
                    WHERE LEASE_NO IN (:ref_lease_no, :lease_no)
                """),
                {"ref_lease_no": ref_lease_no, "lease_no": lease_no},
            ).fetchall()

            lessee_of = {(r.LEASE_NO or '').strip(): (r.LESSEE_NO or '').strip() for r in rows}
            if ref_lease_no not in lessee_of or lease_no not in lessee_of:
                return jsonify({"error": "not_found"}), 404
            if lessee_of[ref_lease_no] != lessee_of[lease_no]:
                return jsonify({"error": "different_customer"}), 200

        row = db.session.execute(
            text(CONTRACT_SQL.get(payment_for, SQL_OTHERS)), {"lease_no": lease_no}
        ).fetchone()

        if row is None:
            return jsonify({"error": "not_found"}), 404

        return jsonify({
            "customer_name": _col(row, 'LESSEE_NM'),
            "due_dt": _col(row, 'GIRODT'),
            "nominal": _num(_col(row, 'nominal', None)),
            "overdue": _num(_col(row, 'overdue', None)),
            "period": str(_col(row, 'period')),
            "et_paid": _num(_col(row, 'et_paid', None)),
        })
    except Exception:
        current_app.logger.exception("Espay contract lookup error")
        return jsonify({"error": "not_found"}), 404

@espay_bp.route('/invoice/overdue', methods=['POST'])
@login_required
def calculate_overdue():
    data = request.get_json(silent=True) or {}
    lease_no = (data.get('lease_no') or '').strip()

    try:
        od_period = int(data.get('od_period') or 0)
        period_from = int(data.get('period') or 0)
    except (TypeError, ValueError):
        return jsonify({"error": "not_found"}), 404

    if not lease_no or od_period <= 0:
        return jsonify({"error": "not_found"}), 404

    period_to = period_from + od_period - 1

    try:
        row = db.session.execute(
            text("""
                SELECT CAST(SUM(nominal) AS DECIMAL(20,0)) AS nominal
                FROM BGIRO WITH(NOLOCK)
                WHERE lease_no = :lease_no AND period BETWEEN :period_from AND :period_to
                GROUP BY lease_no
            """),
            {"lease_no": lease_no, "period_from": period_from, "period_to": period_to},
        ).fetchone()

        if row is None or row.nominal is None:
            return jsonify({"error": "not_found"}), 404
        return jsonify({"nominal": _num(row.nominal)})
    except Exception:
        current_app.logger.exception("Espay overdue calc error")
        return jsonify({"error": "not_found"}), 404

@espay_bp.route('/invoice/<int:id_data>', methods=['DELETE'])
@login_required
def delete_invoice(id_data):
    try:
        row = db.session.execute(
            text("""
                SELECT b.Member_Code
                FROM ESPAY_Invoice a WITH(NOLOCK)
                    INNER JOIN ESPAY_Member_Contract b WITH(NOLOCK) ON a.lease_no = b.Lease_No
                WHERE a.id = :id_data
            """),
            {"id_data": id_data},
        ).fetchone()

        if row is None or not row.Member_Code:
            return jsonify({"message": "Invoice not found."}), 404

        ok, message = _espay_post("/DeleteInvoice", {"member_code": row.Member_Code})
        return jsonify({"message": message}), (200 if ok else 502)
    except Exception:
        current_app.logger.exception("Espay delete error")
        return jsonify({"message": "Failed to delete invoice."}), 500

@espay_bp.route('/invoice', methods=['POST'])
@login_required
def create_invoice():
    data = request.get_json(silent=True) or {}

    lease_no = (data.get('lease_no') or '').strip()
    payment_for = (data.get('payment_for') or '').strip().upper()
    if not lease_no or not payment_for:
        return jsonify({"message": "Contract No. and Payment For are required."}), 400

    od_period = '' if payment_for == 'ET' else (data.get('od_period') or '')

    payload = {
        "lease_no": lease_no,
        "trans_type": payment_for,
        "period": str(data.get('period') or ''),
        "total_amount": str(data.get('tot_inv') or ''),
        "overdue": str(od_period),
        "et_dt": _to_iso_date(data.get('et_date')),
        "invoice_detail": json.dumps([]),
    }

    ok, message = _espay_post("/SendInvoiceForm", payload)
    return jsonify({"message": message}), (200 if ok else 502)

@espay_bp.route('/invoice/multiple', methods=['POST'])
@login_required
def create_multiple_invoices():
    data = request.get_json(silent=True) or {}
    rows = data.get('rows') or []

    if not rows:
        return jsonify({"message": "At least one contract is required."}), 400

    total = 0
    invoice_detail = []
    for item in rows:
        try:
            amount = float(item.get('total') or 0)
        except (TypeError, ValueError):
            amount = 0
        total += amount
        invoice_detail.append({
            "lease_no": (item.get('lease_no') or '').strip(),
            "period": str(item.get('period') or ''),
            "total": str(int(amount)),
        })

    last = rows[-1]

    payload = {
        "lease_no": (last.get('lease_no') or '').strip(),
        "trans_type": "MULTIPLE",
        "period": str(last.get('period') or ''),
        "total_amount": str(int(total)),
        "overdue": "",
        "et_dt": "",
        "invoice_detail": json.dumps(invoice_detail),
    }

    ok, message = _espay_post("/SendInvoiceForm", payload)
    return jsonify({"message": message}), (200 if ok else 502)

@espay_bp.route('/invoice/export')
@login_required
def export_invoices():
    search_by = request.args.get('search_by', 'All')
    search_val = request.args.get('search_val')
    search_val_to = request.args.get('search_val_to') or None

    try:
        rows = _fetch_invoices(search_by, search_val, search_val_to)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ESPAY Invoice"
        headers = ["No.", "Invoice/Due/ET Date", "Contract No.", "Customer Name",
                   "Amount", "Expired Date", "Payment Type", "Period", "Status", "Payment Date"]
        ws.append(headers)

        for i, row in enumerate(rows, start=1):
            amount = row["tot_amt"]
            try:
                amount = float(amount)
            except (TypeError, ValueError):
                pass
            ws.append([
                i, row["invoice_date"], row["contract_no"], row["customer_name"],
                amount, row["expired_date"], row["payment_type"],
                row["period"], row["invoice_status"], row["payment_date"],
            ])

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer,
            as_attachment=True,
            download_name="espay_invoice.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception:
        current_app.logger.exception("Espay export error")
        return jsonify({"message": "Failed to export invoices."}), 500